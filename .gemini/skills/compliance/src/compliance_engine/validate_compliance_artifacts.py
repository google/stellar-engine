#!/usr/bin/env python3
# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""
ATO Package Post-Generation Validator, OpenXML Inspector, & Public Sector Submission Engine

This script performs senior assessor verification across all compliance deliverables:
1. Text-based Markdown (.md) and YAML (.yaml) files.
2. OpenXML Word Policy Manuals (.docx) integrity & XML structure.
3. Macro-enabled Excel Workbooks (.xlsm) data validation & formatting audit.
4. Cross-format parity checking (YAML vs Excel, Markdown vs DOCX).
5. DISA STIG / SRG Checklist Evaluator referencing https://www.stigviewer.com/stigs
6. Complete Public Sector / ISSM ATO Submission Checklist & Operational Evidence Roadmap (ACAS, STIGs, ATC controls, PIA, ISA, SAAR, TTX, ATO Memo).
7. Option `--fix`: Re-runs dual-format generator to synchronize code drift.
8. Updates and compiles master `Path_to_Authorization.md` and `Path_to_Authorization.docx` inside `ato_artifacts/`.
"""

import argparse
from datetime import datetime
import logging
import os
from pathlib import Path
import re
import select
import shutil
import subprocess
import sys
from typing import Any, Dict, FrozenSet, List, Mapping, Optional, Sequence, Set, TextIO, Tuple, Union
import zipfile

from .baseline_registry import (
    REGIME_SPEC_BY_REGIME,
    audit_baseline_control_parity,
    build_regime_template_flags,
    get_regime_baseline,
)
from .file_helpers import (
    current_utc_datetime,
    UnsupportedRegimeError,
    classify_regime,
    ensure_path_within_boundary,
    is_dod_regime,
    is_supported_regime,
    require_supported_regime,
    resolve_cli_target,
    run_bounded_command,
    get_scripts_dir,
    get_skill_root,
    get_templates_dir,
    read_json_file,
    read_text_file,
    read_yaml_file,
    resolve_path,
    validate_system_inventory_schema,
    write_text_file,
)

from . import safe_xml as ET

from .audit_log import get_audit_logger, AuditEvent, AuditOutcome, audit_operation

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    openpyxl = None
    OPENPYXL_AVAILABLE = False

from . import docx_generator
from . import stig_resolver
from .oscal_generator import SUPPORTED_OSCAL_VERSIONS
from .template_engine import (
    PENDING_PLACEHOLDER_RE,
    TemplateEngine,
    _sanitize_markdown_inline_value,
    evaluate_template_conditionals,
)
from .semantic_linter import (
    evaluate_control_substance,
    run_semantic_linter,
)
from .poam_rules import (
    base_control_id,
    extract_normalized_control_ids,
    firewall_rule_exposes_ports,
    has_cmek_linkage,
    is_public_ingress_allow_rule,
    normalize_control_id,
    resolve_max_kms_rotation_days,
)

logger = logging.getLogger(__name__)

SKILL_BASE = str(get_skill_root())
TEMPLATES_DIR = str(get_templates_dir())
SCRIPTS_DIR = str(get_scripts_dir())

#: Data-driven baseline control expectations per regime across all 11 supported regimes.
REGIME_CONTROL_EXPECTATIONS: Dict[str, Dict[str, Any]] = {
    regime.value: {
        "min_ssp_controls": len(get_regime_baseline(regime)["control_ids"]),
        "min_sctm_xlsm_controls": len(get_regime_baseline(regime)["control_ids"]),
    }
    for regime in REGIME_SPEC_BY_REGIME
}

_GENERIC_LOGICAL_LABELS: FrozenSet[str] = frozenset({"default", "main", "this"})
_RUNBOOK_RUNTIME_TOKEN_RE = re.compile(r"<(?:[A-Z][A-Z0-9_]*|FILL IN)\s*:\s*[^>]+>", re.IGNORECASE)
_NIST_PARAMETER_SLOT_RE = re.compile(r"\[(?:assignment|selection)\s*:\s*[^\]]*\]", re.IGNORECASE)


def resolve_system_identity(sys_info: Mapping[str, Any]) -> Tuple[str, str]:
    """Returns canonical (system_name, system_abbreviation) with truthful config-required fallbacks."""
    name = str(sys_info.get("system_name") or "").strip() if isinstance(sys_info, Mapping) else ""
    abbr = str(sys_info.get("system_abbreviation") or "").strip() if isinstance(sys_info, Mapping) else ""
    return (
        name or "[CONFIG_REQUIRED: System Name]",
        abbr or "[CONFIG_REQUIRED: System Abbreviation]",
    )


def _find_pending_placeholders_in_line(line: str) -> List[str]:
    """Returns non-overlapping pending placeholders, runtime <TOKEN: ...> hints, and NIST [Assignment:/Selection:] slots."""
    if not line:
        return []
    spans: List[Tuple[int, int, str]] = []
    for rx in (PENDING_PLACEHOLDER_RE, _RUNBOOK_RUNTIME_TOKEN_RE, _NIST_PARAMETER_SLOT_RE):
        for m in rx.finditer(line):
            s, e = m.span()
            if not any(not (e <= es or s >= ee) for es, ee, _ in spans):
                spans.append((s, e, m.group(0)))
    spans.sort(key=lambda t: t[0])
    return [tok for _, _, tok in spans]


def _read_optional_lowercase(path: Union[str, Path], lower: bool = True) -> str:
    """Reads an optional deliverable, degrading to an empty string.

    An artifact that is genuinely absent and one that exists but cannot be read
    both yield ``""``, which would silently reduce reported coverage. The two
    cases are therefore distinguished in the log: absence is expected and
    recorded at DEBUG, while an existing-but-unreadable artifact is a real
    operator-actionable condition and is recorded at WARNING.

    Args:
        path: Filesystem path of the deliverable to read.
        lower: Whether to lowercase the contents for case-insensitive matching.

    Returns:
        The file contents, or ``""`` when unavailable.
    """
    if not os.path.exists(path):
        logger.debug("Optional deliverable '%s' is not present; treating as empty.", path)
        return ""
    try:
        text = read_text_file(str(path))
    except (OSError, ValueError) as err:
        logger.warning(
            "Deliverable '%s' exists but could not be read (%s); reconciliation coverage "
            "for this artifact will be under-reported.",
            path,
            err,
        )
        return ""
    return text.lower() if lower else text


# Text renditions of a deliverable that may carry asset cross-references.
# Binary renditions (.docx / .xlsm) are generated from these and are audited
# separately for OpenXML integrity, so they are deliberately excluded here.
_DELIVERABLE_TEXT_SUFFIXES: Tuple[str, ...] = (".md", ".yaml", ".yml", ".json")


def _read_deliverable_renditions(
    ato_dir: str, subfolder: str, basenames: List[str]
) -> str:
    """Reads every text rendition of a deliverable and returns their concatenation.

    A deliverable is emitted in several renditions (Markdown, YAML, JSON) and the
    asset-level evidence is not always present in all of them. Reading a single
    hard-coded filename therefore silently loses cross-references and
    under-reports reconciliation coverage, which understates how much of the
    architecture an assessor can actually trace through the package.

    Args:
        ato_dir: Absolute path to the ``ato_artifacts`` directory.
        subfolder: Deliverable subfolder name (for example ``"SSP"``).
        basenames: Candidate file basenames without extension, in preference
            order. Every candidate that exists is read, not just the first.

    Returns:
        The lowercased concatenation of every rendition found, or ``""`` when
        none of the candidates exist or are readable.
    """
    chunks: List[str] = []
    for base in basenames:
        for suffix in _DELIVERABLE_TEXT_SUFFIXES:
            candidate = os.path.join(ato_dir, subfolder, f"{base}{suffix}")
            text = _read_optional_lowercase(candidate)
            if text:
                chunks.append(text)
    if not chunks:
        logger.debug(
            "No readable rendition of deliverable '%s/%s' was found; assets that "
            "are documented only there will be reported as unreconciled.",
            subfolder,
            basenames,
        )
    return "\n".join(chunks)


def evaluate_disa_stig_applicability(
    inventory: Dict[str, Any],
    all_files: Optional[List[str]] = None,
    target_dir: Optional[str] = None,
    update_stigs: bool = False,
    stigs_mode: Optional[str] = None,
    stigs_catalog: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """Evaluates DISA STIG and SRG checklist applicability for the target system.

    Combines foundational cloud mission-owner STIGs with dynamic workload-discovered
    technology checklists, resolves active versions dynamically, and generates STIG Viewer
    reference URLs.

    Args:
        inventory: System inventory dictionary containing infrastructure metadata.
        all_files: Optional list of generated compliance artifact file paths.
        target_dir: Optional workspace directory containing compliance artifacts.
        update_stigs: If True, triggers active pulling of latest versions.
        stigs_mode: Resolution mode ('auto', 'online', or 'offline').
        stigs_catalog: Optional custom catalog source URL or file path.

    Returns:
        A list of dictionaries representing applicable STIG benchmarks with metadata.
    """
    stigs_cfg = inventory.get("disa_stigs") if isinstance(inventory, dict) else {}
    try:
        resolver = stig_resolver.StigResolver(
            target_dir=target_dir,
            config={"disa_stigs": stigs_cfg} if isinstance(stigs_cfg, dict) else None,
            update_mode=stigs_mode,
            catalog_source=stigs_catalog,
        )
        return resolver.evaluate_applicable_stigs(inventory, trigger_pull=update_stigs)
    except (OSError, ValueError, KeyError, RuntimeError) as e:
        logger.error("Failed to evaluate STIGs: %s", e)
        return [{
            "title": "STIG Evaluation Failed",
            "slug": "stig_evaluation_failed",
            "version": "N/A",
            "version_source": "Error",
            "category": "Error",
            "url": "",
            "cyber_exchange_url": "",
            "reason": "STIG resolver crashed during evaluation.",
            "focus": str(e),
            "action": "Check logs.",
            "status": "UNVERIFIED",
        }]


# Characters that the identifier-boundary rule treats as "inside a word". A
# candidate identifier only matches when the characters immediately before and
# after it are outside this class (or the string boundary).
_IDENTIFIER_WORD_CHARS = r"A-Za-z0-9_\-"
# All maximal runs of word characters, used to index a deliverable in one pass.
_WORD_RUN_RE = re.compile(rf"[{_IDENTIFIER_WORD_CHARS}]+")
# An identifier that consists solely of word characters (the common case).
_PURE_WORD_RUN_RE = re.compile(rf"[{_IDENTIFIER_WORD_CHARS}]+\Z")
# Strict IPv4 CIDR. The extractor may serialise `subnets_cidrs` as a stringified
# Python set (for example "{'10.10.0.0/20'}"), so the ranges are recovered with
# an explicit pattern rather than by trusting the container type. Anything that
# does not match is reported as undetermined rather than guessed at.
_CIDR_PATTERN = re.compile(r"\b\d{1,3}(?:\.\d{1,3}){3}/\d{1,2}\b")


class DeliverableIdentifierIndex:
    """Single-pass word-run index over one deliverable, for asset reconciliation.

    Reconciling discovered assets against a deliverable used to run one freshly
    built regex over the entire document per asset, which is O(assets x document
    bytes). Measured on a 441-resource fixture that cost 1.77 s for a single
    audit. This index tokenises the document once into its maximal runs of
    ``[A-Za-z0-9_-]`` and answers the overwhelmingly common case -- an
    identifier made only of word characters -- with an O(1) set membership test.

    The transformation is exact rather than approximate. If an identifier is
    entirely word characters, the boundary rule forces any match to coincide
    with a *maximal* word run, so set membership and the regex agree exactly.
    If an identifier contains separators (``.``, ``/``, ``:``), its leading word
    run must still coincide with a maximal word run of the document, so a failed
    membership test is a definitive non-match; only when that prefilter passes
    is the original regex evaluated. Matching therefore never becomes more
    permissive, which matters because a spurious match here would report an
    undocumented asset as reconciled.
    """

    __slots__ = ("_text", "_word_runs")

    def __init__(self, text: str) -> None:
        """Tokenises the deliverable text a single time.

        Args:
            text: Full deliverable contents. Callers that need case-insensitive
                reconciliation must pass already-lowercased text, matching the
                behaviour of the original matcher, which lowercased only the
                search target and never the document.
        """
        self._text = text or ""
        self._word_runs: FrozenSet[str] = (
            frozenset(_WORD_RUN_RE.findall(self._text)) if self._text else frozenset()
        )

    def contains(self, target: str) -> bool:
        """Reports whether an identifier appears in the deliverable as a whole word.

        Args:
            target: Candidate asset identifier, such as a VM name, bucket name,
                port number, or fully qualified container image reference.

        Returns:
            True when the identifier occurs delimited by non-word characters or
            string boundaries; False when it is absent, empty, or occurs only as
            a substring of a longer identifier.
        """
        if not target or not self._text:
            return False
        target_str = str(target).strip()
        if not target_str:
            return False
        needle = target_str.lower()
        if needle in _GENERIC_LOGICAL_LABELS:
            return False

        if _PURE_WORD_RUN_RE.match(needle):
            return needle in self._word_runs

        leading_run = _WORD_RUN_RE.match(needle)
        if leading_run is not None and leading_run.group(0) not in self._word_runs:
            # The identifier begins with word characters that do not form any
            # maximal word run in the document, so no boundary-delimited match
            # can exist. Skipping the scan here cannot hide a real match.
            return False

        escaped = re.escape(needle)
        pattern = (
            rf"(?:^|(?<=[^{_IDENTIFIER_WORD_CHARS}])){escaped}"
            rf"(?:$|(?=[^{_IDENTIFIER_WORD_CHARS}]))"
        )
        return bool(re.search(pattern, self._text))


def _reconcile_simple_asset_group(
    cat_key: str,
    items: Sequence[Any],
    asset_label: str,
    indexes: Tuple[DeliverableIdentifierIndex, ...],
    breakdown: Dict[str, Dict[str, int]],
    matched_assets: List[str],
    discrepancies: List[str],
    undoc_msg: str,
    empty_msg: Optional[str] = None,
    name_keys: Tuple[str, ...] = ("name",),
) -> None:
    """Reconciles a homogeneous asset collection against deliverable token indexes."""
    for item in items:
        breakdown[cat_key]["discovered"] += 1
        if isinstance(item, dict):
            name = next((str(item.get(k) or "").strip() for k in name_keys if item.get(k) and str(item.get(k)).strip()), "")
        else:
            name = str(item or "").strip()
        if name and any(idx.contains(name) for idx in indexes):
            breakdown[cat_key]["matched"] += 1
            matched_assets.append(f"{asset_label}: {name}")
        elif name:
            discrepancies.append(undoc_msg.format(name=name))
        elif empty_msg:
            discrepancies.append(empty_msg)


def _reconcile_subnets_and_service_accounts(
    infra: Mapping[str, Any],
    net: Mapping[str, Any],
    ssp_index: DeliverableIdentifierIndex,
    ppsm_index: DeliverableIdentifierIndex,
    hwsw_index: DeliverableIdentifierIndex,
    breakdown: Dict[str, Dict[str, int]],
    matched_assets: List[str],
    discrepancies: List[str],
) -> None:
    """Reconciles subnet CIDR ranges and multi-key IAM service account entries."""
    raw_cidrs = net.get("subnets_cidrs")
    cidr_source_text = (
        " ".join(str(c) for c in raw_cidrs)
        if isinstance(raw_cidrs, (list, tuple, set))
        else str(raw_cidrs or "").strip()
    )
    parsed_cidrs = sorted(set(_CIDR_PATTERN.findall(cidr_source_text)))
    if cidr_source_text and not parsed_cidrs:
        breakdown["subnets"]["discovered"] += 1
        discrepancies.append(
            f"Unreconciled Subnet Range: 'subnets_cidrs' is populated ({cidr_source_text!r}) "
            "but yields no parseable CIDR; subnet documentation coverage is [NOT DETERMINED FROM SOURCE]."
        )
    for cidr in parsed_cidrs:
        breakdown["subnets"]["discovered"] += 1
        if ssp_index.contains(cidr) or ppsm_index.contains(cidr) or hwsw_index.contains(cidr):
            breakdown["subnets"]["matched"] += 1
            matched_assets.append(f"Subnet CIDR Range: {cidr}")
        else:
            discrepancies.append(f"Undocumented Subnet CIDR Range: '{cidr}' not found in SSP, PPSM, or HW/SW inventory.")

    for sa in infra.get("service_accounts", []) or []:
        breakdown["service_accounts"]["discovered"] += 1
        if isinstance(sa, dict):
            candidates = [
                str(sa.get(k) or "").strip()
                for k in ("account_id", "name", "email", "resource_name")
                if sa.get(k) and str(sa.get(k)).strip()
            ]
        else:
            s_val = str(sa or "").strip()
            candidates = [s_val] if s_val else []
        sa_label = candidates[0] if candidates else ""
        if not sa_label:
            discrepancies.append(
                "Unreconciled IAM Service Account: service account discovered with no resolvable identifier "
                "(account_id/name/email empty)."
            )
        elif any(ssp_index.contains(c) or hwsw_index.contains(c) for c in candidates):
            breakdown["service_accounts"]["matched"] += 1
            matched_assets.append(f"Service Account: {sa_label}")
        else:
            discrepancies.append(f"Undocumented IAM Service Account: '{sa_label}' not found in SSP or HW/SW inventory.")


def audit_inventory_artifact_alignment(
    target_dir: str, inventory: Dict[str, Any], ato_dir: str
) -> Dict[str, Any]:
    """Reconciles discovered system architecture against generated ATO artifacts."""
    del target_dir
    if not inventory:
        inventory = {}
    infra = inventory.get("infrastructure_components") or {}
    net = inventory.get("network_architecture") or {}
    apps = inventory.get("application_components") or {}

    hwsw_index = DeliverableIdentifierIndex(_read_deliverable_renditions(ato_dir, "HW_SW_Inventory", ["Hardware_Software_Inventory"]))
    ppsm_index = DeliverableIdentifierIndex(_read_deliverable_renditions(ato_dir, "PPSM", ["PPSM_Ports_Protocols_Services"]))
    ssp_index = DeliverableIdentifierIndex(_read_deliverable_renditions(ato_dir, "SSP", ["SSP_System_Security_Plan", "System_Security_Plan"]))
    fips_index = DeliverableIdentifierIndex(
        _read_deliverable_renditions(ato_dir, "FIPS_Cryptography", ["FIPS_Cryptographic_Matrix", "FIPS_140_3_Cryptographic_Matrix"])
    )

    cat_keys = (
        "compute", "storage", "databases", "vpcs", "subnets", "firewall_rules", "kms_keys",
        "service_accounts", "cloud_run_services", "cloud_functions", "artifact_registries",
        "logging_sinks", "pubsub_topics", "nat_gateways", "service_perimeters",
        "exposed_ports", "applications", "container_images",
    )
    breakdown: Dict[str, Dict[str, int]] = {k: {"discovered": 0, "matched": 0} for k in cat_keys}
    discrepancies: List[str] = []
    matched_assets: List[str] = []

    specs: List[Tuple[str, Sequence[Any], str, Tuple[DeliverableIdentifierIndex, ...], str, Optional[str], Tuple[str, ...]]] = [
        ("compute", infra.get("compute_instances") or [], "Compute Instance", (hwsw_index, ssp_index),
         "Undocumented Compute Instance: '{name}' not found in HW/SW inventory or SSP.", None, ("name",)),
        ("compute", infra.get("gke_clusters") or [], "Kubernetes Cluster", (hwsw_index, ssp_index),
         "Undocumented Kubernetes Cluster: '{name}' not found in HW/SW inventory or SSP.", None, ("name",)),
        ("storage", infra.get("storage_buckets") or [], "Storage Bucket", (hwsw_index, ssp_index, fips_index),
         "Undocumented Storage Bucket: '{name}' not found in HW/SW inventory, SSP, or FIPS matrix.", None, ("name",)),
        ("databases", infra.get("databases") or [], "Database Instance", (hwsw_index, ssp_index, fips_index),
         "Undocumented Database: '{name}' not found in HW/SW inventory or SSP.", None, ("name",)),
        ("vpcs", net.get("vpcs") or [], "VPC Network", (hwsw_index, ssp_index),
         "Undocumented VPC Network: '{name}' not found in HW/SW inventory or SSP.", None, ("name",)),
        ("firewall_rules", net.get("firewall_rules") or [], "Firewall Rule", (ppsm_index, ssp_index, hwsw_index),
         "Undocumented Boundary Firewall Rule: '{name}' not found in PPSM, SSP, or HW/SW inventory.",
         "Unreconciled Boundary Firewall Rule: rule discovered with no resolvable name; identifier is [NOT DETERMINED FROM SOURCE].", ("name",)),
        ("kms_keys", infra.get("kms_keys") or [], "KMS Cryptographic Key", (fips_index, hwsw_index, ssp_index),
         "Undocumented KMS Cryptographic Key: '{name}' not found in FIPS matrix, HW/SW inventory, or SSP.",
         "Unreconciled KMS Cryptographic Key: key discovered with no resolvable name; identifier is [NOT DETERMINED FROM SOURCE].", ("name",)),
        ("cloud_run_services", infra.get("cloud_run_services") or [], "Cloud Run Service", (hwsw_index, ssp_index, ppsm_index),
         "Undocumented Cloud Run Service: '{name}' not found in SSP or HW/SW inventory.",
         "Unreconciled Cloud Run Service: resource discovered with no resolvable name.", ("name",)),
        ("cloud_functions", infra.get("cloud_functions") or [], "Cloud Function", (hwsw_index, ssp_index, ppsm_index),
         "Undocumented Cloud Function: '{name}' not found in SSP or HW/SW inventory.",
         "Unreconciled Cloud Function: resource discovered with no resolvable name.", ("name",)),
        ("artifact_registries", infra.get("artifact_registries") or [], "Artifact Registry", (hwsw_index, ssp_index),
         "Undocumented Artifact Registry: '{name}' not found in SSP or HW/SW inventory.",
         "Unreconciled Artifact Registry: resource discovered with no resolvable name.", ("name",)),
        ("logging_sinks", infra.get("logging_sinks") or [], "Cloud Logging Sink", (hwsw_index, ssp_index),
         "Undocumented Cloud Logging Sink: '{name}' not found in SSP or HW/SW inventory.",
         "Unreconciled Cloud Logging Sink: resource discovered with no resolvable name.", ("name",)),
        ("pubsub_topics", infra.get("pubsub_topics") or [], "Pub/Sub Topic", (hwsw_index, ssp_index),
         "Undocumented Pub/Sub Topic: '{name}' not found in SSP or HW/SW inventory.",
         "Unreconciled Pub/Sub Topic: resource discovered with no resolvable name.", ("name",)),
        ("nat_gateways", (infra.get("nat_gateways") or []) + (net.get("nat_gateways") or []), "Cloud NAT Gateway", (hwsw_index, ssp_index, ppsm_index),
         "Undocumented Cloud NAT Gateway: '{name}' not found in SSP or HW/SW inventory.",
         "Unreconciled Cloud NAT Gateway: resource discovered with no resolvable name.", ("name",)),
        ("service_perimeters", (infra.get("service_perimeters") or []) + (net.get("service_perimeters") or []), "VPC-SC Service Perimeter",
         (hwsw_index, ssp_index, ppsm_index), "Undocumented VPC-SC Service Perimeter: '{name}' not found in SSP or HW/SW inventory.",
         "Unreconciled VPC-SC Service Perimeter: resource discovered with no resolvable name.", ("name",)),
        ("exposed_ports", apps.get("exposed_ports") or net.get("application_ports") or [], "PPSM Port", (ppsm_index,),
         "Unmapped Boundary Ingress Port: '{name}' not found in PPSM Matrix.", None, ("port",)),
        ("applications", apps.get("applications") or [], "Application", (hwsw_index, ssp_index),
         "Undocumented Application: '{name}' not found in HW/SW inventory.", None, ("name",)),
        ("container_images", apps.get("container_images") or [], "Container Image", (hwsw_index, ssp_index),
         "Undocumented Container Image: '{name}' not found in HW/SW inventory.", None, ("image", "base_image", "name")),
    ]
    for cat_key, items_list, asset_label, indexes, undoc_msg, empty_msg, name_keys in specs:
        _reconcile_simple_asset_group(
            cat_key, items_list, asset_label, indexes, breakdown, matched_assets, discrepancies, undoc_msg, empty_msg, name_keys
        )

    _reconcile_subnets_and_service_accounts(infra, net, ssp_index, ppsm_index, hwsw_index, breakdown, matched_assets, discrepancies)

    total_discovered = sum(v["discovered"] for v in breakdown.values())
    total_matched = sum(v["matched"] for v in breakdown.values())
    coverage_score: Optional[float] = round((total_matched / total_discovered) * 100.0, 1) if total_discovered > 0 else None

    return {
        "total_discovered_assets": total_discovered,
        "reconciled_assets_count": total_matched,
        "coverage_score_percent": coverage_score,
        "breakdown": breakdown,
        "matched_assets": matched_assets,
        "discrepancies": discrepancies,
    }


def _format_coverage(alignment_res: Dict[str, Any]) -> str:
    """Renders reconciliation coverage for reports.

    Args:
        alignment_res: Result of :func:`audit_inventory_artifact_alignment`.

    Returns:
        ``"N/A (0 assets discovered)"`` when nothing was discovered, otherwise
        ``"<pct>% (<matched>/<discovered> assets)"``.
    """
    discovered = alignment_res.get("total_discovered_assets", 0) or 0
    matched = alignment_res.get("reconciled_assets_count", 0) or 0
    pct = alignment_res.get("coverage_score_percent")
    if not discovered or pct is None:
        return "N/A (0 assets discovered)"
    return f"{pct:.1f}% ({matched}/{discovered} assets)"


def _status_note(results: List[Dict[str, Any]], ok_text: str) -> str:
    """Summarizes per-deliverable audit statuses without overstating them.

    Args:
        results: Audit result dictionaries carrying a ``status`` key.
        ok_text: What a PASS means for this deliverable type.

    Returns:
        ``"<passed>/<total> passed: <ok_text>"``, plus the count of results
        that failed or could not be read.
    """
    total = len(results)
    passed = sum(1 for r in results if r.get("status") == "PASS")
    failed = sum(1 for r in results if r.get("status") in ("FAIL", "UNVERIFIED", "ERROR"))
    note = f"{passed}/{total} passed: {ok_text}"
    if failed:
        note += f"; {failed} failed or unreadable"
    return note


# A control heading in the SSP, anchored so that a request for "AC-17" cannot
# latch onto "### AC-17(2)". The lookahead rejects any character that could
# extend the control identifier.
_SSP_CONTROL_HEADING_TEMPLATE = r"^###[ \t]+{ctrl}(?![A-Za-z0-9_\-(])[^\n]*\n"
# Any subsequent level-3 heading, which terminates the current control section.
_SSP_NEXT_HEADING_RE = re.compile(r"^###[ \t]", re.MULTILINE)
# A GitHub-flavoured task-list checkbox. The SSP renders the implementation
# status list inside a Markdown table cell with "<br>" separators, so the label
# is terminated by "<", "|" or a newline rather than by the end of a line.
_SSP_STATUS_CHECKBOX_RE = re.compile(r"-\s*\[(?P<mark>[ xX])\]\s*(?P<label>[^<|\n]+)")
# NIST catalog assignment placeholders. Their presence means the control
# statement was never tailored to this system, so it cannot substantiate an
# implementation claim.
_UNTAILORED_PARAMETER_RE = re.compile(r"\[assignment:|\[selection:", re.IGNORECASE)


def _extract_ssp_control_section(ssp_text: str, ctrl_id: str) -> str:
    """Extracts the body of a single control section from the SSP Markdown.

    The section is located by an anchored heading match rather than a substring
    search. A substring search for ``"### AC-17"`` also matches
    ``"### AC-17(2)"``, which either attributes another control's narrative to
    this one or, in the original implementation, silently reused the previous
    loop iteration's section as this control's evidence.

    Args:
        ssp_text: Full System Security Plan Markdown, in original case.
        ctrl_id: NIST control identifier, for example ``"IA-2(1)"``.

    Returns:
        The section body with its heading line removed and surrounding
        whitespace stripped, or ``""`` when the control has no section.
    """
    if not ssp_text or not ctrl_id:
        return ""
    heading_re = re.compile(
        _SSP_CONTROL_HEADING_TEMPLATE.format(ctrl=re.escape(ctrl_id)),
        re.MULTILINE,
    )
    match = heading_re.search(ssp_text)
    if match is None:
        return ""
    body_start = match.end()
    next_heading = _SSP_NEXT_HEADING_RE.search(ssp_text, body_start)
    body_end = next_heading.start() if next_heading else len(ssp_text)
    return ssp_text[body_start:body_end].strip()


def _parse_ssp_implementation_status(section: str) -> Tuple[FrozenSet[str], FrozenSet[str]]:
    """Parses the SSP 'Implementation Status (check all that apply)' checkboxes.

    The status must be read from the checkbox state, not from the presence of a
    word. Every control section lists all five candidate statuses, so a plain
    substring test for ``"planned"`` matches the *unchecked* ``- [ ] Planned``
    box in every section and permanently suppresses verification.

    Args:
        section: A single control section body from the SSP.

    Returns:
        A ``(checked, unchecked)`` pair of lowercased status labels. Both are
        empty when the section contains no recognisable status block, which the
        caller must treat as "status not determined" rather than as success.
    """
    checked: Set[str] = set()
    unchecked: Set[str] = set()
    for match in _SSP_STATUS_CHECKBOX_RE.finditer(section or ""):
        label = match.group("label").strip().rstrip("|").strip().lower()
        if not label:
            continue
        if match.group("mark").lower() == "x":
            checked.add(label)
        else:
            unchecked.add(label)
    return frozenset(checked), frozenset(unchecked)


_DANGEROUS_PUBLIC_PORTS: Dict[int, str] = {
    21: "FTP",
    23: "Telnet",
    5432: "PostgreSQL",
    3306: "MySQL",
    1433: "SQL Server",
    1521: "Oracle",
    27017: "MongoDB",
    6379: "Redis",
}
_ADMIN_PUBLIC_PORTS: Dict[int, str] = {22: "SSH", 3389: "RDP"}


def evaluate_firewall_exposure(inventory: Dict[str, Any]) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Returns (CAT I, CAT II) findings for public ingress to sensitive, admin and HTTP ports."""
    cat_1: List[Dict[str, Any]] = []
    cat_2: List[Dict[str, Any]] = []
    net = inventory.get("network_architecture") or {}
    for rule in net.get("firewall_rules", []) or []:
        if not is_public_ingress_allow_rule(rule):
            continue
        r_name = rule.get("name", "unnamed-firewall-rule")
        for p_int in firewall_rule_exposes_ports(rule, set(_DANGEROUS_PUBLIC_PORTS)):
            cat_1.append({
                "id": "CAT1-FW-001",
                "component": f"Firewall Rule '{r_name}'",
                "severity": "CAT I (Critical)",
                "description": f"Direct 0.0.0.0/0 ingress allowed to {_DANGEROUS_PUBLIC_PORTS[p_int]} database/sensitive port ({p_int}).",
                "remediation": "Restrict source IP range to private subnet CIDRs or deploy Cloud IAP zero-trust proxy bastions.",
            })
        for p_int in firewall_rule_exposes_ports(rule, set(_ADMIN_PUBLIC_PORTS)):
            cat_1.append({
                "id": "CAT1-FW-002",
                "component": f"Firewall Rule '{r_name}'",
                "severity": "CAT I (Critical)",
                "description": (
                    f"Direct 0.0.0.0/0 ingress allowed to administrative remote access port "
                    f"({_ADMIN_PUBLIC_PORTS[p_int]} port {p_int})."
                ),
                "remediation": "Enforce Google Cloud IAP (35.235.240.0/20) and remove public 0.0.0.0/0 ingress.",
            })
        if firewall_rule_exposes_ports(rule, {80}):
            cat_2.append({
                "id": "CAT2-FW-001",
                "component": f"Firewall Rule '{r_name}'",
                "severity": "CAT II (Medium)",
                "description": "Plaintext HTTP port 80 exposed to 0.0.0.0/0.",
                "remediation": "Configure automated HTTP-to-HTTPS redirect (TLS 1.3) on target proxy / load balancer."
            })
    return cat_1, cat_2


def _load_sctm_controls_map(
    sctm_yaml: str,
    ato_dir: str,
    cat_1_findings: List[Dict[str, str]],
) -> Dict[str, Dict[str, Any]]:
    """Loads the SCTM control map from YAML, recording a CAT I finding on syntax errors."""
    sctm_controls_map: Dict[str, Dict[str, Any]] = {}
    if not os.path.exists(sctm_yaml):
        return sctm_controls_map
    try:
        sctm_data = read_yaml_file(sctm_yaml, allowed_boundary=ato_dir)
        for fam in sctm_data.get("control_families", []) or []:
            for ctrl in fam.get("controls", []) or []:
                c_id = ctrl.get("id")
                if c_id:
                    sctm_controls_map[str(c_id)] = ctrl
                    if normalize_control_id is not None:
                        sctm_controls_map[normalize_control_id(str(c_id))] = ctrl
    except (OSError, ValueError, TypeError) as err:
        logger.error("Failed to parse SCTM YAML at '%s': %s", sctm_yaml, err)
        cat_1_findings.append({
            "id": "CAT1-SCTM-PARSE",
            "component": "SCTM_Burndown_Matrix.yaml",
            "severity": "CAT I (Critical)",
            "description": f"Failed to parse SCTM YAML, content UNVERIFIED: {err}",
            "remediation": "Fix YAML syntax.",
        })
    return sctm_controls_map


def _audit_cmek_and_key_rotation(
    inventory: Dict[str, Any],
    infra: Mapping[str, Any],
    cat_2_findings: List[Dict[str, str]],
) -> None:
    """Audits CMEK key presence and opt-in cryptoperiod rotation policies."""
    kms_keys = infra.get("kms_keys", []) or []
    has_persistent_storage_or_kms = (
        any("kms" in str(s).lower() for s in infra.get("services_enabled", []) or [])
        or bool(infra.get("storage_buckets"))
        or bool(infra.get("databases"))
        or bool(infra.get("compute_instances"))
        or bool(infra.get("gke_clusters"))
    )
    if not has_cmek_linkage(infra) and has_persistent_storage_or_kms:
        cat_2_findings.append({
            "id": "CAT2-CRYPTO-001",
            "component": "Cloud KMS Key Rings",
            "severity": "CAT II (Medium)",
            "description": "No customer-managed encryption keys (CMEK) discovered in active inventory.",
            "remediation": "Provision Cloud KMS CMEK key rings and associate with persistent storage resources.",
        })
        return

    max_rot_days = resolve_max_kms_rotation_days(inventory)
    if max_rot_days is None:
        return
    max_rot_sec = max_rot_days * 86400
    for k in kms_keys:
        if not isinstance(k, dict):
            continue
        rot = str(k.get("rotation_period", "")).strip()
        k_name = k.get("name", "kms-key")
        if not rot or rot in ("null", "none"):
            cat_2_findings.append({
                "id": "CAT2-CRYPTO-002",
                "component": f"Cloud KMS Key '{k_name}'",
                "severity": "CAT II (Medium)",
                "description": (
                    f"KMS key '{k_name}' missing automated key rotation policy required by "
                    f"organization-defined cryptoperiod ({max_rot_days} days per NIST SP 800-57 Part 1)."
                ),
                "remediation": f"Configure automated key rotation ({max_rot_sec}s) in Terraform.",
            })
        else:
            m = re.match(r"^(\d+)s?$", rot)
            if m and int(m.group(1)) > max_rot_sec:
                cat_2_findings.append({
                    "id": "CAT2-CRYPTO-003",
                    "component": f"Cloud KMS Key '{k_name}'",
                    "severity": "CAT II (Medium)",
                    "description": (
                        f"KMS key '{k_name}' rotation period ({rot}) exceeds organization-defined "
                        f"{max_rot_days}-day cryptoperiod (NIST SP 800-57 Part 1)."
                    ),
                    "remediation": f"Reduce key rotation period to {max_rot_days} days ({max_rot_sec}s).",
                })


def _assess_ssp_atc_section(ssp_section: str) -> Tuple[str, str, bool]:
    """Evaluates an SSP control section for implementation status and placeholder hygiene."""
    affirming = ("implemented", "inherited")
    blocking = ("not implemented", "partially implemented", "planned", "not applicable")
    details = ssp_section[:160].replace("\n", " ").strip() + "..."
    checked, unchecked = _parse_ssp_implementation_status(ssp_section)
    affirmed = sorted(
        label for label in checked
        if any(a in label for a in affirming) and not any(b in label for b in blocking)
    )
    blocked = sorted(label for label in checked if any(b in label for b in blocking))
    if not checked and not unchecked:
        return "Documented in SSP (implementation status [NOT DETERMINED FROM SOURCE])", details, False
    if blocked or not affirmed:
        checked_desc = ", ".join(sorted(checked)) if checked else "no status box checked"
        return f"Documented in SSP (implementation status: {checked_desc})", details, False
    if len(ssp_section) < 50:
        return "Documented in SSP (implementation narrative too short to assess)", details, False
    if PENDING_PLACEHOLDER_RE.search(ssp_section) or "{{" in ssp_section:
        return "Documented in SSP (narrative contains unresolved placeholders)", details, False
    if _UNTAILORED_PARAMETER_RE.search(ssp_section):
        return "Documented in SSP (untailored NIST assignment/selection parameters remain)", details, False
    return f"Implemented in SSP ({', '.join(affirmed)})", details, True


def _cross_check_atc_posture(
    ctrl_id: str,
    status: str,
    details: str,
    cat_1_findings: List[Dict[str, str]],
    cat_2_findings: List[Dict[str, str]],
    semantic_linter_report: Optional[Any],
) -> Tuple[bool, str, str]:
    """Cross-checks an ATC control against live firewall/CMEK and semantic drift findings."""
    if ctrl_id == "SC-7":
        boundary_findings = [f for f in cat_1_findings if "FW" in f.get("id", "")]
        if boundary_findings:
            return False, "Boundary Ingress Exposure (CAT 1 Finding)", f"Boundary violation: {boundary_findings[0]['description']}"
    elif ctrl_id == "AC-17":
        admin_ingress = [f for f in cat_1_findings if "CAT1-FW-002" in f.get("id", "")]
        if admin_ingress:
            return False, "Insecure Remote Admin Ingress (CAT 1 Finding)", f"Remote access exposure: {admin_ingress[0]['description']}"
    elif ctrl_id == "SC-28":
        cmek_missing = [f for f in cat_2_findings if "CAT2-CRYPTO-001" in f.get("id", "")]
        if cmek_missing:
            return False, "CMEK Cryptographic Key Absent (CAT 2 Finding)", "Persistent data stores lack Customer-Managed Encryption Keys."
    if semantic_linter_report is not None:
        norm_ctrl = normalize_control_id(ctrl_id)
        base_ctrl = norm_ctrl.split("(")[0]
        for df in getattr(semantic_linter_report, "drift_findings", []) or []:
            df_ctrl = normalize_control_id(str(df.control_id))
            if df_ctrl in (norm_ctrl, base_ctrl) or df_ctrl.split("(")[0] == base_ctrl:
                return False, f"Architectural Drift ({df.finding_id})", df.description
    return True, status, details


def _applicable_atc_controls(inventory: Dict[str, Any]) -> List[Tuple[str, str, str]]:
    """Returns the subset of ATC_CONNECTION_CONTROLS applicable to the active regime."""
    sys_info = (inventory.get("system_information") or {}) if isinstance(inventory, dict) else {}
    active_regime = classify_regime(sys_info)
    active_baseline_ids = (
        {normalize_control_id(c) for c in get_regime_baseline(active_regime)["control_ids"]}
        if is_supported_regime(active_regime)
        else None
    )
    dod_enclave = is_dod_regime(sys_info)
    selected: List[Tuple[str, str, str]] = []
    for ctrl_id, ctrl_name, ctrl_desc in ATC_CONNECTION_CONTROLS:
        norm_atc_id = normalize_control_id(ctrl_id)
        if norm_atc_id in ("AU-4(1)", "CA-3(6)") and not dod_enclave:
            continue
        if active_baseline_ids is not None and norm_atc_id not in active_baseline_ids:
            continue
        selected.append((ctrl_id, ctrl_name, ctrl_desc))
    return selected


def _audit_atc_connection_controls(
    sctm_controls_map: Dict[str, Dict[str, Any]],
    ssp_text: str,
    inventory: Dict[str, Any],
    cat_1_findings: List[Dict[str, str]],
    cat_2_findings: List[Dict[str, str]],
    semantic_linter_report: Optional[Any],
) -> Tuple[List[Dict[str, Any]], int]:
    """Verifies the applicable ATC connection controls across SCTM and SSP."""
    records: List[Dict[str, Any]] = []
    verified_count = 0
    absent_controls: List[str] = []

    for ctrl_id, ctrl_name, ctrl_desc in _applicable_atc_controls(inventory):
        ctrl_entry = sctm_controls_map.get(ctrl_id)
        if ctrl_entry is None and normalize_control_id is not None:
            ctrl_entry = sctm_controls_map.get(normalize_control_id(ctrl_id))
        is_substantive = False
        evidence_text = ""
        evidence_source = "None"

        if ctrl_entry:
            evidence_source = "SCTM"
            status = ctrl_entry.get("status") or "Documented in SCTM"
            details = ctrl_entry.get("implementation_details") or ""
            evidence_text = details
            stat_lower = str(status).lower()
            is_impl = (
                any(stat_lower.startswith(p) for p in ("implemented", "inherited", "satisfied", "automated"))
                and "not implemented" not in stat_lower
                and "planned" not in stat_lower
            )
            is_substantive = bool(is_impl and len(details) >= 25 and not PENDING_PLACEHOLDER_RE.search(details) and "{{" not in details)
        else:
            ssp_section = _extract_ssp_control_section(ssp_text, ctrl_id)
            if ssp_section:
                evidence_source = "SSP"
                evidence_text = ssp_section
                status, details, is_substantive = _assess_ssp_atc_section(ssp_section)
            else:
                status = "Absent from SCTM and SSP"
                details = (
                    "[NOT DETERMINED FROM SOURCE] No implementation statement for this mandatory "
                    "ATC connection control exists in the SCTM or the SSP."
                )

        if is_substantive:
            if evaluate_control_substance is None:
                is_substantive = False
                status = "Not verified (narrative substance check unavailable: semantic_linter not importable)"
            else:
                sub_ok, sub_status, sub_reason = evaluate_control_substance(ctrl_id=ctrl_id, narrative=evidence_text, inventory=inventory)
                if not sub_ok:
                    is_substantive = False
                    status = f"Incomplete Technical Narrative ({sub_status})"
                    details = f"{details} [Assessor Finding: {sub_reason}]"

        if is_substantive:
            is_substantive, status, details = _cross_check_atc_posture(
                ctrl_id, status, details, cat_1_findings, cat_2_findings, semantic_linter_report
            )

        if not details:
            details = "[NOT DETERMINED FROM SOURCE] No implementation narrative recorded for this control."

        if is_substantive:
            ver_status = "VERIFIED"
            verified_count += 1
        elif evidence_source == "None":
            ver_status = "MISSING"
            absent_controls.append(ctrl_id)
        else:
            ver_status = "PARTIAL"

        records.append({
            "id": ctrl_id,
            "name": ctrl_name,
            "description": ctrl_desc,
            "status": status,
            "details": details,
            "verification": ver_status,
            "evidence_source": evidence_source,
        })

    if absent_controls:
        cat_2_findings.append({
            "id": "CAT2-ATC-001",
            "component": "ATC Connection Controls (SCTM / SSP)",
            "severity": "CAT II (Medium)",
            "description": (
                f"{len(absent_controls)} of the {len(ATC_CONNECTION_CONTROLS)} ATC connection controls have no "
                f"implementation statement in either the SCTM or the SSP: {', '.join(absent_controls)}. "
                "Their implementation status is [NOT DETERMINED FROM SOURCE]."
            ),
            "remediation": (
                "Author implementation statements for the listed controls in the SCTM (or the SSP) "
                "before requesting an Authorization to Connect; an ATC cannot be granted on undocumented controls."
            ),
        })
    return records, verified_count


def _audit_poam_currency_and_placeholders(
    poam_yaml: str,
    ato_dir: str,
    unresolved_tokens: List[Dict[str, Any]],
    config_required_vars: List[Dict[str, Any]],
    disc_list: List[str],
    cat_1_findings: List[Dict[str, str]],
    cat_2_findings: List[Dict[str, str]],
) -> None:
    """Audits unresolved macro tokens, pending config variables, asset discrepancies, and POA&M milestone currency."""
    if unresolved_tokens:
        cat_1_findings.append({
            "id": "CAT1-SYNTAX-001",
            "component": "ATO Deliverable Templates",
            "severity": "CAT I (Critical)",
            "description": f"{len(unresolved_tokens)} unresolved template macro tokens ({{{{ ... }}}}) detected in generated deliverables.",
            "remediation": "Re-run compliance generator to hydrate all template tokens prior to formal package submission.",
        })
    if config_required_vars:
        cat_2_findings.append({
            "id": "CAT2-GOV-001",
            "component": "compliance_config.yaml",
            "severity": "CAT II (Medium)",
            "description": f"{len(config_required_vars)} pending institutional configuration variables requiring human assignment.",
            "remediation": "Fill in lead personnel roles and governance identifiers in compliance_config.yaml.",
        })
    if disc_list:
        sample_disc = "; ".join(disc_list[:2])
        cat_2_findings.append({
            "id": "CAT2-INVENTORY-001",
            "component": "Hardware & Software Inventory",
            "severity": "CAT II (Medium)",
            "description": f"{len(disc_list)} discovered system assets not fully cross-referenced in ATO deliverables ({sample_disc}).",
            "remediation": "Synchronize Hardware/Software inventory with latest system_inventory.json via --fix.",
        })

    if not os.path.exists(poam_yaml):
        return
    try:
        poam_data = read_yaml_file(poam_yaml, allowed_boundary=ato_dir)
        today_d = current_utc_datetime().replace(tzinfo=None).date()
        overdue_items: List[str] = []
        unparseable_dates: List[str] = []
        for item in poam_data.get("poam_items", []) or []:
            if not isinstance(item, dict) or item.get("status") != "Ongoing":
                continue
            sched = item.get("scheduled_completion_date")
            if not sched:
                continue
            try:
                d = datetime.strptime(str(sched).strip()[:10], "%Y-%m-%d").date()
            except (ValueError, TypeError) as date_err:
                logger.warning("POA&M item '%s' has an unparseable scheduled_completion_date %r: %s", item.get("item_id"), sched, date_err)
                unparseable_dates.append(f"{item.get('item_id')} ({sched!r})")
                continue
            if d <= today_d:
                overdue_items.append(f"{item.get('item_id')} (due {sched})")
        if overdue_items:
            cat_2_findings.append({
                "id": "CAT2-POAM-001",
                "component": "Plan of Action & Milestones (POA&M)",
                "severity": "CAT II (Medium)",
                "description": f"POA&M contains {len(overdue_items)} overdue Ongoing milestones past scheduled date: {', '.join(overdue_items[:2])}.",
                "remediation": "Update milestone progress or request Authorizing Official extension in eMASS.",
            })
        if unparseable_dates:
            cat_2_findings.append({
                "id": "CAT2-POAM-002",
                "component": "Plan of Action & Milestones (POA&M)",
                "severity": "CAT II (Medium)",
                "description": (
                    f"{len(unparseable_dates)} Ongoing POA&M milestones have an unparseable "
                    f"scheduled completion date and could not be assessed for currency: {', '.join(unparseable_dates[:2])}."
                ),
                "remediation": "Correct the affected dates to ISO 8601 (YYYY-MM-DD) so milestone currency can be verified.",
            })
    except (OSError, ValueError, TypeError) as err:
        logger.error("Failed evaluating POA&M milestones at '%s': %s", poam_yaml, err)
        cat_1_findings.append({
            "id": "CAT1-POAM-PARSE",
            "component": "Plan_of_Action_and_Milestones.yaml",
            "severity": "CAT I (Critical)",
            "description": f"Failed to parse POA&M YAML, content UNVERIFIED: {err}",
            "remediation": "Fix YAML syntax.",
        })


def _audit_dod_telemetry_boundary(
    sys_info: Mapping[str, Any],
    inventory: Dict[str, Any],
    sctm_controls_map: Dict[str, Dict[str, Any]],
    cat_2_findings: List[Dict[str, str]],
    cat_3_findings: List[Dict[str, str]],
) -> None:
    """Checks DoD IL4/IL5 telemetry boundary rules against unaccredited SCC usage."""
    if not (is_dod_regime(sys_info) or is_dod_regime(sys_info.get("compliance_framework"))):
        return
    allow_unaccredited = (
        sys_info.get("allow_unaccredited_scc_in_il5") is True
        or (inventory.get("security_operations") or {}).get("allow_unaccredited_scc_in_il5") is True
    )
    ext_keywords = ("cssp", "siem", "export sink", "log router", "external", "bridge", "cloud logging", "chronicle", "secops")
    for t_id in ("AU-6", "CA-7", "SI-4", "IR-4"):
        ctrl_rec = sctm_controls_map.get(t_id)
        if not ctrl_rec:
            continue
        combined_txt = f"{ctrl_rec.get('implementation_details') or ''} {ctrl_rec.get('codebase_evidence') or ''}".lower()
        if not any(s in combined_txt for s in ("security command center", " scc ", "scc premium")):
            continue
        if any(k in combined_txt for k in ext_keywords):
            continue
        if allow_unaccredited:
            cat_3_findings.append({
                "id": f"CAT3-DOD-TEL-{t_id}-ETP",
                "component": f"Telemetry Control '{t_id}' (AO Exception-to-Policy)",
                "severity": "CAT III (Advisory)",
                "description": (
                    f"Control {t_id} utilizes Security Command Center in DoD IL4/IL5 under an approved "
                    "Authorizing Official (AO) Exception-to-Policy (allow_unaccredited_scc_in_il5: true)."
                ),
                "remediation": "Maintain signed AO ETP memorandum in eMASS package artifacts.",
            })
        else:
            cat_2_findings.append({
                "id": f"CAT2-DOD-TEL-{t_id}",
                "component": f"Telemetry Control '{t_id}'",
                "severity": "CAT II (Medium)",
                "description": (
                    f"Control {t_id} references Security Command Center (SCC) without routing via Cloud Logging "
                    "export sinks to external accredited CSSP/SIEM. Centralized telemetry export to an accredited "
                    "CSSP/SIEM is required per DoDI 8530.01."
                ),
                "remediation": (
                    "Route telemetry and audit logs via Cloud Logging export sinks to an accredited external "
                    "CSSP/SIEM (e.g., DISA, C5ISR) or note an accredited cross-boundary bridge "
                    "(or set allow_unaccredited_scc_in_il5: true with AO Exception-to-Policy)."
                ),
            })


def _merge_deliverable_and_linter_findings(
    broken_yaml: List[Dict[str, Any]],
    inventory_corrupt: bool,
    inventory_error: str,
    linter_crashed: bool,
    linter_crash_error: str,
    linter_strict: bool,
    semantic_linter_report: Optional[Any],
    cat_1_findings: List[Dict[str, str]],
    cat_2_findings: List[Dict[str, str]],
) -> None:
    """Merges YAML parse failures, inventory errors, and semantic linter CAT I findings into assessor findings."""
    for b_y in broken_yaml:
        cat_1_findings.append({
            "id": "CAT1-YAML-001",
            "component": os.path.basename(b_y.get("file", "YAML Deliverable")),
            "severity": "CAT I (Critical)",
            "description": f"Malformed YAML syntax in deliverable '{b_y.get('file')}': {b_y.get('error', 'Parse error')}",
            "remediation": "Correct YAML syntax to ensure valid structured data for automated ingestion and compliance auditing.",
            "source": "assessor",
        })
    if inventory_corrupt:
        cat_1_findings.append({
            "id": "CAT1-INVENTORY-001",
            "component": "system_inventory.json",
            "severity": "CAT I (Critical)",
            "description": f"Corrupt or unreadable system_inventory.json: {inventory_error}",
            "remediation": "Re-run extract_system_data.py to regenerate a valid system_inventory.json.",
            "source": "assessor",
        })
    if linter_crashed:
        if linter_strict:
            cat_1_findings.append({
                "id": "CAT1-LINTER-CRASH",
                "component": "Semantic Linter",
                "severity": "CAT I (Critical)",
                "description": f"Semantic linter raised an unexpected error: {linter_crash_error}",
                "remediation": "Inspect logs and fix the artifact or inventory defect causing the linter failure.",
                "source": "semantic_linter",
            })
        else:
            cat_2_findings.append({
                "id": "CAT2-LINTER-CRASH",
                "component": "Semantic Linter",
                "severity": "CAT II (Medium)",
                "description": f"Semantic linter raised an error (non-blocking under --no-linter-strict): {linter_crash_error}",
                "remediation": "Inspect logs and fix the linter error.",
                "source": "semantic_linter",
            })
    if semantic_linter_report is None:
        return
    existing_ids = {f.get("id") for f in cat_1_findings}
    for lf in getattr(semantic_linter_report, "all_findings", []) or []:
        sev_str = str(getattr(lf, "severity", "")).upper()
        if re.search(r"\bCAT\s*I\b", sev_str) and not re.search(r"\bCAT\s*(?:II|III)\b", sev_str):
            if lf.finding_id not in existing_ids:
                cat_1_findings.append({
                    "id": lf.finding_id,
                    "component": lf.artifact,
                    "severity": lf.severity,
                    "description": lf.description,
                    "remediation": lf.remediation,
                    "source": "semantic_linter",
                })
                existing_ids.add(lf.finding_id)


def audit_senior_compliance_quality(
    target_dir: str,
    inventory: Dict[str, Any],
    ato_dir: str,
    alignment_res: Dict[str, Any],
    excel_results: List[Dict[str, Any]],
    docx_results: List[Dict[str, Any]],
    oscal_results: List[Dict[str, Any]],
    unresolved_tokens: List[Dict[str, Any]],
    config_required_vars: List[Dict[str, Any]],
    stigs_required: Optional[List[Dict[str, Any]]] = None,
    rmf_action_items: Optional[List[Dict[str, Any]]] = None,
    yaml_results: Optional[List[Dict[str, Any]]] = None,
    semantic_linter_report: Optional[Any] = None,
    linter_strict: bool = True,
    linter_crashed: bool = False,
    linter_crash_error: str = "",
    inventory_corrupt: bool = False,
    inventory_error: str = "",
) -> Dict[str, Any]:
    """Performs Level 4 Senior Compliance Assessor Quality & Security Posture Audit."""
    del target_dir
    if not inventory:
        inventory = {}
    infra = inventory.get("infrastructure_components") or {}
    sys_info = inventory.get("system_information") or {}

    sctm_yaml = os.path.join(ato_dir, "SCTM", "SCTM_Burndown_Matrix.yaml")
    poam_yaml = os.path.join(ato_dir, "POAM", "Plan_of_Action_and_Milestones.yaml")
    ssp_md = os.path.join(ato_dir, "SSP", "SSP_System_Security_Plan.md")
    if not os.path.exists(ssp_md):
        ssp_md = os.path.join(ato_dir, "SSP", "System_Security_Plan.md")

    cat_1_findings: List[Dict[str, str]] = []
    cat_2_findings: List[Dict[str, str]] = []
    cat_3_findings: List[Dict[str, str]] = []

    sctm_controls_map = _load_sctm_controls_map(sctm_yaml, ato_dir, cat_1_findings)
    ssp_text = _read_optional_lowercase(ssp_md, lower=False)

    fw_cat_1, fw_cat_2 = evaluate_firewall_exposure(inventory)
    cat_1_findings.extend(fw_cat_1)
    cat_2_findings.extend(fw_cat_2)

    _audit_cmek_and_key_rotation(inventory, infra, cat_2_findings)
    atc_records, verified_atc_count = _audit_atc_connection_controls(
        sctm_controls_map, ssp_text, inventory, cat_1_findings, cat_2_findings, semantic_linter_report
    )
    _audit_poam_currency_and_placeholders(
        poam_yaml, ato_dir, unresolved_tokens, config_required_vars,
        alignment_res.get("discrepancies", []), cat_1_findings, cat_2_findings,
    )
    cat_1_findings.extend(audit_baseline_control_parity(ato_dir, inventory))
    _audit_dod_telemetry_boundary(sys_info, inventory, sctm_controls_map, cat_2_findings, cat_3_findings)

    if stigs_required:
        cat_3_findings.append({
            "id": "CAT3-STIG-001",
            "component": "DISA STIG Checklists (.ckl)",
            "severity": "CAT III (Advisory)",
            "description": f"{len(stigs_required)} DISA STIG/SRG checklists identified for completion in desktop STIG Viewer.",
            "remediation": "Download official STIG compilation from DoD Cyber Exchange and export completed .ckl files.",
        })
    if rmf_action_items:
        cat_3_findings.append({
            "id": "CAT3-GOV-001",
            "component": "RMF Action Alerts",
            "severity": "CAT III (Advisory)",
            "description": f"{len(rmf_action_items)} policy action highlights require review by appointed RMF roles.",
            "remediation": "Review highlighted action items ([WARNING: ...] / [INFORMATIONAL: ...]) in Policies_and_Procedures/.",
        })

    failing_statuses = ("FAIL", "UNVERIFIED")
    broken_excel = [r for r in excel_results if r.get("status") in failing_statuses]
    broken_docx = [r for r in docx_results if r.get("status") in failing_statuses]
    broken_oscal = [r for r in oscal_results if r.get("status") in failing_statuses]
    broken_yaml = [r for r in (yaml_results or []) if r.get("status") == "FAIL"]

    _merge_deliverable_and_linter_findings(
        broken_yaml, inventory_corrupt, inventory_error, linter_crashed, linter_crash_error,
        linter_strict, semantic_linter_report, cat_1_findings, cat_2_findings,
    )

    broken_deliverables = broken_excel + broken_docx + broken_oscal + broken_yaml
    blocking_cat_1 = [f for f in cat_1_findings if f.get("source") != "semantic_linter" or linter_strict]

    atc_total_scoped = len(atc_records) or len(ATC_CONNECTION_CONTROLS)
    if blocking_cat_1 or broken_deliverables or unresolved_tokens or inventory_corrupt:
        recommendation = "Remediation Required Prior to Assessment Submission"
        rec_badge = "caution"
    elif verified_atc_count == atc_total_scoped and not cat_2_findings and not cat_1_findings and not linter_crashed:
        recommendation = "Technical Package Complete - Ready for Independent Assessor (3PAO / SCA) Review"
        rec_badge = "success"
    else:
        recommendation = "Conditional Package Readiness - Ready for Interim Assessment (IATT / Review)"
        rec_badge = "warning"

    return {
        "readiness_score": None,
        "authorization_status": recommendation,
        "recommendation": recommendation,
        "rec_badge": rec_badge,
        "verified_atc_count": verified_atc_count,
        "atc_total": atc_total_scoped,
        "atc_records": atc_records,
        "cat_1_findings": cat_1_findings,
        "cat_2_findings": cat_2_findings,
        "cat_3_findings": cat_3_findings,
    }


# ------------------------------------------------------------------------------
# ATC (Authorization to Connect) connection controls
# ------------------------------------------------------------------------------
# Provenance: this list arrived with the original contribution as "the 14 ATC
# controls" with no cited DoD source; treat it as the engine's working review
# set, not as a quotation of a DoD issuance. The original used Rev 4 titles and
# included IA-2(3) and IA-2(4), which NIST SP 800-53 Rev 5 withdrew and
# incorporated into IA-2(1) and IA-2(2). Under Rev 5 the same requirements are
# these 12 controls. Titles match the Rev 5 catalog; the third field is the
# engine's review focus, not control text.
ATC_CONNECTION_CONTROLS: List[Tuple[str, str, str]] = [
    ("AC-17", "Remote Access",
     "Encrypted VPN/IAP tunnels, multi-factor authentication, and centralized access logging."),
    ("AC-17(2)", "Protection of Confidentiality and Integrity Using Encryption",
     "FIPS-validated cryptographic protection (TLS) for all remote sessions."),
    ("IA-2(1)", "Multi-factor Authentication to Privileged Accounts",
     "MFA for privileged accounts, network and local access (Rev 5 incorporated IA-2(3))."),
    ("IA-2(2)", "Multi-factor Authentication to Non-privileged Accounts",
     "MFA for non-privileged accounts, network and local access (Rev 5 incorporated IA-2(4))."),
    ("IA-5(1)", "Password-based Authentication",
     "Password composition, storage, and change rules where passwords are still used."),
    ("IR-8", "Incident Response Plan",
     "Documented and tested incident handling procedures with CISA/DoD reporting timelines."),
    ("IR-9", "Information Spillage Response",
     "Spillage containment, forensic isolation, and sanitization procedures."),
    ("RA-5", "Vulnerability Monitoring and Scanning",
     "Credentialed vulnerability scans and a defined remediation timeline."),
    ("SC-7", "Boundary Protection",
     "Default-deny perimeter firewall rules, VPC peering, and subnet isolation."),
    ("SC-8", "Transmission Confidentiality and Integrity",
     "TLS encryption on internal and external communication paths."),
    ("SC-28", "Protection of Information at Rest",
     "Cloud KMS CMEK encryption for persistent disks, buckets, and databases."),
    ("SI-2", "Flaw Remediation",
     "Patch management that remediates Critical/High vulnerabilities within a defined timeline."),
]


def _walk_artifact_files(ato_dir: str) -> List[str]:
    """Enumerates files under the artifact directory without following symlinks.

    ``glob.glob(..., recursive=True)`` descends into symlinked directories, so a
    symlink planted inside ``ato_artifacts/`` (for example ``docs -> ../../``) would
    pull blueprint-owned files into passes that rewrite artifacts in place. Walking
    with ``followlinks=False`` keeps enumeration inside the real artifact tree
    (CWE-59, CWE-22).

    Args:
        ato_dir: Root artifact directory to enumerate.

    Returns:
        Sorted list of regular-file paths beneath ``ato_dir``.
    """
    if not os.path.isdir(ato_dir):
        return []
    collected: List[str] = []
    for root, dirnames, filenames in os.walk(ato_dir, followlinks=False):
        # Drop symlinked subdirectories outright rather than merely not following
        # them, so they are never reported as containers of auditable artifacts.
        dirnames[:] = [d for d in dirnames if not os.path.islink(os.path.join(root, d))]
        for fname in filenames:
            fpath = os.path.join(root, fname)
            if os.path.isfile(fpath) and not os.path.islink(fpath):
                collected.append(fpath)
    return sorted(collected)


def hydrate_example_data_in_artifacts(
    md_and_yaml_files: List[str],
    allowed_boundary: Optional[str] = None,
) -> int:
    """Tags pending RMF placeholders with visible sample example badges.

    Replaces unconfigured variable placeholders with high-visibility callout badges
    in Markdown documents, and valid quoted YAML scalars in YAML matrices, to prevent
    premature submission of incomplete compliance artifacts without corrupting YAML syntax.

    Args:
        md_and_yaml_files: List of file paths to Markdown and YAML artifacts.
        allowed_boundary: Root directory every rewrite must stay inside. This pass
            rewrites files in place, so without a boundary a symlinked directory
            planted inside the artifact tree would redirect the write onto
            blueprint-owned source files (CWE-59).

    Returns:
        Total number of artifact files updated with tagged example badges.
    """
    total_tagged_items = 0
    sample_badge_style = "background-color: #fef08a; border: 2px dashed #ca8a04; color: #854d0e; font-weight: bold; padding: 2px 6px; border-radius: 4px;"

    for fpath in md_and_yaml_files:
        try:
            content = read_text_file(fpath, allowed_boundary=allowed_boundary)
        except OSError as err:
            logger.warning("Failed reading file %s for example hydration: %s", fpath, err)
            continue

        orig_content = content
        is_yaml = fpath.endswith((".yaml", ".yml"))

        content = TemplateEngine.hydrate_legacy_placeholders(
            content, is_yaml=is_yaml, badge_style=sample_badge_style
        )

        if content != orig_content:
            total_tagged_items += 1
            write_text_file(fpath, content, allowed_boundary=allowed_boundary)

    return total_tagged_items


def _inspect_zip_archive_safety(
    archive_path: str,
    max_archive_bytes: int = 50 * 1024 * 1024,
    max_entries: int = 5000,
    max_total_uncompressed: int = 200 * 1024 * 1024,
    max_entry_uncompressed: int = 100 * 1024 * 1024,
    max_compression_ratio: float = 100.0,
) -> None:
    """Validates an OpenXML (.docx/.xlsm/.xlsx) ZIP archive against zip-bomb limits."""
    if os.path.islink(archive_path):
        raise ValueError(f"Archive '{os.path.basename(archive_path)}' is a symbolic link")
    file_size = os.path.getsize(archive_path)
    if file_size > max_archive_bytes:
        raise ValueError(
            f"Archive '{os.path.basename(archive_path)}' size ({file_size} bytes) "
            f"exceeds safe size/compression limits ({max_archive_bytes} bytes)"
        )
    with zipfile.ZipFile(archive_path, "r") as zf:
        infos = zf.infolist()
        if len(infos) > max_entries:
            raise ValueError(
                f"Archive '{os.path.basename(archive_path)}' entry count ({len(infos)}) "
                f"exceeds safe size/compression limits ({max_entries} entries)"
            )
        total_uncompressed = 0
        for info in infos:
            total_uncompressed += info.file_size
            if info.file_size > max_entry_uncompressed:
                raise ValueError(
                    f"Archive entry '{info.filename}' uncompressed size ({info.file_size} bytes) "
                    f"exceeds safe size/compression limits ({max_entry_uncompressed} bytes)"
                )
            if (
                info.compress_size > 0
                and info.file_size > 1024 * 1024
                and (info.file_size / info.compress_size) > max_compression_ratio
            ):
                raise ValueError(
                    f"Archive entry '{info.filename}' compression ratio "
                    f"({info.file_size / info.compress_size:.1f}x) exceeds safe size/compression limits"
                )
        if total_uncompressed > max_total_uncompressed:
            raise ValueError(
                f"Archive '{os.path.basename(archive_path)}' total uncompressed size "
                f"({total_uncompressed} bytes) exceeds safe size/compression limits ({max_total_uncompressed} bytes)"
            )


def _count_populated_sheet_rows(ws: Any, start_row: int, key_cols: Tuple[int, ...]) -> int:
    """Counts worksheet rows >= start_row that have at least one non-empty primary-key cell."""
    max_r = int(getattr(ws, "max_row", 0) or 0)
    count = 0
    for r in range(start_row, max_r + 1):
        try:
            has_key = any(
                ws.cell(row=r, column=c).value is not None
                and str(ws.cell(row=r, column=c).value).strip() != ""
                for c in key_cols
            )
        except (AttributeError, TypeError, ValueError, IndexError):
            continue
        if has_key:
            count += 1
    return count


def _scan_sheet_for_unresolved_tokens(
    ws: Any,
    sheet_label: str,
    item_res: Dict[str, Any],
    max_rows: int = 2000,
    max_cols: Optional[int] = None,
    unresolved_out: Optional[List[Dict[str, Any]]] = None,
    placeholders_out: Optional[List[Dict[str, Any]]] = None,
    rel_file: Optional[str] = None,
) -> None:
    """Scans populated worksheet cells across all columns for unresolved ``{{ TOKEN }}`` macros and pending placeholders."""
    token_re = re.compile(r"\{\{\s*[A-Z0-9_]+\s*\}\}")
    row_limit = min(int(getattr(ws, "max_row", 0) or 0), max_rows)
    raw_col_max = int(getattr(ws, "max_column", 0) or 0)
    col_limit = min(raw_col_max, max_cols) if max_cols is not None else min(raw_col_max, 256)
    doc_label = rel_file or item_res.get("file", sheet_label)
    for r in range(1, row_limit + 1):
        for c in range(1, col_limit + 1):
            try:
                val = ws.cell(row=r, column=c).value
            except (AttributeError, TypeError, ValueError, IndexError) as err:
                item_res["issues"].append(
                    f"{sheet_label} Row {r} Col {c}: Failed reading cell during token scan ({err})."
                )
                item_res["status"] = "FAIL"
                continue
            if not isinstance(val, str) or not val:
                continue
            if "{{" in val:
                m = token_re.search(val)
                if m:
                    item_res["issues"].append(
                        f"{sheet_label} Row {r} Col {c}: Unresolved macro token '{m.group(0)}' in cell."
                    )
                    item_res["status"] = "FAIL"
                    if unresolved_out is not None:
                        unresolved_out.append({
                            "file": doc_label,
                            "line": r,
                            "token": m.group(0),
                            "context": val.strip()[:120],
                        })
            if placeholders_out is not None:
                for ph in _find_pending_placeholders_in_line(val):
                    placeholders_out.append({
                        "file": doc_label,
                        "line": r,
                        "variable": ph,
                        "context": val.strip()[:120],
                    })


def _audit_hwsw_workbook(
    wb: Any,
    rel_file: str,
    item_res: Dict[str, Any],
    unresolved_out: Optional[List[Dict[str, Any]]],
    placeholders_out: Optional[List[Dict[str, Any]]],
) -> None:
    """Audits Hardware_Software_Inventory.xlsm sheets, populated rows, and dropdown lists."""
    allowed_hw: Set[Any] = set()
    allowed_sw: Set[Any] = set()
    if "(U) Lists" in wb.sheetnames:
        lists_ws = wb["(U) Lists"]
        allowed_hw = {lists_ws.cell(row=r, column=1).value for r in range(2, lists_ws.max_row + 1) if lists_ws.cell(row=r, column=1).value is not None}
        allowed_sw = {lists_ws.cell(row=r, column=3).value for r in range(2, lists_ws.max_row + 1) if lists_ws.cell(row=r, column=3).value is not None}

    for sheet_name, count_key, key_cols, allowed_vals, tab_label in (
        ("Hardware", "Hardware Assets", (2, 3), allowed_hw, "Tab_HWType"),
        ("Software", "Software Assets", (2, 4), allowed_sw, "Tab_SWType"),
    ):
        if sheet_name not in wb.sheetnames:
            item_res["issues"].append(f"Missing required '{sheet_name}' sheet.")
            item_res["status"] = "FAIL"
            continue
        ws = wb[sheet_name]
        _scan_sheet_for_unresolved_tokens(ws, sheet_name, item_res, unresolved_out=unresolved_out, placeholders_out=placeholders_out, rel_file=rel_file)
        data_rows = _count_populated_sheet_rows(ws, 8, key_cols)
        item_res["rows_count"][count_key] = data_rows
        if data_rows == 0:
            item_res["issues"].append(f"{sheet_name} sheet contains no data rows (empty from row 8+).")
            if item_res["status"] != "FAIL":
                item_res["status"] = "WARN"
        elif allowed_vals:
            for r in range(8, ws.max_row + 1):
                if not ws.cell(row=r, column=1).value:
                    continue
                val_type = ws.cell(row=r, column=2).value
                type_label = "Component Type" if sheet_name == "Hardware" else "Software Type"
                if val_type and val_type not in allowed_vals:
                    item_res["issues"].append(f"{sheet_name} Row {r}: {type_label} '{val_type}' is not in {tab_label} dropdown.")
                    if item_res["status"] != "FAIL":
                        item_res["status"] = "WARN"


def _audit_poam_and_ppsm_workbook(
    wb: Any,
    rel_name: str,
    rel_file: str,
    item_res: Dict[str, Any],
    unresolved_out: Optional[List[Dict[str, Any]]],
    placeholders_out: Optional[List[Dict[str, Any]]],
) -> None:
    """Audits POA&M or PPSM workbook sheets, populated rows, and milestone dates."""
    if "Plan_of_Action_and_Milestones" in rel_name:
        if "POA&M" not in wb.sheetnames:
            item_res["issues"].append("Missing required 'POA&M' sheet.")
            item_res["status"] = "FAIL"
            return
        ws_poam = wb["POA&M"]
        _scan_sheet_for_unresolved_tokens(ws_poam, "POA&M", item_res, unresolved_out=unresolved_out, placeholders_out=placeholders_out, rel_file=rel_file)
        item_res["rows_count"]["POA&M Milestones"] = _count_populated_sheet_rows(ws_poam, 8, (1, 2))
        today_d = current_utc_datetime().replace(tzinfo=None).date()
        for r in range(8, ws_poam.max_row + 1):
            if not ws_poam.cell(row=r, column=1).value and not ws_poam.cell(row=r, column=2).value:
                continue
            p_status = str(ws_poam.cell(row=r, column=6).value or "").strip()
            sched_d = ws_poam.cell(row=r, column=7).value
            if p_status == "Ongoing" and sched_d:
                try:
                    sched_parsed = datetime.strptime(str(sched_d).strip()[:10], "%Y-%m-%d").date()
                except (ValueError, TypeError) as date_err:
                    item_res["issues"].append(
                        f"POA&M Row {r}: Scheduled Completion Date ({sched_d!r}) is not a valid ISO 8601 date ({date_err}); milestone currency is UNVERIFIED."
                    )
                    if item_res["status"] != "FAIL":
                        item_res["status"] = "WARN"
                else:
                    if sched_parsed <= today_d:
                        item_res["issues"].append(f"POA&M Row {r}: Ongoing finding has past/today Scheduled Completion Date ({sched_d}).")
                        if item_res["status"] != "FAIL":
                            item_res["status"] = "WARN"
    elif "PPSM" in rel_name:
        if "PPSM" not in wb.sheetnames:
            item_res["issues"].append("Missing required 'PPSM' sheet.")
            item_res["status"] = "FAIL"
            return
        ws_ppsm = wb["PPSM"]
        _scan_sheet_for_unresolved_tokens(ws_ppsm, "PPSM", item_res, unresolved_out=unresolved_out, placeholders_out=placeholders_out, rel_file=rel_file)
        item_res["rows_count"]["PPSM Boundary Records"] = _count_populated_sheet_rows(ws_ppsm, 9, (2, 3))


def _check_sctm_row_date(
    r: int,
    ctrl_id: Any,
    status_val: str,
    est_date: Any,
    today_d: Any,
    item_res: Dict[str, Any],
) -> Tuple[int, int]:
    """Returns (missing_est_date_inc, invalid_date_logic_inc) for one SCTM row."""
    if not est_date or not str(est_date).strip():
        return 1, 0
    try:
        d_parsed = datetime.strptime(str(est_date).strip(), "%m/%d/%Y").date()
    except (ValueError, TypeError) as date_err:
        item_res["issues"].append(
            f"SCTM Row {r} ({ctrl_id}): Estimated Completion Date ({est_date!r}) is not a valid "
            f"MM/DD/YYYY date ({date_err}); status/date consistency is UNVERIFIED."
        )
        if item_res["status"] != "FAIL":
            item_res["status"] = "WARN"
        return 0, 1

    if status_val in ("Implemented", "Inherited", "Hybrid", "Compensated") and d_parsed > today_d:
        item_res["issues"].append(f"SCTM Row {r} ({ctrl_id}): {status_val} control has future Estimated Completion Date ({est_date}).")
        if item_res["status"] != "FAIL":
            item_res["status"] = "WARN"
        return 0, 1
    if status_val == "Planned" and d_parsed <= today_d:
        item_res["issues"].append(f"SCTM Row {r} ({ctrl_id}): Planned control has past/today Estimated Completion Date ({est_date}).")
        if item_res["status"] != "FAIL":
            item_res["status"] = "WARN"
        return 0, 1
    return 0, 0


def _audit_sctm_workbook(
    wb: Any,
    rel_file: str,
    item_res: Dict[str, Any],
    unresolved_out: Optional[List[Dict[str, Any]]],
    placeholders_out: Optional[List[Dict[str, Any]]],
) -> None:
    """Audits SCTM_Burndown_Matrix.xlsm Template sheet invariants and date logic."""
    if "Template" not in wb.sheetnames:
        item_res["issues"].append("Missing required 'Template' sheet.")
        item_res["status"] = "FAIL"
        return
    ws_sctm = wb["Template"]
    _scan_sheet_for_unresolved_tokens(ws_sctm, "SCTM", item_res, unresolved_out=unresolved_out, placeholders_out=placeholders_out, rel_file=rel_file)
    active_controls = 0
    missing_est_date = 0
    missing_slcm_comments = 0
    missing_resp_entities = 0
    invalid_date_logic = 0
    consecutive_blanks = 0
    today_d = current_utc_datetime().replace(tzinfo=None).date()

    for r in range(7, ws_sctm.max_row + 1):
        ctrl_id = ws_sctm.cell(row=r, column=1).value
        if not ctrl_id or not str(ctrl_id).strip():
            consecutive_blanks += 1
            if consecutive_blanks > 10:
                break
            continue
        consecutive_blanks = 0
        active_controls += 1
        status_val = str(ws_sctm.cell(row=r, column=5).value or "").strip()
        if status_val == "Not Applicable":
            na_just = ws_sctm.cell(row=r, column=9).value
            if not na_just or not str(na_just).strip():
                item_res["issues"].append(f"SCTM Row {r} ({ctrl_id}): N/A status requires N/A Justification (Col I).")
                if item_res["status"] != "FAIL":
                    item_res["status"] = "WARN"
            est_date_na = ws_sctm.cell(row=r, column=10).value
            if est_date_na and str(est_date_na).strip():
                item_res["issues"].append(f"SCTM Row {r} ({ctrl_id}): N/A status requires blank Estimated Completion Date (Col J).")
                if item_res["status"] != "FAIL":
                    item_res["status"] = "WARN"
            continue

        m_inc, inv_inc = _check_sctm_row_date(r, ctrl_id, status_val, ws_sctm.cell(row=r, column=10).value, today_d, item_res)
        missing_est_date += m_inc
        invalid_date_logic += inv_inc
        if not ws_sctm.cell(row=r, column=19).value or not str(ws_sctm.cell(row=r, column=19).value).strip():
            missing_slcm_comments += 1
        if not ws_sctm.cell(row=r, column=12).value or not str(ws_sctm.cell(row=r, column=12).value).strip():
            missing_resp_entities += 1

    item_res["rows_count"]["Controls Monitored"] = active_controls
    item_res["rows_count"]["Estimated Dates Filled"] = active_controls - missing_est_date
    item_res["rows_count"]["SLCM Comments Filled"] = active_controls - missing_slcm_comments
    for count_val, msg in (
        (missing_est_date, f"SCTM contains {missing_est_date} controls missing Estimated Completion Date."),
        (missing_slcm_comments, f"SCTM contains {missing_slcm_comments} controls missing SLCM Comments."),
        (missing_resp_entities, f"SCTM contains {missing_resp_entities} controls missing Responsible Entities."),
        (invalid_date_logic, f"SCTM contains {invalid_date_logic} controls with invalid date logic (future for implemented or past for planned)."),
    ):
        if count_val > 0:
            item_res["issues"].append(msg)
            if item_res["status"] != "FAIL":
                item_res["status"] = "WARN"


def audit_excel_workbooks(
    ato_dir: str,
    unresolved_out: Optional[List[Dict[str, Any]]] = None,
    placeholders_out: Optional[List[Dict[str, Any]]] = None,
) -> List[Dict[str, Any]]:
    """Audits generated macro-enabled Excel workbooks for data integrity."""
    if not OPENPYXL_AVAILABLE or openpyxl is None:
        logger.warning("openpyxl is not installed; skipping Excel workbook audit")
        return [{
            "file": "Excel_Workbooks",
            "path": ato_dir,
            "status": "SKIP",
            "sheets": [],
            "rows_count": {},
            "issues": ["openpyxl is not installed; install openpyxl via 'pip install openpyxl' to audit Excel workbooks."],
        }]
    all_walked = _walk_artifact_files(ato_dir) if os.path.exists(ato_dir) else []
    xl_files = [f for f in all_walked if f.endswith((".xlsm", ".xlsx"))]
    excel_audit_results: List[Dict[str, Any]] = []

    for xl_path in xl_files:
        if "Google Services FedRamp Package" in xl_path:
            continue
        rel_name = os.path.basename(xl_path)
        rel_file = os.path.relpath(xl_path, ato_dir)
        item_res: Dict[str, Any] = {"file": rel_name, "path": xl_path, "status": "PASS", "sheets": [], "rows_count": {}, "issues": []}
        wb = None
        try:
            _inspect_zip_archive_safety(xl_path)
            wb = openpyxl.load_workbook(xl_path, data_only=True, keep_vba=True)
            item_res["sheets"] = wb.sheetnames
            if "Hardware_Software_Inventory" in rel_name:
                _audit_hwsw_workbook(wb, rel_file, item_res, unresolved_out, placeholders_out)
            elif "Plan_of_Action_and_Milestones" in rel_name or "PPSM" in rel_name:
                _audit_poam_and_ppsm_workbook(wb, rel_name, rel_file, item_res, unresolved_out, placeholders_out)
            elif "SCTM" in rel_name:
                _audit_sctm_workbook(wb, rel_file, item_res, unresolved_out, placeholders_out)
        except (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile) as e:
            item_res["status"] = "UNVERIFIED"
            item_res["issues"].append(f"Failed to load workbook with openpyxl, content UNVERIFIED: {str(e)}")
        finally:
            if wb is not None:
                try:
                    wb.close()
                except (OSError, ValueError, AttributeError) as err:
                    logger.debug("Failed closing workbook %s: %s", xl_path, err)
        excel_audit_results.append(item_res)
    return excel_audit_results


def _verify_docx_hyperlinks_and_tokens(
    docx_zip: zipfile.ZipFile,
    namelist: List[str],
    root: Any,
    paragraphs: List[Any],
    d_path: str,
    rel_name: str,
    item_res: Dict[str, Any],
    unresolved_out: Optional[List[Dict[str, Any]]],
    placeholders_out: Optional[List[Dict[str, Any]]],
) -> None:
    """Scans DOCX paragraphs for tokens/placeholders and verifies hyperlink relationship targets."""
    if os.path.basename(d_path) != "Path_to_Authorization.docx" and (unresolved_out is not None or placeholders_out is not None):
        docx_token_re = re.compile(r"\{\{\s*[A-Z0-9_]+\s*\}\}")
        for p_idx, p_elem in enumerate(paragraphs, 1):
            p_text = "".join(p_elem.itertext()).strip()
            if not p_text:
                continue
            if unresolved_out is not None and "{{" in p_text:
                for m in docx_token_re.finditer(p_text):
                    unresolved_out.append({"file": rel_name, "line": p_idx, "token": m.group(0), "context": p_text[:120]})
                    item_res["status"] = "FAIL"
            if placeholders_out is not None:
                for ph in _find_pending_placeholders_in_line(p_text):
                    placeholders_out.append({"file": rel_name, "line": p_idx, "variable": ph, "context": p_text[:120]})

    hyperlinks = root.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink")
    item_res["hyperlinks_count"] = len(hyperlinks)
    if "word/_rels/document.xml.rels" in namelist:
        rels_info = docx_zip.getinfo("word/_rels/document.xml.rels")
        if rels_info.file_size > 100 * 1024 * 1024 or (rels_info.compress_size > 0 and rels_info.file_size / rels_info.compress_size > 100):
            raise ValueError("word/_rels/document.xml.rels exceeds safe size/compression limits")
        rels_root = ET.fromstring(docx_zip.read("word/_rels/document.xml.rels"))
        rels_map = {
            r.attrib["Id"]: {"target": r.attrib.get("Target") or "", "type": r.attrib.get("Type") or ""}
            for r in rels_root.findall(".//{http://schemas.openxmlformats.org/package/2006/relationships}Relationship")
            if r.attrib.get("Id")
        }
        item_res["relationships_count"] = len(rels_map)
        rel_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        broken_links = [
            hl.attrib.get(rel_attr) or "missing_id"
            for hl in hyperlinks
            if not hl.attrib.get(rel_attr) or not rels_map.get(hl.attrib[rel_attr], {}).get("target")
        ]
        if broken_links:
            item_res["status"] = "FAIL"
            item_res["issues"].append(
                f"Hyperlink relationship integrity error: {len(broken_links)} hyperlink(s) lack valid target relationships ({broken_links[:5]})."
            )
    elif hyperlinks:
        item_res["status"] = "FAIL"
        item_res["issues"].append("Missing 'word/_rels/document.xml.rels' with active hyperlinks present.")


def audit_docx_policies(
    ato_dir: str,
    unresolved_out: Optional[List[Dict[str, Any]]] = None,
    placeholders_out: Optional[List[Dict[str, Any]]] = None,
) -> List[Dict[str, Any]]:
    """Audits generated Word (.docx) policy documents for OpenXML validity."""
    docx_audit_results: List[Dict[str, Any]] = []
    docx_files = [f for f in _walk_artifact_files(ato_dir) if f.endswith(".docx")] if os.path.exists(ato_dir) else []

    for d_path in sorted(docx_files):
        rel_name = os.path.relpath(d_path, ato_dir)
        item_res: Dict[str, Any] = {
            "file": rel_name, "status": "PASS", "has_cover": False, "has_headings": False,
            "hyperlinks_count": 0, "relationships_count": 0, "issues": [],
        }
        try:
            _inspect_zip_archive_safety(d_path)
            with zipfile.ZipFile(d_path, "r") as docx_zip:
                namelist = docx_zip.namelist()
                if "word/document.xml" not in namelist:
                    item_res["status"] = "FAIL"
                    item_res["issues"].append("Missing 'word/document.xml' in OpenXML archive.")
                    docx_audit_results.append(item_res)
                    continue
                doc_info = docx_zip.getinfo("word/document.xml")
                if doc_info.file_size > 100 * 1024 * 1024 or (doc_info.compress_size > 0 and doc_info.file_size / doc_info.compress_size > 100):
                    raise ValueError("word/document.xml exceeds safe size/compression limits")
                root = ET.fromstring(docx_zip.read("word/document.xml"))
                tables = root.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl")
                paragraphs = root.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")
                item_res["has_cover"] = len(tables) >= 1
                item_res["has_headings"] = len(paragraphs) >= 5
                _verify_docx_hyperlinks_and_tokens(docx_zip, namelist, root, paragraphs, d_path, rel_name, item_res, unresolved_out, placeholders_out)
        except (zipfile.BadZipFile, OSError, ET.ParseError, ValueError, KeyError) as e:
            item_res["status"] = "UNVERIFIED"
            item_res["issues"].append(f"Failed to parse OpenXML structure, content UNVERIFIED: {str(e)}")
        docx_audit_results.append(item_res)
    return docx_audit_results


def audit_oscal_packages(ato_dir: str) -> List[Dict[str, Any]]:
    """Audits generated NIST OSCAL deliverables for schema integrity and control coverage.

    Verifies system-security-plan and component-definition JSON deliverables exist,
    contain recognized NIST OSCAL metadata (e.g. 1.2.3 or 1.1.0), and define implemented
    requirements mapping to system components.

    Args:
        ato_dir: Absolute path to the ato_artifacts directory.

    Returns:
        A list of audit result dictionaries detailing OSCAL deliverable status.
    """
    oscal_audit_results: List[Dict[str, Any]] = []
    oscal_dir = os.path.join(ato_dir, "OSCAL_SSP")
    if not os.path.exists(oscal_dir):
        return oscal_audit_results

    for filename in ["system_security_plan.oscal.json", "component_definition.oscal.json"]:
        filepath = os.path.join(oscal_dir, filename)
        if not os.path.exists(filepath):
            continue
        rel_name = os.path.relpath(filepath, ato_dir)
        item_res: Dict[str, Any] = {
            "file": rel_name,
            "status": "PASS",
            "controls_count": 0,
            "components_count": 0,
            "oscal_version": None,
            "issues": [],
        }
        try:
            data = read_json_file(filepath, allowed_boundary=ato_dir)
            if "system-security-plan" in data:
                ssp = data["system-security-plan"]
                ver = ssp.get("metadata", {}).get("oscal-version")
                item_res["oscal_version"] = ver
                if ver not in SUPPORTED_OSCAL_VERSIONS:
                    item_res["status"] = "WARN"
                    item_res["issues"].append(
                        f"Unrecognized OSCAL version '{ver}'. Expected one of: {', '.join(SUPPORTED_OSCAL_VERSIONS[-4:])}"
                    )
                comps = ssp.get("system-implementation", {}).get("components", [])
                item_res["components_count"] = len(comps)
                reqs = ssp.get("control-implementation", {}).get("implemented-requirements", [])
                item_res["controls_count"] = len(reqs)
                if len(reqs) < 5:
                    item_res["status"] = "WARN"
                    item_res["issues"].append(f"Low control count ({len(reqs)}) in OSCAL SSP.")
            elif "component-definition" in data:
                compdef = data["component-definition"]
                ver = compdef.get("metadata", {}).get("oscal-version")
                item_res["oscal_version"] = ver
                if ver not in SUPPORTED_OSCAL_VERSIONS:
                    item_res["status"] = "WARN"
                    item_res["issues"].append(
                        f"Unrecognized OSCAL version '{ver}'. Expected one of: {', '.join(SUPPORTED_OSCAL_VERSIONS[-4:])}"
                    )
                comps = compdef.get("components", [])
                item_res["components_count"] = len(comps)
            else:
                item_res["status"] = "FAIL"
                item_res["issues"].append("Missing valid root OSCAL element.")
        except (OSError, ValueError, TypeError) as e:
            item_res["status"] = "UNVERIFIED"
            item_res["issues"].append(f"Failed parsing OSCAL JSON, content UNVERIFIED: {e}")
        oscal_audit_results.append(item_res)

    return oscal_audit_results


def audit_yaml_syntax_integrity(ato_dir: str) -> List[Dict[str, Any]]:
    """Performs deep YAML syntax integrity verification across all YAML deliverables.

    Reads each .yaml and .yml deliverable under ato_artifacts using read_yaml_file()
    to ensure valid syntax, proper indentation, and readable structured data.

    Args:
        ato_dir: Absolute path to the ato_artifacts directory.

    Returns:
        List of audit dictionaries for each audited YAML deliverable.
    """
    yaml_results: List[Dict[str, Any]] = []
    all_files = _walk_artifact_files(ato_dir)
    yaml_files = sorted([f for f in all_files if f.endswith((".yaml", ".yml")) and os.path.isfile(f)])

    for yf in yaml_files:
        rel = os.path.relpath(yf, ato_dir)
        item: Dict[str, Any] = {
            "file": rel,
            "abs_path": yf,
            "status": "PASS",
            "error": None,
        }
        try:
            parsed = read_yaml_file(yf, allowed_boundary=ato_dir)
            if parsed is None:
                item["status"] = "PASS"
                item["note"] = "Empty or null YAML document"
            elif not isinstance(parsed, (dict, list)):
                item["status"] = "WARN"
                item["note"] = f"Top-level YAML is scalar ({type(parsed).__name__})"
        except (OSError, ValueError, TypeError) as e:
            item["status"] = "FAIL"
            item["error"] = str(e)
            logger.error("YAML syntax error in deliverable '%s': %s", rel, e)
        yaml_results.append(item)

    return yaml_results


def _backup_artifact_package(ato_dir: str) -> str:
    """Snapshots an existing ``ato_artifacts/`` tree before ``--fix`` overwrites it."""
    if os.path.islink(ato_dir):
        raise PermissionError(f"Refusing to back up symlinked artifact directory: '{ato_dir}'")
    backup_dir = f"{ato_dir.rstrip(os.sep)}_backup_{current_utc_datetime().replace(tzinfo=None).strftime('%Y%m%dT%H%M%S_%f')}"
    if os.path.lexists(backup_dir):
        raise FileExistsError(f"Backup destination already exists: '{backup_dir}'")

    def _ignore_symlinks(src: str, names: List[str]) -> List[str]:
        skipped: List[str] = []
        for name in names:
            candidate = os.path.join(src, name)
            if os.path.islink(candidate):
                logger.warning("Skipping symbolic link during artifact backup: '%s'", candidate)
                skipped.append(name)
        return skipped

    try:
        shutil.copytree(
            ato_dir,
            backup_dir,
            symlinks=True,
            ignore=_ignore_symlinks,
            ignore_dangling_symlinks=True,
        )
        logger.warning(
            "--fix regenerates every deliverable and will overwrite manual edits. "
            "Existing package backed up to '%s'.",
            backup_dir,
        )
    except (OSError, shutil.Error) as err:
        logger.error(
            "Could not back up the existing package to '%s': %s. Aborting the "
            "drift repair rather than overwriting artifacts with no recovery path.",
            backup_dir,
            err,
        )
        raise
    return backup_dir


def _load_pta_template_sections(replacements: Mapping[str, str]) -> Dict[str, str]:
    """Loads top-level ``## `` sections from ``Path_to_Authorization_Template.md``.

    Args:
        replacements: Mapping of placeholder keys (without braces) to values.

    Returns:
        Dictionary mapping section headings (without leading ``## ``) to their
        hydrated Markdown section text (including the ``## <Heading>`` line).
    """
    template_path = os.path.join(TEMPLATES_DIR, "pta", "Path_to_Authorization_Template.md")
    raw = read_text_file(template_path)
    pta_regime = classify_regime(replacements.get("IMPACT_LEVEL") or "DoD IL5")
    raw = evaluate_template_conditionals(raw, build_regime_template_flags(pta_regime, {}))
    for key, val in replacements.items():
        val_str = str(val)
        raw = re.sub(r"\{\{\s*" + re.escape(key) + r"\s*\}\}", lambda _m: val_str, raw)
    cleaned_lines = [line for line in raw.splitlines() if line.strip() != "---"]
    cleaned = "\n".join(cleaned_lines)
    sections: Dict[str, str] = {}
    parts = re.split(r"(?m)^##\s+", cleaned)
    for part in parts[1:]:
        nl = part.find("\n")
        if nl == -1:
            heading = part.strip()
            body = ""
        else:
            heading = part[:nl].strip()
            body = part[nl + 1 :].strip()
        sections[heading] = f"## {heading}\n\n{body}\n"
    return sections


def _strip_pta_summary_echo_sections(text: str) -> str:
    """Removes validator-generated placeholder and token echo tables from Path_to_Authorization.md."""
    out_lines: List[str] = []
    skipping = False
    for line in text.splitlines():
        if line.startswith("## Critical Unresolved Syntax Tokens") or line.startswith(
            "## Pending Institutional Configuration Variables"
        ):
            skipping = True
            continue
        if skipping and line.startswith("## "):
            skipping = False
        if not skipping:
            out_lines.append(line)
    return "\n".join(out_lines)


def _read_oscal_parity_info(ato_dir: str) -> Tuple[str, List[Dict[str, Any]]]:
    """Reads system-name and implemented-requirements from OSCAL SSP JSON if present."""
    oscal_path = os.path.join(ato_dir, "OSCAL_SSP", "system_security_plan.oscal.json")
    if not os.path.exists(oscal_path):
        return "", []
    try:
        o_data = read_json_file(oscal_path, allowed_boundary=ato_dir)
    except (OSError, ValueError, TypeError):
        return "", []
    ssp_obj = (o_data.get("system-security-plan") or {}) if isinstance(o_data, dict) else {}
    o_name = str((ssp_obj.get("system-characteristics") or {}).get("system-name") or "").strip()
    if o_name.startswith("[CONFIG_REQUIRED") or o_name == "Cloud Foundation Platform":
        o_name = ""
    reqs = (ssp_obj.get("control-implementation") or {}).get("implemented-requirements") or []
    return o_name, reqs


def _collect_parity_names_and_open_controls(
    ato_dir: str,
    inv_name: str,
) -> Tuple[Dict[str, str], Set[str], Dict[str, Any], List[Dict[str, Any]]]:
    """Collects system names, open POA&M base controls, SCTM YAML data, and OSCAL requirements across deliverables."""
    system_names_found: Dict[str, str] = {}
    if inv_name and not inv_name.startswith("[CONFIG_REQUIRED"):
        system_names_found["system_inventory.json"] = inv_name

    open_poam_controls: Set[str] = set()
    yaml_specs = (
        ("POAM/Plan_of_Action_and_Milestones.yaml", "POAM.yaml"),
        ("SCTM/SCTM_Burndown_Matrix.yaml", "SCTM.yaml"),
        ("HW_SW_Inventory/Hardware_Software_Inventory.yaml", "HWSW.yaml"),
        ("PPSM/PPSM_Ports_Protocols_Services.yaml", "PPSM.yaml"),
    )
    sctm_yaml_data: Dict[str, Any] = {}
    for rel_p, label in yaml_specs:
        full_p = os.path.join(ato_dir, rel_p)
        if not os.path.exists(full_p):
            continue
        try:
            y_data = read_yaml_file(full_p, allowed_boundary=ato_dir)
        except (OSError, ValueError, TypeError):
            continue
        if not isinstance(y_data, dict):
            continue
        meta_name = str((y_data.get("system_metadata") or {}).get("system_name") or "").strip()
        if meta_name and not meta_name.startswith("[CONFIG_REQUIRED"):
            system_names_found[label] = meta_name
        if label == "POAM.yaml":
            for it in y_data.get("poam_items") or y_data.get("items") or []:
                if isinstance(it, dict) and str(it.get("status") or "Ongoing").strip().lower() in ("ongoing", "open", "delayed"):
                    raw_c = it.get("control_identifier") or it.get("control") or it.get("aps") or ""
                    open_poam_controls.update(base_control_id(cid) for cid in extract_normalized_control_ids(raw_c))
        elif label == "SCTM.yaml":
            sctm_yaml_data = y_data

    o_name, oscal_reqs = _read_oscal_parity_info(ato_dir)
    if o_name:
        system_names_found["OSCAL_SSP"] = o_name
    return system_names_found, open_poam_controls, sctm_yaml_data, oscal_reqs


def _find_sctm_xlsm_poam_conflicts(ato_dir: str, open_poam_controls: Set[str]) -> List[str]:
    """Finds controls marked Implemented in SCTM XLSM despite open POA&M items."""
    sctm_xl = os.path.join(ato_dir, "SCTM", "SCTM_Burndown_Matrix.xlsm")
    if not (OPENPYXL_AVAILABLE and openpyxl is not None and os.path.exists(sctm_xl)):
        return []
    conflicts: List[str] = []
    wb = None
    try:
        wb = openpyxl.load_workbook(sctm_xl, data_only=True, keep_vba=True)
        if "Template" in wb.sheetnames:
            ws = wb["Template"]
            for r in range(7, int(ws.max_row or 0) + 1):
                cid = ws.cell(row=r, column=1).value
                if cid and base_control_id(cid) in open_poam_controls:
                    if str(ws.cell(row=r, column=5).value or "").strip().lower() == "implemented":
                        conflicts.append(
                            f"SCTM XLSM Row {r} marks control '{base_control_id(cid)}' as Implemented despite an open POA&M item."
                        )
    except (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile):
        pass
    finally:
        if wb is not None:
            try:
                wb.close()
            except (OSError, ValueError, AttributeError):
                pass
    return conflicts


def _find_open_poam_status_conflicts(
    ato_dir: str,
    open_poam_controls: Set[str],
    sctm_yaml_data: Dict[str, Any],
    oscal_reqs: List[Dict[str, Any]],
) -> List[str]:
    """Finds controls marked Implemented in OSCAL, SCTM YAML, or SCTM XLSM despite open POA&M items."""
    if not open_poam_controls:
        return []
    conflicts: List[str] = []
    for req in oscal_reqs:
        if not isinstance(req, dict):
            continue
        bc = base_control_id(req.get("control-id") or "")
        if bc in open_poam_controls:
            for comp in req.get("by-components") or []:
                if str((comp.get("implementation-status") or {}).get("state") or "").strip().lower() == "implemented":
                    conflicts.append(f"OSCAL SSP marks control '{bc}' as implemented despite an open POA&M item.")
    for fam in sctm_yaml_data.get("control_families") or []:
        if not isinstance(fam, dict):
            continue
        for ctrl in fam.get("controls") or []:
            if isinstance(ctrl, dict):
                bc = base_control_id(ctrl.get("id") or "")
                if bc in open_poam_controls and str(ctrl.get("status") or "").strip().lower() == "implemented":
                    conflicts.append(f"SCTM YAML marks control '{bc}' as Implemented despite an open POA&M item.")
    conflicts.extend(_find_sctm_xlsm_poam_conflicts(ato_dir, open_poam_controls))
    return conflicts


def audit_cross_artifact_parity(ato_dir: str, inventory: Dict[str, Any]) -> Dict[str, Any]:
    """Audits cross-artifact consistency for system identity, open-POA&M control status, and regime scope."""
    sys_info = (inventory.get("system_information") or {}) if isinstance(inventory, dict) else {}
    inv_name, _ = resolve_system_identity(sys_info)
    system_names_found, open_poam_controls, sctm_yaml_data, oscal_reqs = _collect_parity_names_and_open_controls(ato_dir, inv_name)
    status_conflicts = _find_open_poam_status_conflicts(ato_dir, open_poam_controls, sctm_yaml_data, oscal_reqs)

    scope_issues: List[str] = []
    regime = classify_regime(sys_info)
    expectations = REGIME_CONTROL_EXPECTATIONS.get(regime.value, {})
    ssp_md = os.path.join(ato_dir, "SSP", "SSP_System_Security_Plan.md")
    if expectations and os.path.exists(ssp_md):
        ssp_txt = _read_optional_lowercase(ssp_md, lower=False)
        ssp_ctrl_cnt = len(re.findall(r"^###\s+[A-Z]{2}-\d+", ssp_txt, re.MULTILINE))
        min_ssp = int(expectations.get("min_ssp_controls", 0))
        if 0 < ssp_ctrl_cnt < min_ssp:
            scope_issues.append(f"SSP control count ({ssp_ctrl_cnt}) is below {regime.value} expectation ({min_ssp}).")

    system_name_consistent = len(set(system_names_found.values())) <= 1
    conflicts: List[str] = []
    if not system_name_consistent:
        conflicts.append(f"System name mismatch across artifacts: {system_names_found}")
    conflicts.extend(status_conflicts)
    return {
        "system_name_consistent": system_name_consistent,
        "system_names_found": system_names_found,
        "status_conflicts": status_conflicts,
        "scope_issues": scope_issues,
        "conflicts": conflicts,
    }


def _load_validator_inventory(inventory_path: str, target_dir: str) -> Tuple[Dict[str, Any], bool, str]:
    """Loads and validates system_inventory.json, returning (inventory, is_corrupt, error_msg)."""
    if not os.path.exists(inventory_path):
        return {}, False, ""
    try:
        loaded_inv = read_json_file(inventory_path, allowed_boundary=target_dir)
    except (OSError, ValueError, TypeError) as err:
        logger.error("Failed to read system inventory '%s': %s", inventory_path, err)
        return {}, True, str(err)
    if not isinstance(loaded_inv, dict):
        msg = f"Expected JSON object in '{inventory_path}', got {type(loaded_inv).__name__}"
        logger.error("%s", msg)
        return {}, True, msg
    try:
        validate_system_inventory_schema(loaded_inv, source_path=inventory_path)
    except (ValueError, TypeError) as err:
        logger.warning("System inventory validation warning: %s", err)
    return loaded_inv, False, ""


def _run_drift_preflight_repair(
    target_dir: str,
    ato_dir: str,
    inventory: Dict[str, Any],
    policy_format: Optional[str],
    data_format: Optional[str],
    oscal_format: Optional[str],
    oscal_version: Optional[str],
    run_scanners: bool,
) -> bool:
    """Backs up ato_artifacts and runs generate_compliance_artifacts.py; returns True if failed."""
    generator_script = os.path.join(SKILL_BASE, "scripts", "generate_compliance_artifacts.py")
    if os.path.isdir(ato_dir):
        _backup_artifact_package(ato_dir)
    logger.info("Pre-flight sync: Synchronizing generated compliance package with live codebase inventory before audit...")
    export_prefs = inventory.get("export_preferences") or inventory.get("compliance_config", {})
    p_fmt = policy_format or export_prefs.get("policy_formats")
    d_fmt = data_format or export_prefs.get("structured_data_formats")
    o_fmt = oscal_format or export_prefs.get("oscal_formats") or export_prefs.get("oscal_format")
    o_ver = oscal_version or export_prefs.get("oscal_version")

    gen_cmd = (
        [sys.executable, generator_script]
        if os.path.isfile(generator_script)
        else [sys.executable, "-m", "compliance_engine.generate_compliance_artifacts"]
    )
    if p_fmt:
        gen_cmd.append(f"--policy-format={p_fmt}")
    if d_fmt:
        gen_cmd.append(f"--data-format={d_fmt}")
    if o_fmt:
        gen_cmd.append(f"--oscal-format={o_fmt}")
    if o_ver:
        gen_cmd.append(f"--oscal-version={o_ver}")
    if run_scanners:
        gen_cmd.append("--run-scanners")
    gen_cmd.extend(["--", str(resolve_path(target_dir))])

    try:
        with audit_operation(AuditEvent.EXTERNAL_COMMAND, obj=target_dir, detail={"command": str(gen_cmd)}):
            res = run_bounded_command(gen_cmd, timeout=600, max_output_bytes=8 * 1024 * 1024)
            if res.stdout.strip():
                logger.info("%s", res.stdout.rstrip())
            if res.stderr.strip():
                logger.info("%s", res.stderr.rstrip())
            if res.returncode != 0:
                logger.warning("Artifact generator exited with non-zero code %d during drift repair.", res.returncode)
                raise subprocess.SubprocessError(f"Generator failed with code {res.returncode}")
    except subprocess.TimeoutExpired:
        logger.error("Artifact generator timed out after 600 seconds during drift repair.")
        return True
    except (subprocess.SubprocessError, OSError) as err:
        logger.error("Failed executing artifact generator during drift repair: %s", err)
        return True
    return False


def _scan_markdown_and_yaml_tokens(
    audited_text_files: List[str],
    ato_dir: str,
    own_report: str,
    fill_examples: bool,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Scans Markdown and YAML deliverables for unresolved {{ ... }} macros, pending placeholders, and action items."""
    if fill_examples:
        logger.info("Relabeling [CONFIG_REQUIRED: ...] placeholders as [EXAMPLE_DATA_REQUIRED: ...]...")
        cnt = hydrate_example_data_in_artifacts(audited_text_files, allowed_boundary=ato_dir)
        logger.info("Relabeled placeholders in %d artifact files.", cnt)

    unresolved_tokens: List[Dict[str, Any]] = []
    config_required_vars: List[Dict[str, Any]] = []
    rmf_action_items: List[Dict[str, Any]] = []
    token_regex = re.compile(r"\{\{\s*[A-Z0-9_]+\s*\}\}")
    action_item_regex = re.compile(
        r"\[(?:WARNING|INFORMATIONAL):\s*([^\]\[]+)\]|<mark\b[^>]*>([^<]{0,8192})</mark>",
        re.DOTALL,
    )

    for fpath in audited_text_files:
        rel_path = os.path.relpath(fpath, ato_dir)
        try:
            content = read_text_file(fpath, allowed_boundary=ato_dir)
        except OSError as err:
            logger.warning("Failed reading file %s during token audit: %s", fpath, err)
            continue
        if os.path.normcase(fpath) == own_report:
            content = _strip_pta_summary_echo_sections(content)
        for idx, line in enumerate(content.splitlines(), 1):
            for m in token_regex.finditer(line):
                unresolved_tokens.append({"file": rel_path, "line": idx, "token": m.group(0), "context": line.strip()})
            for ph in _find_pending_placeholders_in_line(line):
                config_required_vars.append({"file": rel_path, "line": idx, "variable": ph, "context": line.strip()})
            for m in action_item_regex.finditer(line):
                txt = (m.group(1) or m.group(2) or "").strip()
                if txt:
                    rmf_action_items.append({"file": rel_path, "line": idx, "action": txt})
    return unresolved_tokens, config_required_vars, rmf_action_items


_EXPECTED_RUNBOOKS: Tuple[str, ...] = (
    "IR_IAM_Compromised_Credentials_Runbook",
    "IR_Compute_Resource_Compromise_Runbook",
    "IR_KMS_CMEK_Compromise_Runbook",
    "IR_Network_Intrusion_Runbook",
    "IR_VPC_Service_Controls_Violation_Runbook",
)


def _audit_runbooks_presence(ato_dir: str) -> List[Dict[str, Any]]:
    """Checks presence of the expected incident response runbooks."""
    runbooks_dir = os.path.join(ato_dir, "Incident_Response_Runbooks")
    if not os.path.exists(runbooks_dir):
        return []
    results: List[Dict[str, Any]] = []
    for rb_name in _EXPECTED_RUNBOOKS:
        has_md = os.path.exists(os.path.join(runbooks_dir, f"{rb_name}.md"))
        has_docx = os.path.exists(os.path.join(runbooks_dir, f"{rb_name}.docx"))
        results.append({"runbook": rb_name, "has_md": has_md, "has_docx": has_docx, "status": "PASS" if has_md else "MISSING"})
    return results


_RECON_CATEGORIES_META: Tuple[Tuple[str, str], ...] = (
    ("compute", "Compute & Virtual Machines"),
    ("storage", "Cloud Storage Buckets"),
    ("databases", "Database Instances & Warehouses"),
    ("vpcs", "Virtual Private Clouds (VPCs)"),
    ("subnets", "Subnet CIDR Ranges"),
    ("firewall_rules", "Boundary Firewall Rules"),
    ("kms_keys", "KMS Cryptographic Keys"),
    ("service_accounts", "IAM Service Accounts"),
    ("cloud_run_services", "Cloud Run Services"),
    ("cloud_functions", "Cloud Functions"),
    ("artifact_registries", "Artifact Registries"),
    ("logging_sinks", "Cloud Logging Sinks"),
    ("pubsub_topics", "Pub/Sub Topics"),
    ("nat_gateways", "Cloud NAT Gateways"),
    ("service_perimeters", "VPC-SC Service Perimeters"),
    ("exposed_ports", "Boundary Ingress Ports (PPSM)"),
    ("applications", "Workload Applications"),
    ("container_images", "Container & Workload Images"),
)


def _escape_pta_table_cell(val: Any) -> str:
    """Normalizes CR/LF and escapes unescaped pipes for Markdown table cells."""
    text = re.sub(r"[\r\n]+", " ", str(val or "")).strip()
    return re.sub(r"(?<!\\)\|", r"\|", text)


def _escape_pta_code_cell(val: Any) -> str:
    """Escapes a string for placement inside a backticked Markdown table cell (`...`)."""
    return _escape_pta_table_cell(str(val or "").replace("`", "'"))


def _clean_pta_gov_field(val: Any) -> str:
    """Sanitizes a governance header field for safe table and line-start Markdown rendering."""
    return _sanitize_markdown_inline_value(str(val or ""), on_table_line=True, at_line_start=True)


def _extract_pta_governance_metadata(inventory: Dict[str, Any]) -> Dict[str, str]:
    """Resolves system, organization, FIPS-199, and personnel strings for PTA rendering."""
    sys_info = inventory.get("system_information", {}) if inventory else {}
    sys_name, sys_abbr = resolve_system_identity(sys_info)
    org_name = sys_info.get("organization") or "Department of Defense / Enterprise"
    impact_level = sys_info.get("impact_level") or "IL5 / FedRAMP High"
    fips_199 = (
        f"Confidentiality: {sys_info.get('confidentiality_impact') or 'High'} / "
        f"Integrity: {sys_info.get('integrity_impact') or 'High'} / "
        f"Availability: {sys_info.get('availability_impact') or 'High'}"
    )
    baseline = sys_info.get("compliance_baseline") or "NIST SP 800-53 Rev. 5 (FedRAMP High / DoD IL5)"
    roles_info = inventory.get("personnel_roles", {})
    raw_gov = {
        "sys_name": sys_name,
        "sys_abbr": sys_abbr,
        "org_name": org_name,
        "impact_level": str(impact_level),
        "fips_199": fips_199,
        "baseline": str(baseline),
        "issm_name": roles_info.get("issm", {}).get("name") or sys_info.get("issm_name") or "[CONFIG_REQUIRED: Lead ISSM Name]",
        "issm_title": roles_info.get("issm", {}).get("title") or sys_info.get("issm_title") or "Information System Security Manager (ISSM)",
        "so_name": roles_info.get("system_owner", {}).get("name") or sys_info.get("system_owner_name") or "[CONFIG_REQUIRED: System Owner Name]",
        "so_title": roles_info.get("system_owner", {}).get("title") or sys_info.get("system_owner_title") or "System Owner / Program Manager",
        "ao_name": (
            roles_info.get("authorizing_official", {}).get("name")
            or sys_info.get("authorizing_official_name")
            or "[CONFIG_REQUIRED: Authorizing Official Name]"
        ),
        "ao_title": roles_info.get("authorizing_official", {}).get("title") or sys_info.get("authorizing_official_title") or "Authorizing Official (AO)",
        "target_date": (
            sys_info.get("target_authorization_date")
            or sys_info.get("effective_date")
            or current_utc_datetime().replace(tzinfo=None).strftime("%B %d, %Y")
        ),
    }
    return {k: _clean_pta_gov_field(v) for k, v in raw_gov.items()}


def _render_pta_header_and_dashboard(
    gov: Dict[str, str],
    bundle: Dict[str, Any],
) -> List[str]:
    """Renders the PTA governance table, executive summary, and multi-level readiness dashboard."""
    yaml_results = bundle["yaml_results"]
    docx_results = bundle["docx_results"]
    excel_results = bundle["excel_results"]
    oscal_results = bundle["oscal_results"]
    runbook_results = bundle["runbook_results"]
    stigs_required = bundle["stigs_required"]
    senior_audit = bundle["senior_audit"]
    alignment_results = bundle["alignment_results"]
    unresolved_tokens = bundle["unresolved_tokens"]
    config_required_vars = bundle["config_required_vars"]
    rmf_action_items = bundle["rmf_action_items"]

    yaml_status_note = _status_note(yaml_results, "parsed as valid YAML")
    total_docx_hl = sum(r.get("hyperlinks_count", 0) for r in docx_results)
    docx_status_note = _status_note(docx_results, "OpenXML structure and hyperlink targets checked")
    atc_total = senior_audit.get("atc_total", len(ATC_CONNECTION_CONTROLS))
    recommendation = senior_audit["recommendation"]
    verified_atc = senior_audit["verified_atc_count"]
    reconciled_cnt = alignment_results["reconciled_assets_count"]
    total_disc_cnt = alignment_results["total_discovered_assets"]
    packaging_results = docx_results + excel_results + oscal_results + yaml_results
    packaging_failed = [r for r in packaging_results if r.get("status") in ("FAIL", "UNVERIFIED")]

    lines: List[str] = [
        "# Path to Authorization (PTA) Strategy, Master ATO Roadmap & Package Validation Audit\n",
        "## Document Governance & Accreditation Baseline\n",
        "| Governance Metric | Policy Standard & Specification |",
        "| :--- | :--- |",
        "| **Document Title** | Path to Authorization (PTA) Strategy & Master ATO Roadmap |",
        f"| **Target System Name** | {gov['sys_name']} ({gov['sys_abbr']}) |",
        f"| **Security Categorization** | {gov['fips_199']} ({gov['impact_level']}) |",
        f"| **Governing Entity** | {gov['org_name']} |",
        f"| **Compliance Baseline** | {gov['baseline']} |",
        f"| **Document Owner** | {gov['issm_name']} ({gov['issm_title']}) |",
        f"| **System Owner** | {gov['so_name']} ({gov['so_title']}) |",
        f"| **Approval Authority** | {gov['ao_name']} ({gov['ao_title']}) |",
        f"| **Target Authorization Date** | {gov['target_date']} |",
        f"| **Audit Timestamp** | {current_utc_datetime().replace(tzinfo=None).strftime('%B %d, %Y at %H:%M:%S')} |\n",
        "> [!IMPORTANT]",
        "> **HOW THIS CODEBASE ACCELERATES YOUR PATH TO ATO**:",
        (
            "> Executing this compliance automation provisions the **foundational technical deliverables (Phase 3)** and "
            "extracts live architecture facts from Terraform. However, achieving formal Authorization to Operate (ATO) "
            "requires completing the end-to-end operational, governance, and assessment itinerary below. Running code "
            "alone does not grant an ATO—it provides the accelerated baseline and technical evidence needed to achieve it.\n"
        ),
        "> [!NOTE]",
        "> **eMASS & Enterprise GRC Direct Authoring & Import Workflow**:",
        (
            "> In many Department of Defense (DoD) and Federal authorization environments, the System Security Plan (SSP) "
            "is authored and maintained directly within enterprise GRC platforms such as **eMASS**, **CSAM**, **Xacta**, "
            "or the **FedRAMP Repository**."
        ),
        (
            "> The SSP documentation, SCTM burndown matrices, and policy statements generated by this engine provide the "
            "authoritative technical implementation narratives, infrastructure parameters, and control allocations "
            "formatted for direct manual entry or bulk automated ingestion (via eMASS TR/TRX and NIST OSCAL XML/JSON "
            "exports) into the GRC system of record.\n"
        ),
        "## Executive Artifact Audit & Integrity Summary",
        f"- **Markdown & YAML Inventory Files Audited**: `{bundle['total_files_checked']}` files",
        f"- **Structured YAML Deliverables Audited**: `{len(yaml_results)}` files (`{yaml_status_note}`)",
        (
            f"- **Word Policy Manuals Audited (.docx)**: `{len(docx_results)}` documents "
            f"(`{docx_status_note}`; {total_docx_hl} hyperlinks with relationship targets)"
        ),
    ]
    if oscal_results:
        oscal_ver_str = oscal_results[0].get("oscal_version") or "[NOT DETERMINED FROM SOURCE]"
        oscal_status_note = _status_note(oscal_results, "root element, version, and control count checked")
        lines.append(
            f"- **NIST OSCAL Packages (v{oscal_ver_str})**: `{len(oscal_results)}` deliverables "
            f"(`{oscal_status_note}`; not validated against the OSCAL JSON schema)"
        )
    lines.extend([
        f"- **Unresolved Macro Templates (`{{ ... }}`)**: `{len(unresolved_tokens)}` syntax gaps",
        f"- **Pending Institutional Variables**: `{len(config_required_vars)}` placeholders",
        f"- **Designated RMF Team Human Action Highlights**: `{len(rmf_action_items)}` review items requiring manual input",
    ])
    if runbook_results:
        passed_runbooks = sum(1 for r in runbook_results if r["status"] == "PASS")
        lines.append(f"- **Incident Response Runbooks Audited**: `{passed_runbooks}/{len(_EXPECTED_RUNBOOKS)}` scenarios present")
    cat_summary = f"{len(senior_audit['cat_1_findings'])} CAT I / {len(senior_audit['cat_2_findings'])} CAT II"
    lines.extend([
        f"- **Applicable DISA STIG Checklists Identified**: `{len(stigs_required)}` STIG benchmarks",
        f"- **Authorization Readiness Determination**: {recommendation}",
        f"- **Dynamic Architecture Reconciliation Coverage**: {_format_coverage(alignment_results)}",
        f"- **ATC Connection Controls Verified**: `{verified_atc}/{atc_total}` verified\n",
        "## Senior Compliance Assessor Multi-Level Quality & Technical Evidence Readiness Dashboard\n",
        (
            "Automated checks run by `validate_compliance_artifacts.py` over the generated deliverables. "
            "They support, but do not replace, a control assessment under NIST SP 800-53A Rev. 5:\n"
        ),
        "| Assessor Evaluation Dimension | Evaluation / Determination | Target Baseline | Compliance Assessor Assessment Status |",
        "| :--- | :--- | :--- | :--- |",
        f"| **Authorization Readiness Gate** | **{recommendation}** | Ready for Independent Review | **{recommendation}** |",
        (
            f"| **Level 1: Multi-Format Packaging Integrity** | `{len(packaging_failed)} failed or unreadable of "
            f"{len(packaging_results)}` | 0 failed | OpenXML structure (.docx/.xlsm), YAML syntax, and OSCAL root "
            "element and version checked |"
        ),
        (
            f"| **Level 2: Dynamic Architecture Reconciliation** | {_format_coverage(alignment_results)} | "
            f"All discovered assets reconciled | Reconciled {reconciled_cnt} of {total_disc_cnt} discovered cloud "
            "infrastructure & application assets |"
        ),
        (
            f"| **Level 3: DISA STIG & SRG Baseline Alignment** | `{len(stigs_required)} Benchmarks` | "
            "Complete Technology Coverage | Mapped to STIG Viewer desktop workflow & DoD Cyber Exchange |"
        ),
        (
            f"| **Level 4: ATC Connection Controls** | `{verified_atc} / {atc_total} Verified` | "
            f"{atc_total}/{atc_total} Implemented | Verified implementation depth across SCTM & SSP |"
        ),
        (
            f"| **Level 5: Continuous Monitoring & POA&M** | `{cat_summary}` | 0 CAT I Findings | "
            "Assessed boundary exposure, CMEK rotation, and milestone currency |"
        ),
        "",
    ])
    return lines


def _format_finding_severity_display(raw_s: str) -> str:
    """Formats raw finding severity into eMASS CAT display strings."""
    if re.search(r"\bCAT\s*III\b", raw_s, re.I) or "LOW" in raw_s.upper() or "ADVISORY" in raw_s.upper():
        return "Advisory (eMASS CAT III)"
    if re.search(r"\bCAT\s*II\b", raw_s, re.I) or "MEDIUM" in raw_s.upper() or "MODERATE" in raw_s.upper():
        return "Medium (eMASS CAT II)"
    if re.search(r"\bCAT\s*I\b", raw_s, re.I) or "CRITICAL" in raw_s.upper() or "HIGH" in raw_s.upper():
        return "Critical (eMASS CAT I)"
    return raw_s


def _render_pta_reconciliation_and_findings(bundle: Dict[str, Any]) -> List[str]:
    """Renders Level 2 reconciliation breakdown, assessor findings, and semantic linter report."""
    alignment_results = bundle["alignment_results"]
    senior_audit = bundle["senior_audit"]
    semantic_linter_report = bundle["semantic_linter_report"]
    reconciled_cnt = alignment_results["reconciled_assets_count"]
    total_disc_cnt = alignment_results["total_discovered_assets"]
    bd = alignment_results.get("breakdown", {})

    lines: List[str] = [
        "### Dynamic System Architecture & Inventory Cross-Reconciliation",
        (
            "The validation engine checks whether each cloud asset discovered in `system_inventory.json` is named in the "
            "System Security Plan (SSP), Hardware/Software Inventory, PPSM, or FIPS matrix across the categories below:\n"
        ),
        "| Asset Category | Discovered in Inventory | Documented in ATO Package | Reconciliation Coverage | Status |",
        "| :--- | :--- | :--- | :--- | :--- |",
    ]
    for cat_key, cat_label in _RECON_CATEGORIES_META:
        c_disc = bd.get(cat_key, {}).get("discovered", 0)
        c_match = bd.get(cat_key, {}).get("matched", 0)
        c_pct = ("%.1f%%" % ((c_match / c_disc) * 100.0)) if c_disc > 0 else "N/A"
        c_status = "None discovered" if c_disc == 0 else ("Reconciled" if c_match >= c_disc else "Review Needed")
        lines.append(f"| **{cat_label}** | `{c_disc}` | `{c_match}` | {c_pct} | {c_status} |")

    total_status = "None discovered" if total_disc_cnt == 0 else ("Complete" if reconciled_cnt >= total_disc_cnt else "Review Needed")
    total_pct = "N/A" if total_disc_cnt == 0 else "%.1f%%" % alignment_results["coverage_score_percent"]
    lines.extend([
        f"| **Total Assets Reconciled** | **`{total_disc_cnt}`** | **`{reconciled_cnt}`** | **`{total_pct}`** | **`{total_status}`** |",
        "",
    ])

    discs = alignment_results.get("discrepancies", [])
    if discs:
        lines.extend(["> [!NOTE]", "> **Assessor Inventory Alignment Notes**:"])
        lines.extend(f"> - {disc}" for disc in discs[:5])
        if len(discs) > 5:
            lines.append(f"> - *...and {len(discs) - 5} additional alignment observations.*")
        lines.append("")

    all_findings = senior_audit.get("cat_1_findings", []) + senior_audit.get("cat_2_findings", []) + senior_audit.get("cat_3_findings", [])
    if all_findings:
        lines.extend([
            "### Senior Compliance Assessor Findings & Remediation Plan",
            (
                "The following security findings and documentation gaps were identified during multi-level assessment. "
                "Findings are categorized by **NIST SP 800-30 / FedRAMP Risk Level** (Critical, High, Medium, Low) and "
                "mapped to **DoD eMASS POA&M Severity Categories** (CAT I, CAT II, CAT III) for formal package intake. "
                "Technical DISA STIG benchmark findings are tracked separately via STIG Viewer checklists (.ckl):\n"
            ),
            "| Finding ID | Affected Component | Severity (NIST / eMASS) | Description | Assessor Remediation Guidance |",
            "| :--- | :--- | :--- | :--- | :--- |",
        ])
        for f in all_findings:
            sev_display = _format_finding_severity_display(f.get("severity_label") or f.get("severity", ""))
            lines.append(f"| `{f['id']}` | **{f['component']}** | `{sev_display}` | {f['description']} | {f['remediation']} |")
        lines.append("")

    if semantic_linter_report:
        lines.extend([semantic_linter_report.to_markdown(), ""])
    return lines


def _render_pta_atc_stigs_and_openxml(
    bundle: Dict[str, Any],
    pta_sections: Dict[str, str],
) -> List[str]:
    """Renders RMF crosswalk sections, ATC table, ACAS rules, DISA STIG table, and Excel/DOCX audit tables."""
    senior_audit = bundle["senior_audit"]
    stigs_required = bundle["stigs_required"]
    excel_results = bundle["excel_results"]
    docx_results = bundle["docx_results"]
    lines: List[str] = []

    for sec_title in (
        "NIST SP 800-37 Rev. 2 RMF 7-Step Crosswalk",
        "Master ATO Journey & Complete Accreditation Itinerary",
        "Strategic RMF Considerations & Authorizing Official (AO) Engagement",
        "Federal & DoD Privacy Compliance Requirements (PIA, PCIL, SORN)",
    ):
        if sec_title in pta_sections:
            lines.append(pta_sections[sec_title])

    lines.extend([
        "## ATC (Authorization to Connect) Connection Controls Verification",
        (
            "For systems connecting to DoD enterprise networks or requesting an Authorization to Connect (ATC), the "
            f"following {len(ATC_CONNECTION_CONTROLS)} controls must have complete implementation statements in the SCTM "
            "and zero unmitigated High/Very High residual risks. This is the engine's working review set (no DoD issuance "
            "is cited for it); IA-2(3) and IA-2(4) from older 14-control lists were withdrawn in NIST SP 800-53 Rev 5 and "
            "are covered by IA-2(1) and IA-2(2).\n"
        ),
        "| Control ID | Control Name | Review Focus | Verification Status | Evidence Source | Implementation Evidence in SCTM / SSP |",
        "| :--- | :--- | :--- | :--- | :--- | :--- |",
    ])
    for atc in senior_audit.get("atc_records", []):
        ver = atc.get("verification", "MISSING")
        assessed_status = atc.get("status") or "[NOT DETERMINED FROM SOURCE]"
        source = atc.get("evidence_source", "None")
        source_label = "None (control absent from package)" if source == "None" else source
        details = atc.get("details", "")
        if len(details) > 90:
            details = details[:87] + "..."
        lines.append(f"| **{atc['id']}** | {atc['name']} | {atc['description']} | `{ver}`: {assessed_status} | `{source_label}` | {details} |")
    lines.append("")

    acas_title = 'ACAS Vulnerability Scan "Good Data" Verification Rules'
    if acas_title in pta_sections:
        lines.append(pta_sections[acas_title])

    dynamic_updates = [s for s in stigs_required if any(k in s.get("version_source", "") for k in ("Override", "Live", "Cached"))]
    lines.extend([
        "## Mandatory DISA STIG & SRG Checklist Compliance Roadmap",
        "> [!IMPORTANT]",
        "> **AUTHORITATIVE DISA STIG SOURCE & DESKTOP STIG VIEWER APPLICATION**:",
        (
            "> 1. **Official STIG Downloads**: Official DISA STIG compilation packages and checklist benchmarks must be "
            "downloaded from the DoD Cyber Exchange at [https://public.cyber.mil/stigs/downloads/](https://public.cyber.mil/stigs/downloads/)."
        ),
        (
            "> 2. **DISA STIG Viewer Application**: The **DISA STIG Viewer desktop application** (downloadable from DoD "
            "Cyber Exchange at [https://public.cyber.mil/stigs/srg-stig-tools/](https://public.cyber.mil/stigs/srg-stig-tools/)) "
            "is the required software used to import, complete, evaluate, and export `.ckl` checklist files for eMASS ingestion."
        ),
        (
            "> 3. **Dynamic STIG Version Resolver**: All STIG/SRG checklist versions are dynamically resolved across "
            "institutional overrides, remote feeds, and authoritative baselines."
        ),
        (
            "> 4. **Public Web Reference**: The [STIG Viewer website](https://www.stigviewer.com/stigs) links provided "
            "below are reference links for web browsing and quick lookup of rule requirements and fix scripts.\n"
        ),
        (
            f"Based on the infrastructure components discovered in Terraform code, the cybersecurity team must complete, "
            f"verify, and document `{len(stigs_required)}` STIG Benchmark compliance checklists (`.ckl` files created in "
            f"the STIG Viewer desktop app). Dynamically updated or overridden versions: `{len(dynamic_updates)}`.\n"
        ),
        "| DISA STIG / SRG Benchmark | Version | Source | Public Reference Link | Triggering Component / Scope | Action Required for Assessor Team | Status |",
        "| :--- | :--- | :--- | :--- | :--- | :--- | :--- |",
    ])
    for s in stigs_required:
        src_label = s.get("version_source") or "Authoritative Baseline"
        lines.append(f"| **{s['title']}** | `{s['version']}` | {src_label} | [{s['slug']}]({s['url']}) | {s['reason']} | {s['action']} | `{s['status']}` |")
    lines.extend([
        "",
        "## Excel Workbooks (.xlsm) Audit Results",
        "| Workbook File | Audit Status | Sheet Structure | Populated Rows | Status Details |",
        "| :--- | :--- | :--- | :--- | :--- |",
    ])
    for xr in excel_results:
        sheets_str = ", ".join(xr["sheets"][:3])
        rows_str = ", ".join(f"{k}: {v}" for k, v in xr["rows_count"].items()) if xr["rows_count"] else "Preserved"
        details = ", ".join(xr["issues"]) if xr["issues"] else "No issues found"
        lines.append(f"| `{xr['file']}` | `{xr['status']}` | {sheets_str} | {rows_str} | {details} |")
    total_docx_hl = sum(r.get("hyperlinks_count", 0) for r in docx_results)
    lines.extend([
        "",
        "## Word Policy Documents (.docx) Audit Results",
        f"- Total DOCX Documents Checked: `{len(docx_results)}` files",
        f"- OpenXML Packaging: `{_status_note(docx_results, 'word/document.xml parsed')}` ({total_docx_hl} hyperlinks with relationship targets)",
    ])
    for dr in docx_results:
        if dr.get("status") != "PASS":
            lines.append(f"  - `{dr.get('file')}`: `{dr.get('status')}` {'; '.join(dr.get('issues', []))}")
    lines.append("")
    return lines


def _render_pta_variables_and_memo(
    gov: Dict[str, str],
    bundle: Dict[str, Any],
    pta_sections: Dict[str, str],
) -> List[str]:
    """Renders unresolved tokens, pending config variables, action highlights, ATO request memo, and WBS/References."""
    unresolved_tokens = bundle["unresolved_tokens"]
    config_required_vars = bundle["config_required_vars"]
    rmf_action_items = bundle["rmf_action_items"]
    senior_audit = bundle["senior_audit"]
    stigs_required = bundle["stigs_required"]
    lines: List[str] = []

    if unresolved_tokens:
        lines.extend([
            "## Critical Unresolved Syntax Tokens",
            "| Document Path | Line | Token Found | Raw Context Snippet |",
            "| :--- | :--- | :--- | :--- |",
        ])
        for u in unresolved_tokens:
            lines.append(
                f"| `{_escape_pta_code_cell(u['file'])}` | {u['line']} | "
                f"`{_escape_pta_code_cell(u['token'])}` | `{_escape_pta_code_cell(u['context'][:65])}` |"
            )
        lines.append("")

    if config_required_vars:
        grouped_vars: Dict[Tuple[str, str], List[int]] = {}
        for c in config_required_vars:
            grouped_vars.setdefault((c["file"], c["variable"]), []).append(c["line"])
        lines.extend([
            "## Pending Institutional Configuration Variables (`compliance_config.yaml`)",
            f"{len(config_required_vars)} placeholder occurrences in {len(grouped_vars)} distinct document/placeholder pairs.\n",
            "| Document Path | First Line | Occurrences | Placeholder Variable | Action Required |",
            "| :--- | :--- | :--- | :--- | :--- |",
        ])
        for (var_file, variable), var_lines in sorted(grouped_vars.items()):
            lines.append(
                f"| `{_escape_pta_code_cell(var_file)}` | {min(var_lines)} | {len(var_lines)} | "
                f"`{_escape_pta_code_cell(variable)}` | Define in `compliance_config.yaml` and re-run generator |"
            )
        lines.append("")

    exec_matrix_title = "Institutional Policy Manuals & Core Deliverables Human Execution Matrix"
    if exec_matrix_title in pta_sections:
        lines.append(pta_sections[exec_matrix_title])

    if rmf_action_items:
        lines.extend([
            "## Live Policy Customization & Action Item Alerts (Grouped by Document)",
            "The scanner identified the following action alerts across the generated policy manuals and security documentation:\n",
            "| Policy / Document Path | Action Items Count | Key Action Highlights |",
            "| :--- | :--- | :--- |",
        ])
        doc_actions: Dict[str, List[str]] = {}
        for item in rmf_action_items:
            doc_actions.setdefault(item["file"], []).append(item["action"])
        for doc_file in sorted(doc_actions.keys()):
            items = doc_actions[doc_file]
            highlights_str = "; ".join(list(dict.fromkeys(items))[:2])
            if len(highlights_str) > 130:
                highlights_str = highlights_str[:127] + "..."
            lines.append(
                f"| `{_escape_pta_code_cell(doc_file)}` | `{len(items)} items` | "
                f"{_escape_pta_table_cell(highlights_str)} |"
            )
        lines.append("")

    lines.extend([
        "## Sample Executive ATO Determination Request Memo Template\n",
        "```text",
        f"MEMORANDUM FOR: Authorizing Official (AO), {gov['org_name']}",
        f"FROM: Information System Security Manager (ISSM), {gov['sys_name']}",
        f"SUBJECT: Request for Authorization to Operate (ATO) for {gov['sys_name']} ({gov['sys_abbr']})\n",
        "1. PURPOSE:",
        (
            f"This memorandum requests an Authorization to Operate (ATO) for {gov['sys_name']} at {gov['impact_level']} / {gov['baseline']}, "
            "with [CONFIG_REQUIRED: Authorization Term and Continuous Monitoring Cycle].\n"
        ),
        "2. SYSTEM ARCHITECTURE & BOUNDARY:",
        "[CONFIG_REQUIRED: Boundary Summary - regions, encryption at rest, and remote access path as documented in the SSP]\n",
        "3. SECURITY ASSESSMENT & RESIDUAL RISK:",
        "[CONFIG_REQUIRED: Assessor (SCA / 3PAO) and Assessment Dates]",
        (
            f"- Automated validator findings at the time of this report: {len(senior_audit['cat_1_findings'])} CAT I / "
            f"{len(senior_audit['cat_2_findings'])} CAT II / {len(senior_audit['cat_3_findings'])} CAT III."
        ),
        "- Open POA&M Items by Residual Risk: [CONFIG_REQUIRED: Open POA&M Counts by Risk Level]",
        f"- DISA STIG Checklists: {len(stigs_required)} applicable STIG benchmarks identified for completion in STIG Viewer.\n",
        "4. RECOMMENDATION:",
        "[CONFIG_REQUIRED: ISSM Recommendation and Residual Risk Basis]\n",
        "____________________________________________",
        gov["issm_name"],
        gov["issm_title"],
        f"{gov['org_name']}\n",
        "CONCURRENCE:",
        "____________________________________________",
        gov["so_name"],
        gov["so_title"],
        gov["org_name"],
        "```\n",
    ])
    for tail_title in (
        "Work Breakdown Structure (WBS) for ATO",
        "Military Service Branch & Federal Agency Governance Overlays",
        "References",
    ):
        if tail_title in pta_sections:
            lines.append(pta_sections[tail_title])
    return lines


def _write_pta_outputs(
    ato_dir: str,
    report_lines: List[str],
    inventory: Dict[str, Any],
    policy_format: Optional[str],
    fix_drift: bool,
) -> str:
    """Writes Path_to_Authorization.md and/or .docx honoring policy_format; returns report_path."""
    report_path = os.path.join(ato_dir, "Path_to_Authorization.md")
    export_prefs_pta = (inventory.get("export_preferences") or inventory.get("compliance_config") or {}) if inventory else {}
    effective_policy_format = (policy_format or export_prefs_pta.get("policy_formats") or "both").strip().lower()
    write_pta_md = effective_policy_format in ("both", "markdown")
    write_pta_docx = effective_policy_format in ("both", "docx")
    report_text = "\n".join(report_lines)

    if write_pta_md:
        if os.path.islink(report_path):
            raise PermissionError(f"Refusing to overwrite symlinked report: '{report_path}'")
        write_text_file(report_path, report_text, allowed_boundary=ato_dir)
    elif fix_drift and os.path.exists(report_path) and not os.path.islink(report_path):
        os.remove(report_path)

    docx_path = os.path.join(ato_dir, "Path_to_Authorization.docx")
    if write_pta_docx:
        if os.path.islink(docx_path):
            raise PermissionError(f"Refusing to overwrite symlinked DOCX report: '{docx_path}'")
        ensure_path_within_boundary(docx_path, ato_dir, allow_symlinks=False)
        if docx_generator:
            docx_generator.convert_markdown_to_docx(report_text, docx_path, inventory, allowed_boundary=ato_dir)
    elif fix_drift and os.path.exists(docx_path) and not os.path.islink(docx_path):
        os.remove(docx_path)
    return report_path


def _log_validation_summary_and_emit_audit(
    target_dir: str,
    report_path: str,
    inventory: Dict[str, Any],
    bundle: Dict[str, Any],
    failed: bool,
) -> None:
    """Logs the validation summary banner and emits structured audit-log events."""
    yaml_results = bundle["yaml_results"]
    excel_results = bundle["excel_results"]
    docx_results = bundle["docx_results"]
    oscal_results = bundle["oscal_results"]
    runbook_results = bundle["runbook_results"]
    stigs_required = bundle["stigs_required"]
    senior_audit = bundle["senior_audit"]
    alignment_results = bundle["alignment_results"]
    semantic_linter_report = bundle["semantic_linter_report"]
    linter_crashed = bundle["linter_crashed"]
    app_info = inventory.get("application_components", {}) if inventory else {}
    apps_cnt = len(app_info.get("applications", []))
    pkgs_cnt = len(app_info.get("software_packages", []))
    ctr_cnt = len(app_info.get("container_images", []))
    ports_cnt = len(app_info.get("exposed_ports", []))
    total_docx_hl = sum(r.get("hyperlinks_count", 0) for r in docx_results)
    atc_total = senior_audit.get("atc_total", len(ATC_CONNECTION_CONTROLS))

    logger.info("=" * 80)
    logger.info("COMPLIANCE PACKAGE VALIDATION AUDIT COMPLETE")
    logger.info("=" * 80)
    logger.info("  • Markdown/YAML Audited      : %d", bundle["total_files_checked"])
    logger.info("  • Structured YAML Audited    : %d files (%s)", len(yaml_results), _status_note(yaml_results, "parsed as valid YAML"))
    logger.info("  • Excel Workbooks Audited    : %d (%s)", len(excel_results), _status_note(excel_results, "opened and checked"))
    logger.info(
        "  • Word Policy Manuals Audited: %d (%s; %d hyperlinks with relationship targets)",
        len(docx_results), _status_note(docx_results, "OpenXML structure checked"), total_docx_hl,
    )
    if oscal_results:
        oscal_ver_str = oscal_results[0].get("oscal_version") or "unknown"
        logger.info(
            "  • NIST OSCAL (v%s) Audited   : %d packages (%s; no JSON-schema validation)",
            oscal_ver_str, len(oscal_results), _status_note(oscal_results, "structure checked"),
        )
    if apps_cnt > 0 or pkgs_cnt > 0 or ctr_cnt > 0:
        logger.info("  • Application Components     : %d apps, %d packages, %d container images", apps_cnt, pkgs_cnt, ctr_cnt)
    if ports_cnt > 0:
        logger.info("  • App Ingress Ports Mapped   : %d ports in PPSM", ports_cnt)
    if runbook_results:
        passed_rb = sum(1 for r in runbook_results if r["status"] == "PASS")
        logger.info("  • IR Runbooks Audited        : %d/%d scenario workflows", passed_rb, len(_EXPECTED_RUNBOOKS))
    logger.info("  • Unresolved Tokens          : %d", len(bundle["unresolved_tokens"]))
    logger.info("  • Pending Placeholders       : %d", len(bundle["config_required_vars"]))
    dynamic_updates = [s for s in stigs_required if any(k in s.get("version_source", "") for k in ("Override", "Live", "Cached"))]
    logger.info("  • Applicable DISA STIGs      : %d (STIG Viewer catalog mapped, %d active updates/overrides)", len(stigs_required), len(dynamic_updates))
    logger.info("  • Architecture Reconciliation: %s", _format_coverage(alignment_results))
    missing_atc = [r for r in senior_audit.get("atc_records", []) if r.get("verification") == "MISSING"]
    if missing_atc:
        logger.info(
            "  • ATC Connection Controls    : %d/%d Verified (%d absent from SCTM and SSP: %s)",
            senior_audit["verified_atc_count"], atc_total, len(missing_atc), ", ".join(r["id"] for r in missing_atc),
        )
    else:
        logger.info("  • ATC Connection Controls    : %d/%d Verified", senior_audit["verified_atc_count"], atc_total)
    logger.info("  • Technical Readiness Status : %s", senior_audit["recommendation"])
    logger.info(
        "  • Assessor Findings (CAT I/II/III): %d / %d / %d",
        len(senior_audit["cat_1_findings"]), len(senior_audit["cat_2_findings"]), len(senior_audit["cat_3_findings"]),
    )
    if semantic_linter_report:
        logger.info(
            "  • Semantic Linter Audit      : %s (%d/%d passed, score: %.1f%%)",
            semantic_linter_report.summary["verdict"], semantic_linter_report.summary["passed_count"],
            semantic_linter_report.summary["total_artifacts_evaluated"], semantic_linter_report.summary["compliance_score_percent"],
        )
        logger.info(
            "  • Semantic & Drift Findings  : %d CAT I / %d CAT II / %d CAT III (%d drift findings)",
            semantic_linter_report.summary["cat_1_findings_count"], semantic_linter_report.summary["cat_2_findings_count"],
            semantic_linter_report.summary["cat_3_findings_count"], semantic_linter_report.summary["architectural_drift_findings_count"],
        )
    logger.info("  • ISSM Submission Roadmap    : Phase 1-6 itinerary + %d ATC controls", atc_total)
    logger.info("  • Master PTA & Audit Report  : '%s' (.md & .docx)", report_path)
    logger.info("=" * 80)

    audit_logger = get_audit_logger()
    for r in docx_results + excel_results + oscal_results:
        if r.get("status") in ("FAIL", "ERROR", "UNVERIFIED"):
            audit_logger.emit(
                AuditEvent.ARTIFACT_VALIDATED,
                outcome=AuditOutcome.DENIED,
                obj=r.get("file", "unknown"),
                detail={"issues": r.get("issues", [])},
            )

    audit_logger.emit(
        AuditEvent.ARTIFACT_VALIDATED,
        outcome=AuditOutcome.DENIED if failed else AuditOutcome.SUCCESS,
        obj=target_dir,
        detail={
            "cat_1_count": len(senior_audit.get("cat_1_findings", [])),
            "cat_2_count": len(senior_audit.get("cat_2_findings", [])),
            "verified_atc": senior_audit.get("verified_atc_count", 0),
            "linter_validated": semantic_linter_report is not None,
            "linter_verdict": semantic_linter_report.summary["verdict"] if semantic_linter_report else ("CRASHED" if linter_crashed else "N/A"),
            "linter_cat_1_count": semantic_linter_report.summary["cat_1_findings_count"] if semantic_linter_report else 0,
            "linter_drift_count": semantic_linter_report.summary["architectural_drift_findings_count"] if semantic_linter_report else 0,
        },
    )


def validate_compliance_package(
    target_dir: str,
    fix_drift: bool = False,
    fill_examples: bool = False,
    policy_format: Optional[str] = None,
    data_format: Optional[str] = None,
    oscal_format: Optional[str] = None,
    oscal_version: Optional[str] = None,
    update_stigs: bool = False,
    stigs_mode: Optional[str] = None,
    stigs_catalog: Optional[str] = None,
    run_linter: bool = True,
    linter_strict: bool = True,
    run_scanners: bool = False,
) -> bool:
    """Validates the generated compliance package, compiles Path_to_Authorization, and optionally repairs drift."""
    target_dir = str(resolve_path(target_dir))
    ato_dir = str(ensure_path_within_boundary(os.path.join(target_dir, "ato_artifacts"), target_dir))
    inventory_path = str(ensure_path_within_boundary(os.path.join(target_dir, "system_inventory.json"), target_dir))

    if not os.path.exists(ato_dir):
        logger.error("ATO artifacts directory not found at '%s'. Run generator first.", ato_dir)
        return False

    inventory, inventory_corrupt, inventory_error = _load_validator_inventory(inventory_path, target_dir)
    if isinstance(inventory.get("system_information"), Mapping):
        sys_info_check = inventory["system_information"]
        if any(k in sys_info_check for k in ("impact_level", "compliance_baseline", "governance_regime")):
            require_supported_regime(sys_info_check)
    drift_repair_failed = inventory_corrupt and fix_drift
    if fix_drift and not inventory_corrupt and inventory:
        drift_repair_failed = _run_drift_preflight_repair(
            target_dir, ato_dir, inventory, policy_format, data_format, oscal_format, oscal_version, run_scanners
        )

    all_files = _walk_artifact_files(ato_dir)
    md_and_yaml_files = [fp for fp in all_files if fp.endswith((".md", ".yaml"))]
    own_report = os.path.normcase(os.path.join(ato_dir, "Path_to_Authorization.md"))
    unresolved_tokens, config_required_vars, rmf_action_items = _scan_markdown_and_yaml_tokens(
        md_and_yaml_files, ato_dir, own_report, fill_examples
    )

    yaml_results = audit_yaml_syntax_integrity(ato_dir)
    excel_results = audit_excel_workbooks(ato_dir, unresolved_out=unresolved_tokens, placeholders_out=config_required_vars)
    docx_results = audit_docx_policies(ato_dir, unresolved_out=unresolved_tokens, placeholders_out=config_required_vars)
    runbook_results = _audit_runbooks_presence(ato_dir)
    oscal_results = audit_oscal_packages(ato_dir)
    stigs_required = evaluate_disa_stig_applicability(
        inventory, md_and_yaml_files, target_dir=target_dir,
        update_stigs=update_stigs, stigs_mode=stigs_mode, stigs_catalog=stigs_catalog,
    )
    alignment_results = audit_inventory_artifact_alignment(target_dir, inventory, ato_dir)

    semantic_linter_report = None
    linter_crashed = False
    linter_crash_error = ""
    if run_linter and run_semantic_linter is not None:
        logger.info("Executing Deterministic Semantic Linter & Architectural Drift Evaluation...")
        try:
            semantic_linter_report = run_semantic_linter(
                target_dir=target_dir, inventory=inventory, ato_dir=ato_dir, strict=linter_strict,
            )
        except (OSError, ValueError, TypeError, KeyError, RuntimeError) as err:
            linter_crashed = True
            linter_crash_error = str(err)
            logger.error("Semantic linter encountered error: %s", err)

    parity_results = audit_cross_artifact_parity(ato_dir, inventory)
    senior_audit = audit_senior_compliance_quality(
        target_dir=target_dir, inventory=inventory, ato_dir=ato_dir, alignment_res=alignment_results,
        excel_results=excel_results, docx_results=docx_results, oscal_results=oscal_results,
        unresolved_tokens=unresolved_tokens, config_required_vars=config_required_vars,
        stigs_required=stigs_required, rmf_action_items=rmf_action_items, yaml_results=yaml_results,
        semantic_linter_report=semantic_linter_report, linter_strict=linter_strict,
        linter_crashed=linter_crashed, linter_crash_error=linter_crash_error,
        inventory_corrupt=inventory_corrupt, inventory_error=inventory_error,
    )
    if parity_results.get("conflicts"):
        senior_audit["cat_2_findings"].append({
            "id": "CAT2-PARITY-001",
            "component": "Cross-Artifact Parity",
            "severity": "CAT II (Medium)",
            "description": "; ".join(parity_results["conflicts"][:3]),
            "remediation": "Re-run generator to synchronize system identity and control statuses across SSP, SCTM, POA&M, and OSCAL.",
        })

    gov = _extract_pta_governance_metadata(inventory)
    pta_sections = _load_pta_template_sections({"IMPACT_LEVEL": gov["impact_level"]})
    bundle: Dict[str, Any] = {
        "total_files_checked": len(md_and_yaml_files),
        "yaml_results": yaml_results,
        "excel_results": excel_results,
        "docx_results": docx_results,
        "oscal_results": oscal_results,
        "runbook_results": runbook_results,
        "stigs_required": stigs_required,
        "alignment_results": alignment_results,
        "senior_audit": senior_audit,
        "semantic_linter_report": semantic_linter_report,
        "linter_crashed": linter_crashed,
        "unresolved_tokens": unresolved_tokens,
        "config_required_vars": config_required_vars,
        "rmf_action_items": rmf_action_items,
    }
    report_lines = _render_pta_header_and_dashboard(gov, bundle)
    report_lines.extend(_render_pta_reconciliation_and_findings(bundle))
    report_lines.extend(_render_pta_atc_stigs_and_openxml(bundle, pta_sections))
    report_lines.extend(_render_pta_variables_and_memo(gov, bundle, pta_sections))
    report_path = _write_pta_outputs(ato_dir, report_lines, inventory, policy_format, fix_drift)

    has_packaging_failure = any(
        r.get("status") in ("FAIL", "UNVERIFIED", "ERROR")
        for r in (yaml_results + docx_results + excel_results + oscal_results)
    )
    blocking_cat_1 = [
        f for f in senior_audit.get("cat_1_findings", [])
        if f.get("source") != "semantic_linter" or linter_strict
    ]
    has_linter_failure = linter_strict and (
        linter_crashed
        or (semantic_linter_report is not None and not semantic_linter_report.passed)
    )
    failed = (
        drift_repair_failed
        or inventory_corrupt
        or has_packaging_failure
        or bool(unresolved_tokens)
        or bool(blocking_cat_1)
        or has_linter_failure
    )
    _log_validation_summary_and_emit_audit(target_dir, report_path, inventory, bundle, failed)
    return not failed


#: Seconds the example-data prompt waits for an answer before taking the
#: documented default (no example data). Only used on an attended terminal.
EXAMPLE_DATA_PROMPT_TIMEOUT_SECONDS: float = 60.0

_EXAMPLE_DATA_PROMPT = (
    "\n[PROMPT] Relabel the remaining [CONFIG_REQUIRED: ...] placeholders as "
    "[EXAMPLE_DATA_REQUIRED: ...] for an AI agent or reviewer to fill? (y/N): "
)


def _is_attended_terminal(
    stdin: Optional[TextIO] = None,
    stdout: Optional[TextIO] = None,
    environ: Optional[Mapping[str, str]] = None,
) -> bool:
    """Reports whether a person is plausibly present to answer a prompt.

    ``stdin.isatty()`` alone is not enough: agent shells and some CI runners
    attach a pseudo-terminal that nobody types into, and a prompt gated only on
    ``isatty()`` blocked those runs forever. A prompt is offered only when stdin
    and stdout are both terminals, ``CI`` is unset, and ``TERM`` is not
    ``dumb`` (the value agent shells and editor shells set).

    Args:
        stdin: Input stream to inspect. Defaults to ``sys.stdin``.
        stdout: Output stream to inspect. Defaults to ``sys.stdout``.
        environ: Environment to inspect. Defaults to ``os.environ``.

    Returns:
        True only when every check indicates an interactive session.
    """
    stdin = sys.stdin if stdin is None else stdin
    stdout = sys.stdout if stdout is None else stdout
    environ = os.environ if environ is None else environ
    try:
        if not (stdin.isatty() and stdout.isatty()):
            return False
    except (AttributeError, ValueError, OSError):
        # Closed, replaced, or detached streams cannot be prompted on.
        return False
    if environ.get("CI", "").strip().lower() not in ("", "0", "false"):
        return False
    if environ.get("TERM", "").strip().lower() == "dumb":
        return False
    return True


def _prompt_yes_no(prompt: str, timeout_seconds: float, stdin: TextIO, stdout: TextIO) -> bool:
    """Asks a yes/no question and returns True only for an explicit yes.

    On POSIX the wait is bounded with ``select`` so an unattended terminal that
    slipped past :func:`_is_attended_terminal` still cannot hang the run. Where
    the stream cannot be polled (for example a Windows console) the call falls
    back to a plain blocking read, which is the previous behavior.

    Args:
        prompt: Question text, written to ``stdout``.
        timeout_seconds: Maximum seconds to wait for an answer.
        stdin: Stream the answer is read from.
        stdout: Stream the prompt is written to.

    Returns:
        True for "y" or "yes" (any case); False for anything else, end of
        input, or no answer within the timeout.

    Raises:
        SystemExit: With status 130 when the operator presses Ctrl-C.
    """
    stdout.write(prompt)
    stdout.flush()
    try:
        try:
            ready, _, _ = select.select([stdin], [], [], timeout_seconds)
        except (OSError, ValueError, TypeError):
            answer = stdin.readline()
        else:
            if not ready:
                stdout.write("\n")
                stdout.flush()
                logger.info(
                    "No answer within %d seconds; continuing without example data.",
                    int(timeout_seconds),
                )
                return False
            answer = stdin.readline()
    except KeyboardInterrupt:
        # Ctrl-C means abort, not "continue with the default".
        logger.warning("Interrupted at prompt; aborting validation.")
        raise SystemExit(130)
    if not answer:
        # stdin closed mid-prompt: take the documented default (no example data).
        logger.info("No response received on stdin; defaulting to --no-fill-example-data.")
        return False
    return answer.strip().lower() in ("y", "yes")


def _resolve_fill_examples(
    argv: List[str],
    stdin: Optional[TextIO] = None,
    stdout: Optional[TextIO] = None,
    environ: Optional[Mapping[str, str]] = None,
    timeout_seconds: float = EXAMPLE_DATA_PROMPT_TIMEOUT_SECONDS,
) -> bool:
    """Decides whether example data is filled into the package.

    Explicit flags always win. Without a flag the operator is asked only on an
    attended terminal; every other session takes the documented default of no
    example data, so agents and CI never block on stdin.

    Args:
        argv: Command-line arguments (``sys.argv``).
        stdin: Input stream for the prompt. Defaults to ``sys.stdin``.
        stdout: Output stream for the prompt. Defaults to ``sys.stdout``.
        environ: Environment used to detect unattended sessions.
        timeout_seconds: Maximum seconds to wait for an answer.

    Returns:
        True when example data should be filled in.
    """
    if "--fill-example-data" in argv:
        return True
    if "--no-fill-example-data" in argv:
        return False
    stdin = sys.stdin if stdin is None else stdin
    stdout = sys.stdout if stdout is None else stdout
    if not _is_attended_terminal(stdin, stdout, environ):
        logger.info(
            "Non-interactive session: example data is not filled in "
            "(pass --fill-example-data to enable it)."
        )
        return False
    return _prompt_yes_no(_EXAMPLE_DATA_PROMPT, timeout_seconds, stdin, stdout)


def _build_arg_parser() -> argparse.ArgumentParser:
    """Constructs the CLI argument parser for validate_compliance_artifacts."""
    parser = argparse.ArgumentParser(
        prog="validate_compliance_artifacts.py",
        description="Senior Assessor Public Sector Verification Gate & Unified Authorization Playbook Compiler.",
        allow_abbrev=False,
    )
    parser.add_argument(
        "target_dir",
        help="Target blueprint folder containing ato_artifacts/ (required)",
    )
    parser.add_argument(
        "--fix",
        action="store_true",
        help=(
            "Regenerate the ENTIRE package to reconcile drift against live Terraform. "
            "Backs up the existing package to ato_artifacts_backup_<timestamp>/ first."
        ),
    )
    fill_group = parser.add_mutually_exclusive_group()
    fill_group.add_argument(
        "--fill-example-data",
        action="store_true",
        help=(
            "Relabel [CONFIG_REQUIRED: ...] placeholders as "
            "[EXAMPLE_DATA_REQUIRED: ...] for an AI agent or reviewer to fill"
        ),
    )
    fill_group.add_argument(
        "--no-fill-example-data",
        action="store_true",
        help=(
            "Leave placeholders as [CONFIG_REQUIRED: ...] (default). Without either "
            "flag the validator asks only on an interactive terminal."
        ),
    )
    parser.add_argument(
        "--policy-format",
        choices=("both", "docx", "markdown"),
        default=None,
        help="With --fix: override policy format (both, docx, markdown)",
    )
    parser.add_argument(
        "--data-format",
        choices=("both", "excel", "yaml"),
        default=None,
        help="With --fix: override structured data format (both, excel, yaml)",
    )
    parser.add_argument(
        "--oscal-format",
        choices=("both", "json", "yaml", "none"),
        default=None,
        help="With --fix: override OSCAL format (both, json, yaml, none)",
    )
    parser.add_argument(
        "--oscal-version",
        choices=tuple(SUPPORTED_OSCAL_VERSIONS),
        default=None,
        help="With --fix: target OSCAL specification version (e.g. 1.2.3 or 1.1.0)",
    )
    parser.add_argument(
        "--update-stigs",
        action="store_true",
        help="Check and refresh DISA STIG / SRG benchmark checklists",
    )
    parser.add_argument(
        "--stigs-mode",
        choices=("auto", "online", "offline"),
        default=None,
        help="STIG version resolution mode (auto, online, offline)",
    )
    parser.add_argument(
        "--stigs-catalog",
        default=None,
        help="Custom DISA STIG catalog URL or local JSON file path",
    )
    parser.add_argument(
        "--run-scanners",
        action="store_true",
        help=(
            "With --fix: run live scanners (Checkov, Semgrep) on a sanitized private copy "
            "of the target. COMPLIANCE_DISABLE_LIVE_SCANNERS=1 always disables them."
        ),
    )
    parser.add_argument(
        "--no-linter",
        action="store_true",
        help="Disable deterministic semantic linter evaluation",
    )
    parser.add_argument(
        "--no-linter-strict",
        action="store_true",
        help=(
            "Do not fail validation on CAT I semantic linter findings "
            "(structural, token, inventory, and assessor CAT I failures still fail closed)"
        ),
    )
    return parser


def main(argv: Optional[Sequence[str]] = None) -> None:
    """CLI entry point for compliance package validation and STIG evaluation."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    raw_argv = list(sys.argv[1:] if argv is None else argv)
    parser = _build_arg_parser()
    args = parser.parse_args(raw_argv)

    if not args.fix and any(
        val is not None
        for val in (args.policy_format, args.data_format, args.oscal_format, args.oscal_version)
    ):
        parser.error(
            "--policy-format, --data-format, --oscal-format, and --oscal-version require --fix"
        )

    try:
        target_abs = str(resolve_cli_target(args.target_dir))
    except (FileNotFoundError, NotADirectoryError, PermissionError) as err:
        print(f"[ERROR] {err}", file=sys.stderr)
        sys.exit(1)

    fill_examples = _resolve_fill_examples(raw_argv)

    try:
        success = validate_compliance_package(
            target_abs,
            fix_drift=args.fix,
            fill_examples=fill_examples,
            policy_format=args.policy_format,
            data_format=args.data_format,
            oscal_format=args.oscal_format,
            oscal_version=args.oscal_version,
            update_stigs=args.update_stigs,
            stigs_mode=args.stigs_mode,
            stigs_catalog=args.stigs_catalog,
            run_linter=not args.no_linter,
            linter_strict=not args.no_linter_strict,
            run_scanners=args.run_scanners,
        )
    except UnsupportedRegimeError as err:
        print(f"[ERROR] {err}", file=sys.stderr)
        sys.exit(1)
    if not success:
        sys.exit(1)


if __name__ == "__main__":
    main()
