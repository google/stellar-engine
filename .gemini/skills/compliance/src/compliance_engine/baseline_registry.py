# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Multi-regime baseline registry, control-set filtering, and parameter injection."""

from __future__ import annotations

from dataclasses import dataclass
import functools
import json
from pathlib import Path
import re
from typing import Any, Collection, Dict, Final, List, Mapping, Optional, Set, Tuple, Union
import yaml

from .file_helpers import (
    CJIS_CSP_PATO_PLACEHOLDER,
    DEFAULT_CSP_PATO_PACKAGE_ID,
    DOD_IL6_CSP_PATO_PLACEHOLDER,
    DOD_REGIMES,
    HIGH_BASELINE_REGIMES,
    LOW_BASELINE_REGIMES,
    MODERATE_BASELINE_REGIMES,
    Regime,
    classify_regime,
    get_skill_root,
    is_supported_regime,
    require_supported_regime,
)
from .poam_rules import base_control_id, normalize_control_id

_FEDRAMP_PROFILE_BASE: Final[str] = (
    "https://raw.githubusercontent.com/OSCAL-Foundation/fedramp-resources/main/baselines/rev5/json/"
)
_NIST_PROFILE_BASE: Final[str] = (
    "https://raw.githubusercontent.com/usnistgov/oscal-content/main/nist.gov/SP800-53/rev5/json/"
)


@dataclass(frozen=True)
class RegimeSpec:
    """Declarative specification for a supported compliance regime."""

    regime: Regime
    regime_id: str
    display_name: str
    tier: int
    baseline_filename: str
    ssp_template_filename: str
    default_impacts: Tuple[str, str, str]
    is_dod: bool
    is_fedramp: bool
    is_fisma: bool
    oscal_profile_href: str
    oscal_profile_remarks: Optional[str] = None
    csp_inheritance_header: str = ""
    default_csp_pato_package_id: str = DEFAULT_CSP_PATO_PACKAGE_ID


REGIME_SPEC_BY_REGIME: Final[Dict[Regime, RegimeSpec]] = {
    Regime.FEDRAMP_LOW: RegimeSpec(
        regime=Regime.FEDRAMP_LOW,
        regime_id="fedramp_low",
        display_name="FedRAMP Low",
        tier=1,
        baseline_filename="fedramp_rev5_low.json",
        ssp_template_filename="SSP_FedRAMP_High_Template.md",
        default_impacts=("Low", "Low", "Low"),
        is_dod=False,
        is_fedramp=True,
        is_fisma=False,
        oscal_profile_href=_FEDRAMP_PROFILE_BASE + "FedRAMP_rev5_LOW-baseline_profile.json",
        csp_inheritance_header="Inherited from the Google Services FedRAMP P-ATO (Package ID: {{ CSP_PATO_PACKAGE_ID }}).",
    ),
    Regime.FEDRAMP_MODERATE: RegimeSpec(
        regime=Regime.FEDRAMP_MODERATE,
        regime_id="fedramp_moderate",
        display_name="FedRAMP Moderate",
        tier=1,
        baseline_filename="fedramp_rev5_moderate.json",
        ssp_template_filename="SSP_FedRAMP_High_Template.md",
        default_impacts=("Moderate", "Moderate", "Moderate"),
        is_dod=False,
        is_fedramp=True,
        is_fisma=False,
        oscal_profile_href=_FEDRAMP_PROFILE_BASE + "FedRAMP_rev5_MODERATE-baseline_profile.json",
        csp_inheritance_header="Inherited from the Google Services FedRAMP P-ATO (Package ID: {{ CSP_PATO_PACKAGE_ID }}).",
    ),
    Regime.FEDRAMP_HIGH: RegimeSpec(
        regime=Regime.FEDRAMP_HIGH,
        regime_id="fedramp_high",
        display_name="FedRAMP High",
        tier=1,
        baseline_filename="fedramp_rev5_high.json",
        ssp_template_filename="SSP_FedRAMP_High_Template.md",
        default_impacts=("High", "High", "High"),
        is_dod=False,
        is_fedramp=True,
        is_fisma=False,
        oscal_profile_href=_FEDRAMP_PROFILE_BASE + "FedRAMP_rev5_HIGH-baseline_profile.json",
        csp_inheritance_header="Inherited from the Google Services FedRAMP High P-ATO (Package ID: {{ CSP_PATO_PACKAGE_ID }}).",
    ),
    Regime.FISMA_LOW: RegimeSpec(
        regime=Regime.FISMA_LOW,
        regime_id="fisma_low",
        display_name="FISMA Low",
        tier=1,
        baseline_filename="nist_800_53b_low.json",
        ssp_template_filename="SSP_FedRAMP_High_Template.md",
        default_impacts=("Low", "Low", "Low"),
        is_dod=False,
        is_fedramp=False,
        is_fisma=True,
        oscal_profile_href=_NIST_PROFILE_BASE + "NIST_SP-800-53_rev5_LOW-baseline_profile.json",
        csp_inheritance_header=(
            "Inherited from the Google Cloud common infrastructure authorization boundary "
            "(Package ID: {{ CSP_PATO_PACKAGE_ID }})."
        ),
    ),
    Regime.FISMA_MODERATE: RegimeSpec(
        regime=Regime.FISMA_MODERATE,
        regime_id="fisma_moderate",
        display_name="FISMA Moderate",
        tier=1,
        baseline_filename="nist_800_53b_moderate.json",
        ssp_template_filename="SSP_FedRAMP_High_Template.md",
        default_impacts=("Moderate", "Moderate", "Moderate"),
        is_dod=False,
        is_fedramp=False,
        is_fisma=True,
        oscal_profile_href=_NIST_PROFILE_BASE + "NIST_SP-800-53_rev5_MODERATE-baseline_profile.json",
        csp_inheritance_header=(
            "Inherited from the Google Cloud common infrastructure authorization boundary "
            "(Package ID: {{ CSP_PATO_PACKAGE_ID }})."
        ),
    ),
    Regime.FISMA_HIGH: RegimeSpec(
        regime=Regime.FISMA_HIGH,
        regime_id="fisma_high",
        display_name="FISMA High",
        tier=1,
        baseline_filename="nist_800_53b_high.json",
        ssp_template_filename="SSP_FedRAMP_High_Template.md",
        default_impacts=("High", "High", "High"),
        is_dod=False,
        is_fedramp=False,
        is_fisma=True,
        oscal_profile_href=_NIST_PROFILE_BASE + "NIST_SP-800-53_rev5_HIGH-baseline_profile.json",
        csp_inheritance_header=(
            "Inherited from the Google Cloud common infrastructure authorization boundary "
            "(Package ID: {{ CSP_PATO_PACKAGE_ID }})."
        ),
    ),
    Regime.DOD_IL2: RegimeSpec(
        regime=Regime.DOD_IL2,
        regime_id="dod_il2",
        display_name="DoD IL2",
        tier=2,
        baseline_filename="dod_il2.json",
        ssp_template_filename="SSP_IL5_Template.md",
        default_impacts=("Low", "Low", "Low"),
        is_dod=True,
        is_fedramp=False,
        is_fisma=False,
        oscal_profile_href=_FEDRAMP_PROFILE_BASE + "FedRAMP_rev5_MODERATE-baseline_profile.json",
        oscal_profile_remarks=(
            "DoD IL2 provides reciprocity with the FedRAMP Moderate baseline per DISA CSP SRG V1R7 Section 3.7.5 "
            "and Table 3-1 (323 controls); no official DoD Cloud Computing SRG OSCAL overlay is published, "
            "so SRG-specific enclave requirements are documented in the SSP narrative."
        ),
        csp_inheritance_header=(
            "Inherited from the Google Services FedRAMP Moderate/High P-ATO "
            "(Package ID: {{ CSP_PATO_PACKAGE_ID }}) and DoD IL2 PA."
        ),
    ),
    Regime.DOD_IL4: RegimeSpec(
        regime=Regime.DOD_IL4,
        regime_id="dod_il4",
        display_name="DoD IL4",
        tier=2,
        baseline_filename="dod_il4.json",
        ssp_template_filename="SSP_IL5_Template.md",
        default_impacts=("Moderate", "Moderate", "Moderate"),
        is_dod=True,
        is_fedramp=False,
        is_fisma=False,
        oscal_profile_href=_FEDRAMP_PROFILE_BASE + "FedRAMP_rev5_MODERATE-baseline_profile.json",
        oscal_profile_remarks=(
            "DoD IL4 overlays the FedRAMP Moderate baseline per DISA CSP SRG V1R7 Table 3-1 and Appendix D "
            "Table D-1 (334 controls); no official DoD Cloud Computing SRG OSCAL overlay is published, "
            "so SRG-specific parameters and controls are documented in the SSP narrative."
        ),
        csp_inheritance_header=(
            "Inherited from the Google Services FedRAMP Moderate P-ATO "
            "(Package ID: {{ CSP_PATO_PACKAGE_ID }}) and DoD IL4 PA."
        ),
    ),
    Regime.DOD_IL5: RegimeSpec(
        regime=Regime.DOD_IL5,
        regime_id="dod_il5",
        display_name="DoD IL5",
        tier=2,
        baseline_filename="dod_il5.json",
        ssp_template_filename="SSP_IL5_Template.md",
        default_impacts=("High", "High", "High"),
        is_dod=True,
        is_fedramp=False,
        is_fisma=False,
        oscal_profile_href=_FEDRAMP_PROFILE_BASE + "FedRAMP_rev5_HIGH-baseline_profile.json",
        oscal_profile_remarks=(
            "DoD IL5 overlays the FedRAMP High baseline per DISA CSP SRG V1R7 Table 3-1, Appendix D Table D-1, "
            "and CNSSI No. 1253 (624 controls); no official DoD Cloud Computing SRG OSCAL overlay is published, "
            "so SRG-specific parameters and controls are documented in the SSP narrative."
        ),
        csp_inheritance_header=(
            "Inherited from the Google Services FedRAMP High P-ATO "
            "(Package ID: {{ CSP_PATO_PACKAGE_ID }}) and DoD IL5 PA."
        ),
    ),
    Regime.DOD_IL6: RegimeSpec(
        regime=Regime.DOD_IL6,
        regime_id="dod_il6",
        display_name="DoD IL6",
        tier=2,
        baseline_filename="dod_il6.json",
        ssp_template_filename="SSP_IL5_Template.md",
        default_impacts=("High", "High", "High"),
        is_dod=True,
        is_fedramp=False,
        is_fisma=False,
        oscal_profile_href=_FEDRAMP_PROFILE_BASE + "FedRAMP_rev5_HIGH-baseline_profile.json",
        oscal_profile_remarks=(
            "DoD IL6 overlays the FedRAMP High baseline per DISA CSP SRG V1R7 Table 3-1, Appendix D Table D-1 (IL6), "
            "and the CNSSI No. 1253 Classified Information Overlay (627 controls); no official DoD Cloud Computing "
            "SRG OSCAL overlay is published, so SRG-specific classified enclave requirements are documented in the "
            "SSP narrative."
        ),
        csp_inheritance_header=(
            "[CONFIG_REQUIRED: Classified Cloud / GDC Secret Enclave Inheritance Boundary — DoD IL6 PA "
            "(Verify SIPRNet / NSA Type 1 Crypto / Cross-Domain Solution Inheritance)]"
        ),
        default_csp_pato_package_id=DOD_IL6_CSP_PATO_PLACEHOLDER,
    ),
    Regime.CJIS: RegimeSpec(
        regime=Regime.CJIS,
        regime_id="cjis",
        display_name="CJIS",
        tier=3,
        baseline_filename="cjis.json",
        ssp_template_filename="SSP_FedRAMP_High_Template.md",
        default_impacts=("Moderate", "Moderate", "Moderate"),
        is_dod=False,
        is_fedramp=False,
        is_fisma=False,
        oscal_profile_href=_NIST_PROFILE_BASE + "NIST_SP-800-53_rev5_MODERATE-baseline_profile.json",
        oscal_profile_remarks=(
            "CJIS applies the FBI CJIS Security Policy v5.9.4 / v6.0 NIST SP 800-53 Rev. 5 Moderate harmonized "
            "baseline overlay (323 controls); no official FBI CJIS OSCAL profile is published, so CJIS Security "
            "Policy parameter overrides and State CSA obligations are documented in the SSP narrative and "
            "set-parameters."
        ),
        csp_inheritance_header=(
            "Inherited from Google Cloud Assured Workloads for CJIS "
            "(FBI CJIS Security Policy v5.9.4 / v6.0 Data Boundary — {{ CSP_PATO_PACKAGE_ID }})."
        ),
        default_csp_pato_package_id=CJIS_CSP_PATO_PLACEHOLDER,
    ),
}


def get_baselines_dir() -> Path:
    """Return path to ``.gemini/skills/compliance/config/baselines``."""
    return get_skill_root() / "config" / "baselines"


def get_regime_spec(regime: Regime) -> RegimeSpec:
    """Return the :class:`RegimeSpec` for ``regime``."""
    if regime not in REGIME_SPEC_BY_REGIME:
        require_supported_regime(regime.value)
        raise KeyError(f"No RegimeSpec registered for {regime!r}")
    return REGIME_SPEC_BY_REGIME[regime]


@functools.lru_cache(maxsize=16)
def _load_baseline_cached(regime: Regime) -> Dict[str, Any]:
    spec = get_regime_spec(regime)
    path = get_baselines_dir() / spec.baseline_filename
    return json.loads(path.read_text(encoding="utf-8"))


def get_regime_baseline(regime: Regime) -> Dict[str, Any]:
    """Load and return the normalized baseline dictionary for ``regime``."""
    return _load_baseline_cached(regime)


def build_regime_template_flags(
    regime: Regime,
    inventory: Optional[Mapping[str, Any]] = None,
) -> Dict[str, bool]:
    """Build conditional evaluation flags for ``evaluate_template_conditionals``."""
    inv = inventory or {}
    raw_sys = inv.get("system_information")
    sys_info: Mapping[str, Any] = raw_sys if isinstance(raw_sys, Mapping) else {}
    raw_sec = inv.get("security_operations")
    raw_sys_sec = sys_info.get("security_operations")
    sec_ops: Mapping[str, Any] = (
        raw_sec
        if isinstance(raw_sec, Mapping)
        else (raw_sys_sec if isinstance(raw_sys_sec, Mapping) else {})
    )
    scc_enabled = bool(sec_ops.get("scc_enabled") if "scc_enabled" in sec_ops else sys_info.get("scc_enabled"))
    secops_enabled = bool(
        sec_ops.get("secops_enabled") if "secops_enabled" in sec_ops else sys_info.get("secops_enabled")
    )
    cssp_provider = str(sec_ops.get("cssp_provider") or sys_info.get("cssp_provider") or "DISA").strip()
    ext_siem = str(sec_ops.get("external_siem_type") or sys_info.get("external_siem_type") or "").strip()

    is_dod = regime in DOD_REGIMES
    is_fedramp = regime in (Regime.FEDRAMP_LOW, Regime.FEDRAMP_MODERATE, Regime.FEDRAMP_HIGH)
    is_fisma = regime in (Regime.FISMA_LOW, Regime.FISMA_MODERATE, Regime.FISMA_HIGH)
    is_high = regime in HIGH_BASELINE_REGIMES
    is_mod = regime in MODERATE_BASELINE_REGIMES
    is_low = regime in LOW_BASELINE_REGIMES

    return {
        "SCC_ENABLED": scc_enabled,
        "SCC_DISABLED": not scc_enabled,
        "SECOPS_ENABLED": secops_enabled,
        "SECOPS_DISABLED": not secops_enabled,
        "CSSP_ENABLED": bool(cssp_provider and cssp_provider.lower() not in ("none", "not applicable", "n/a", "")),
        "CSSP_DISABLED": not bool(
            cssp_provider and cssp_provider.lower() not in ("none", "not applicable", "n/a", "")
        ),
        "EXTERNAL_SIEM_ENABLED": bool(ext_siem and ext_siem.lower() not in ("none", "not applicable", "n/a", "")),
        "DOD_ENCLAVE": is_dod,
        "DOD_IL2": regime == Regime.DOD_IL2,
        "DOD_IL4": regime == Regime.DOD_IL4,
        "DOD_IL5": regime == Regime.DOD_IL5,
        "DOD_IL6": regime == Regime.DOD_IL6,
        "CJIS_ENCLAVE": regime == Regime.CJIS,
        "FEDRAMP_ENCLAVE": is_fedramp,
        "FISMA_ENCLAVE": is_fisma,
        "CIVILIAN_ENCLAVE": not is_dod,
        "COMMERCIAL_ENCLAVE": not is_dod,
        "HIGH_BASELINE": is_high,
        "MODERATE_BASELINE": is_mod,
        "LOW_BASELINE": is_low,
        "MODERATE_OR_HIGH_BASELINE": is_high or is_mod,
        "ALLOW_UNACCREDITED_SCC": bool(sec_ops.get("allow_unaccredited_scc_in_il5", False)),
    }


def _lookup_param_value(
    baseline: Mapping[str, Any],
    param_id: str,
    fallback: str,
) -> str:
    params = baseline.get("parameters", {})
    if isinstance(params, Mapping) and param_id in params:
        val = params[param_id].get("value")
        if val:
            return str(val)
    return fallback


def _resolve_dod_enclave_labels(regime: Regime, spec: RegimeSpec) -> Tuple[str, str, str, str, str, str, str, str]:
    """Resolves (grc_label, csp_inh_auth, short, label, pa_label, pa_full, base_fr, framework) for regime."""
    is_dod = regime in DOD_REGIMES
    is_fisma = regime in (Regime.FISMA_LOW, Regime.FISMA_MODERATE, Regime.FISMA_HIGH)
    if is_dod:
        grc_label = "DoD RMF / eMASS"
    elif is_fisma:
        grc_label = "FISMA / NIST SP 800-53B"
    elif regime == Regime.CJIS:
        grc_label = "FBI CJIS Security Policy / State CSA"
    else:
        grc_label = "FedRAMP"

    if regime == Regime.FEDRAMP_HIGH:
        csp_inh_auth = "Google Services FedRAMP High P-ATO"
    elif regime in (Regime.FEDRAMP_MODERATE, Regime.FEDRAMP_LOW):
        csp_inh_auth = "Google Services FedRAMP P-ATO"
    elif regime == Regime.CJIS:
        csp_inh_auth = "Google Cloud Assured Workloads for CJIS (FBI CJIS Security Policy v5.9.4 / v6.0 Data Boundary)"
    elif regime == Regime.DOD_IL6:
        csp_inh_auth = spec.csp_inheritance_header
    else:
        csp_inh_auth = "Google Cloud common infrastructure authorization boundary"

    if regime == Regime.DOD_IL2:
        return (
            grc_label, csp_inh_auth, "IL2", "DoD IL2", "DoD IL2 PA",
            "DoD Impact Level 2 Provisional Authorization (PA)", "FedRAMP Moderate/High", "DISA CSP SRG V1R7 (DoD IL2)",
        )
    if regime == Regime.DOD_IL4:
        return (
            grc_label, csp_inh_auth, "IL4", "DoD IL4", "DoD IL4 PA",
            "DoD Impact Level 4 Provisional Authorization (PA)", "FedRAMP Moderate", "DISA CSP SRG V1R7 (DoD IL4)",
        )
    if regime == Regime.DOD_IL6:
        return (
            grc_label,
            csp_inh_auth,
            "IL6",
            "DoD IL6",
            "[CONFIG_REQUIRED: Classified Cloud / GDC Secret Enclave DoD IL6 PA]",
            "[CONFIG_REQUIRED: Classified Cloud / GDC Secret Enclave DoD Impact Level 6 Provisional Authorization (PA)]",
            "[CONFIG_REQUIRED: Classified Cloud / GDC Secret Enclave Baseline]",
            "DISA CSP SRG V1R7 + CNSSI 1253 Classified Overlay (DoD IL6)",
        )
    return (grc_label, csp_inh_auth, "IL5", "DoD IL5", "DoD IL5 PA", "DoD Impact Level 5 Provisional Authorization (PA)", "FedRAMP High", "CNSSI 1253")


def _resolve_fisma_or_fedramp_odps(
    baseline: Mapping[str, Any], is_fisma: bool, is_high: bool
) -> Tuple[str, str, str, str, str, str, str, str]:
    """Resolves the 8 conditional ODP parameter values for FISMA vs FedRAMP/DoD/CJIS."""
    if is_fisma:
        return (
            _lookup_param_value(baseline, "ac-02.02_odp.02", "[CONFIG_REQUIRED: Organization-defined time period for each type of account (ac-02.02_odp.02)]"),
            _lookup_param_value(baseline, "ac-02.03_odp.01", "[CONFIG_REQUIRED: Organization-defined time period (ac-02.03_odp.01)]"),
            _lookup_param_value(baseline, "ac-02.03_odp.02", "[CONFIG_REQUIRED: Organization-defined time period (ac-02.03_odp.02)]"),
            _lookup_param_value(baseline, "ac-10_odp.02", "[CONFIG_REQUIRED: Organization-defined number (ac-10_odp.02)]"),
            _lookup_param_value(baseline, "cm-08.03_odp.01", "[CONFIG_REQUIRED: Organization-defined frequency (cm-08.03_odp.01)]"),
            _lookup_param_value(baseline, "sc-10_odp", "[CONFIG_REQUIRED: Organization-defined time period (sc-10_odp)]"),
            _lookup_param_value(baseline, "sc-45.01_odp.01", "[CONFIG_REQUIRED: Organization-defined frequency (sc-45.01_odp.01)]"),
            _lookup_param_value(baseline, "si-06_odp.05", "[CONFIG_REQUIRED: Organization-defined frequency (si-06_odp.05)]"),
        )
    ac2_3_inact_raw = _lookup_param_value(baseline, "ac-02.03_odp.02", "thirty-five (35) days" if is_high else "ninety (90) days")
    sc45_freq = _lookup_param_value(baseline, "sc-45.01_odp.01", "At least hourly")
    sc45_src = _lookup_param_value(baseline, "sc-45.01_odp.02", "http://tf.nist.gov/tf-cgi/servers.cgi")
    sc45_diff = _lookup_param_value(baseline, "sc-45.01_odp.03", "any difference")
    return (
        _lookup_param_value(baseline, "ac-02.02_odp.02", "no more than 24 hours from last use" if is_high else "no more than 96 hours from last use"),
        _lookup_param_value(baseline, "ac-02.03_odp.01", "24 hours for user accounts"),
        ac2_3_inact_raw.split(" (See")[0].strip(),
        _lookup_param_value(baseline, "ac-10_odp.02", "three (3) sessions for privileged access and two (2) sessions for non-privileged access"),
        _lookup_param_value(baseline, "cm-08.03_odp.01", "automated mechanisms with a maximum five-minute delay in detection"),
        _lookup_param_value(
            baseline, "sc-10_odp",
            "no longer than ten (10) minutes for privileged sessions and no longer than fifteen (15) minutes for user sessions",
        ),
        f"{sc45_freq} against {sc45_src} upon {sc45_diff}",
        _lookup_param_value(baseline, "si-06_odp.05", "at least monthly"),
    )


def build_regime_odp_tokens(
    regime: Regime,
    inventory: Optional[Mapping[str, Any]] = None,
) -> Dict[str, str]:
    """Construct regime-specific ODP and enclave label tokens from ``config/baselines/<regime>.json``."""
    del inventory
    spec = get_regime_spec(regime)
    baseline = get_regime_baseline(regime)
    is_fisma = regime in (Regime.FISMA_LOW, Regime.FISMA_MODERATE, Regime.FISMA_HIGH)
    is_high = regime in HIGH_BASELINE_REGIMES

    grc_label, csp_inh_auth, dod_impact_short, dod_impact_label, dod_pa_label, dod_pa_full, base_fedramp_label, dod_framework_label = (
        _resolve_dod_enclave_labels(regime, spec)
    )
    ac2_2_val, ac2_3_dis, ac2_3_inact, ac10_val, cm8_3_val, sc10_val, sc45_1_val, si6_val = (
        _resolve_fisma_or_fedramp_odps(baseline, is_fisma, is_high)
    )

    return {
        "{{ IMPACT_LEVEL }}": regime.value,
        "{{ REGIME_DISPLAY_NAME }}": regime.value,
        "{{ GRC_STANDARD_LABEL }}": grc_label,
        "{{ CSP_INHERITANCE_AUTHORIZATION }}": csp_inh_auth,
        "{{ CSP_INHERITANCE_HEADER }}": spec.csp_inheritance_header,
        "{{ DOD_IMPACT_SHORT }}": dod_impact_short,
        "{{ DOD_IMPACT_LABEL }}": dod_impact_label,
        "{{ DOD_PA_LABEL }}": dod_pa_label,
        "{{ DOD_PA_FULL_LABEL }}": dod_pa_full,
        "{{ BASE_FEDRAMP_LABEL }}": base_fedramp_label,
        "{{ DOD_BASELINE_FRAMEWORK_LABEL }}": dod_framework_label,
        "{{ ODP_POLICY_REVIEW_FREQ }}": _lookup_param_value(
            baseline,
            "ac-01_odp.05",
            "[CONFIG_REQUIRED: Organization-defined policy review frequency (ac-01_odp.05)]",
        ),
        "{{ ODP_AC_2_REVIEW_FREQ }}": _lookup_param_value(
            baseline,
            "ac-02_odp.10",
            "[CONFIG_REQUIRED: Organization-defined account review frequency (ac-02_odp.10)]",
        ),
        "{{ ODP_AC_2_2_DISABLE_WINDOW }}": ac2_2_val,
        "{{ ODP_AC_2_3_DISABLE_WINDOW }}": ac2_3_dis,
        "{{ ODP_AC_2_3_INACTIVITY_WINDOW }}": ac2_3_inact,
        "{{ ODP_AC_10_SESSIONS }}": ac10_val,
        "{{ ODP_CA_5_POAM_FREQ }}": _lookup_param_value(
            baseline,
            "ca-05_odp",
            "[CONFIG_REQUIRED: Organization-defined frequency (ca-05_odp)]",
        ),
        "{{ ODP_CM_8_3_DETECTION_DELAY }}": cm8_3_val,
        "{{ ODP_CP_4_TEST_FREQ }}": _lookup_param_value(
            baseline,
            "cp-04_odp.01",
            "[CONFIG_REQUIRED: Organization-defined frequency (cp-04_odp.01)]",
        ),
        "{{ ODP_CP_9_BACKUP_FREQ }}": _lookup_param_value(
            baseline,
            "cp-09_odp.02",
            "[CONFIG_REQUIRED: Organization-defined frequency (cp-09_odp.02)]",
        ),
        "{{ ODP_IR_2_TRAINING_WINDOW }}": _lookup_param_value(
            baseline,
            "ir-02_odp.01",
            "[CONFIG_REQUIRED: Organization-defined time period (ir-02_odp.01)]",
        ),
        "{{ ODP_PE_2_REVIEW_FREQ }}": _lookup_param_value(
            baseline,
            "pe-02_odp",
            "[CONFIG_REQUIRED: Organization-defined frequency (pe-02_odp)]",
        ),
        "{{ ODP_PE_6_REVIEW_FREQ }}": _lookup_param_value(
            baseline,
            "pe-06_odp.01",
            "[CONFIG_REQUIRED: Organization-defined frequency (pe-06_odp.01)]",
        ),
        "{{ ODP_PE_8_REVIEW_FREQ }}": _lookup_param_value(
            baseline,
            "pe-08_odp.02",
            "[CONFIG_REQUIRED: Organization-defined frequency (pe-08_odp.02)]",
        ),
        "{{ ODP_PS_4_TERMINATION_WINDOW }}": _lookup_param_value(
            baseline,
            "ps-04_odp.01",
            "[CONFIG_REQUIRED: Organization-defined time period (ps-04_odp.01)]",
        ),
        "{{ ODP_PS_7_NOTIFICATION_WINDOW }}": _lookup_param_value(
            baseline,
            "ps-07_odp.02",
            "[CONFIG_REQUIRED: Organization-defined time period (ps-07_odp.02)]",
        ),
        "{{ ODP_RA_3_ASSESSMENT_FREQ }}": _lookup_param_value(
            baseline,
            "ra-03_odp.03",
            "[CONFIG_REQUIRED: Organization-defined frequency (ra-03_odp.03)]",
        ),
        "{{ ODP_SC_10_INACTIVITY_TIMEOUT }}": sc10_val,
        "{{ ODP_SC_45_1_SYNC_FREQ }}": sc45_1_val,
        "{{ ODP_SI_6_VERIFY_FREQ }}": si6_val,
    }


_SSP_HEADING_RE = re.compile(r"(?m)^(##\s+[^\n]+|###\s+([A-Z]{2}-\d+(?:\(\d+\))?)(?=\s|$)[^\n]*)$")
_POLICY_APPENDIX_ROW_RE = re.compile(
    r"^[ \t]*\|[ \t]*\**`?([A-Z]{2})-0*(\d+)(?:\s*\(0*(\d+)\))?`?\**[ \t]*\|"
)


def _strip_inline_fedramp_high_values(text: str) -> str:
    """Strip inline `` — FedRAMP High: ...`` suffixes inside ``[Assignment: ...]`` / ``[Selection: ...]``."""
    marker = " — FedRAMP High:"
    if marker not in text:
        return text
    out: List[str] = []
    i = 0
    n = len(text)
    while i < n:
        if text[i] == "[" and (
            text.startswith("[Assignment:", i) or text.startswith("[Selection", i)
        ):
            depth = 1
            j = i + 1
            marker_pos = -1
            while j < n and depth > 0:
                if depth == 1 and marker_pos == -1 and text.startswith(marker, j):
                    marker_pos = j
                if text[j] == "[":
                    depth += 1
                elif text[j] == "]":
                    depth -= 1
                j += 1
            if depth == 0 and marker_pos != -1:
                out.append(text[i:marker_pos])
                out.append("]")
                i = j
                continue
        out.append(text[i])
        i += 1
    return "".join(out)


def filter_ssp_control_sections(
    raw_ssp_markdown: str,
    allowed_control_ids: Collection[str],
) -> str:
    """Filter Section 2 of an SSP Markdown template to ``allowed_control_ids``."""
    allowed_set: Set[str] = {normalize_control_id(c) for c in allowed_control_ids}
    sec2_marker = "# 2. Minimum Security Controls"
    idx = raw_ssp_markdown.find(sec2_marker)
    if idx == -1:
        return raw_ssp_markdown

    header_part = raw_ssp_markdown[: idx + len(sec2_marker)]
    sec2_body = raw_ssp_markdown[idx + len(sec2_marker) :]

    matches = list(_SSP_HEADING_RE.finditer(sec2_body))
    if not matches:
        return raw_ssp_markdown

    kept_families: List[Tuple[str, List[str]]] = []
    current_family_header = ""
    current_controls: List[str] = []

    for m_idx, m in enumerate(matches):
        block_start = m.start()
        block_end = matches[m_idx + 1].start() if m_idx + 1 < len(matches) else len(sec2_body)
        block_text = sec2_body[block_start:block_end]
        ctrl_id = m.group(2)
        if ctrl_id is None:
            # Family heading (## 2.x ...)
            if current_family_header and current_controls:
                kept_families.append((current_family_header, current_controls))
            current_family_header = block_text
            current_controls = []
        else:
            norm_cid = normalize_control_id(ctrl_id)
            if norm_cid in allowed_set:
                current_controls.append(block_text)

    if current_family_header and current_controls:
        kept_families.append((current_family_header, current_controls))

    rebuilt_sec2 = ["\n\n\n"]
    for fam_header, ctrl_blocks in kept_families:
        rebuilt_sec2.append(fam_header.lstrip("\n"))
        for cb in ctrl_blocks:
            rebuilt_sec2.append(cb)
    return header_part + "".join(rebuilt_sec2)


def inject_ssp_control_parameters(
    ssp_markdown: str,
    regime: Regime,
    baseline: Mapping[str, Any],
) -> str:
    """Inject regime parameter and guidance rows into each SSP control summary table."""
    sec2_marker = "# 2. Minimum Security Controls"
    idx = ssp_markdown.find(sec2_marker)
    if idx == -1:
        return ssp_markdown

    header_part = ssp_markdown[: idx + len(sec2_marker)]
    sec2_body = ssp_markdown[idx + len(sec2_marker) :]

    formatted_rows: Mapping[str, str] = baseline.get("formatted_parameter_rows", {})
    guidance_by_ctrl: Mapping[str, str] = baseline.get("guidance_by_control", {})
    include_fedramp_guidance = regime in (
        Regime.FEDRAMP_LOW,
        Regime.FEDRAMP_MODERATE,
        Regime.FEDRAMP_HIGH,
    )

    ctrl_heading_re = re.compile(r"(?m)^###\s+([A-Z]{2}-\d+(?:\(\d+\))?)(?=\s|$)[^\n]*$")
    parts: List[str] = []
    matches = list(ctrl_heading_re.finditer(sec2_body))
    if not matches:
        return ssp_markdown

    parts.append(sec2_body[: matches[0].start()])
    for m_idx, m in enumerate(matches):
        start = m.start()
        end = matches[m_idx + 1].start() if m_idx + 1 < len(matches) else len(sec2_body)
        chunk = sec2_body[start:end]
        norm_cid = normalize_control_id(m.group(1))

        # Split off any trailing ## family heading that belongs to the next family
        next_fam_m = re.search(r"(?m)^##\s+", chunk)
        if next_fam_m:
            ctrl_chunk = chunk[: next_fam_m.start()]
            trailing_fam = chunk[next_fam_m.start() :]
        else:
            ctrl_chunk = chunk
            trailing_fam = ""

        table_marker = "| Control Summary Information |"
        if table_marker in ctrl_chunk:
            stmt_part, tbl_marker, summary_part = ctrl_chunk.partition(table_marker)
            if regime != Regime.FEDRAMP_HIGH:
                stmt_part = _strip_inline_fedramp_high_values(stmt_part)

            filtered_lines: List[str] = []
            for line in summary_part.splitlines():
                stripped = line.strip()
                if stripped.startswith("| **") and (
                    "Parameter Values**:" in stripped
                    or stripped.startswith("| **FedRAMP Additional Guidance**:")
                ):
                    continue
                filtered_lines.append(line)

            param_val = formatted_rows.get(norm_cid)
            guide_val = guidance_by_ctrl.get(norm_cid) if include_fedramp_guidance else None
            new_row: Optional[str] = None
            if param_val:
                if guide_val:
                    new_row = (
                        f"| **{regime.value} Parameter Values**: {param_val}"
                        f"<br>**FedRAMP Additional Guidance**: {guide_val} |"
                    )
                else:
                    new_row = f"| **{regime.value} Parameter Values**: {param_val} |"
            elif guide_val:
                new_row = f"| **FedRAMP Additional Guidance**: {guide_val} |"

            if new_row is not None:
                # Insert immediately after the last table row in filtered_lines
                insert_at = len(filtered_lines)
                for idx_line in range(len(filtered_lines) - 1, -1, -1):
                    if filtered_lines[idx_line].strip().startswith("|"):
                        insert_at = idx_line + 1
                        break
                filtered_lines.insert(insert_at, new_row)

            ctrl_chunk = stmt_part + tbl_marker + "\n".join(filtered_lines) + "\n\n\n"

        parts.append(ctrl_chunk + trailing_fam)

    return header_part + "".join(parts)


def _apply_regime_ssp_label_substitutions(
    injected: str,
    regime: Regime,
    spec: RegimeSpec,
    odp_tokens: Mapping[str, str],
) -> str:
    """Applies post-token regime label substitutions in SSP prose for DoD IL2/IL4/IL6 and civilian regimes."""
    if regime in (Regime.DOD_IL2, Regime.DOD_IL4):
        il_num = "2" if regime == Regime.DOD_IL2 else "4"
        injected = injected.replace(
            "DoD Impact Level 5 Provisional Authorization (PA)",
            f"DoD Impact Level {il_num} Provisional Authorization (PA)",
        )
        injected = injected.replace("DoD IL5 Provisional Authorization", f"DoD IL{il_num} Provisional Authorization")
        injected = injected.replace("DoD IL5 PA", f"DoD IL{il_num} PA")
        injected = injected.replace("DoD IL5", f"DoD IL{il_num}")
        injected = injected.replace("Impact Level 5 (IL5", f"Impact Level {il_num} (IL{il_num}")
        injected = injected.replace("Impact Level 5", f"Impact Level {il_num}")
        injected = injected.replace("SRG (IL5)", f"SRG (IL{il_num})")
        injected = injected.replace("Assured Workloads IL5", f"Assured Workloads IL{il_num}")
        injected = injected.replace("primary IL5 boundary sensor", f"primary IL{il_num} boundary sensor")
    elif regime == Regime.DOD_IL6:
        injected = injected.replace("DoD IL5 Provisional Authorization", "DoD IL6 Provisional Authorization")
        injected = injected.replace("DoD IL5 PA", spec.csp_inheritance_header)
        injected = injected.replace("DoD IL5", "DoD IL6")
        injected = injected.replace("Impact Level 5 (IL5", "Impact Level 6 (IL6")
        injected = injected.replace("Impact Level 5", "Impact Level 6")
        injected = injected.replace("SRG (IL5)", "SRG (IL6)")
        injected = injected.replace("Assured Workloads IL5", "[CONFIG_REQUIRED: Classified Cloud / GDC Secret Enclave]")
        injected = injected.replace("Assured Workloads IL6", "[CONFIG_REQUIRED: Classified Cloud / GDC Secret Enclave]")
        injected = injected.replace("primary IL5 boundary sensor", "primary IL6 boundary sensor")
    elif regime in (
        Regime.FEDRAMP_LOW,
        Regime.FEDRAMP_MODERATE,
        Regime.FISMA_LOW,
        Regime.FISMA_MODERATE,
        Regime.FISMA_HIGH,
        Regime.CJIS,
    ):
        csp_auth = odp_tokens["{{ CSP_INHERITANCE_AUTHORIZATION }}"]
        injected = injected.replace("Inherited from the Google Services FedRAMP High P-ATO", f"Inherited from the {csp_auth}")
        injected = injected.replace("Leverages the Google Services FedRAMP High P-ATO", f"Leverages the {csp_auth}")
        injected = injected.replace(
            "Google Cloud FedRAMP High customer responsibility matrix",
            f"Google Cloud {regime.value} customer responsibility matrix",
        )
        injected = injected.replace("within the accredited FedRAMP High boundary", f"within the accredited {regime.value} boundary")
        injected = injected.replace("where enabled in FedRAMP High", f"where enabled in {regime.value}")
        injected = injected.replace(
            "Security Command Center in FedRAMP High / commercial boundaries",
            f"Security Command Center in {regime.value} / commercial boundaries",
        )
    return injected


def render_ssp_for_regime(
    raw_ssp_markdown: str,
    regime: Regime,
    baseline: Optional[Mapping[str, Any]] = None,
) -> str:
    """Filter SSP control sections and inject baseline parameter values for ``regime``."""
    spec = get_regime_spec(regime)
    active_baseline = baseline if baseline is not None else get_regime_baseline(regime)
    filtered = filter_ssp_control_sections(raw_ssp_markdown, active_baseline["control_ids"])
    injected = inject_ssp_control_parameters(filtered, regime, active_baseline)

    if regime == Regime.DOD_IL6:
        il6_hdr = spec.csp_inheritance_header
        injected = injected.replace(
            "Inherited controls cite the Google Services {{ BASE_FEDRAMP_LABEL }} P-ATO "
            "(FedRAMP package {{ CSP_PATO_PACKAGE_ID }}) and the {{ DOD_PA_FULL_LABEL }} "
            "that DISA issued for Google Cloud.",
            f"Inherited controls cite {il6_hdr} (Package ID: {{{{ CSP_PATO_PACKAGE_ID }}}}).",
        )
        injected = injected.replace(
            "Enforced at the Google Cloud Assured Workloads {{ DOD_IMPACT_SHORT }} folder boundary "
            "(`{{ PRIMARY_LOCATION }}`) and inherited from the Google Services {{ DOD_PA_FULL_LABEL }} "
            "(`{{ CSP_PATO_PACKAGE_ID }}`).",
            "[CONFIG_REQUIRED: Document Classified Cloud / GDC Secret enclave data residency and personnel "
            f"security controls (`{{{{ PRIMARY_LOCATION }}}}`)] and {il6_hdr} (`{{{{ CSP_PATO_PACKAGE_ID }}}}`).",
        )
        for prefix in ("Inherited from", "Leverages", "inherited from"):
            injected = injected.replace(
                f"{prefix} the Google Services {{{{ BASE_FEDRAMP_LABEL }}}} P-ATO "
                "(Package ID: {{ CSP_PATO_PACKAGE_ID }}) and {{ DOD_PA_LABEL }}",
                f"{prefix} {il6_hdr} (Package ID: {{{{ CSP_PATO_PACKAGE_ID }}}})",
            )
    elif regime == Regime.CJIS:
        cjis_clause = (
            "Inherited from Google Cloud Assured Workloads for CJIS "
            "(FBI CJIS Security Policy v5.9.4 / v6.0 Data Boundary — {{ CSP_PATO_PACKAGE_ID }})"
        )
        injected = injected.replace(
            "Leverages the {{ CSP_INHERITANCE_AUTHORIZATION }} (Package ID: {{ CSP_PATO_PACKAGE_ID }}).",
            spec.csp_inheritance_header,
        )
        injected = injected.replace(
            "Inherited from the {{ CSP_INHERITANCE_AUTHORIZATION }} (Package ID: {{ CSP_PATO_PACKAGE_ID }})",
            cjis_clause,
        )
        injected = injected.replace(
            "inherited from the Google Services FedRAMP High P-ATO (Package ID: {{ CSP_PATO_PACKAGE_ID }})",
            "inherited from Google Cloud Assured Workloads for CJIS "
            "(FBI CJIS Security Policy v5.9.4 / v6.0 Data Boundary — {{ CSP_PATO_PACKAGE_ID }})",
        )

    odp_tokens = build_regime_odp_tokens(regime)
    for token_key in (
        "{{ REGIME_DISPLAY_NAME }}",
        "{{ CSP_INHERITANCE_AUTHORIZATION }}",
        "{{ CSP_INHERITANCE_HEADER }}",
        "{{ DOD_IMPACT_SHORT }}",
        "{{ DOD_IMPACT_LABEL }}",
        "{{ DOD_PA_LABEL }}",
        "{{ DOD_PA_FULL_LABEL }}",
        "{{ BASE_FEDRAMP_LABEL }}",
        "{{ DOD_BASELINE_FRAMEWORK_LABEL }}",
    ):
        if token_key in injected:
            injected = injected.replace(token_key, odp_tokens[token_key])

    return _apply_regime_ssp_label_substitutions(injected, regime, spec, odp_tokens)


def filter_policy_appendix_rows(
    raw_policy_markdown: str,
    allowed_control_ids: Collection[str],
) -> str:
    """Filter ``## Appendix A`` Markdown table rows to ``allowed_control_ids``."""
    allowed_set: Set[str] = {normalize_control_id(c) for c in allowed_control_ids}
    lines = raw_policy_markdown.splitlines(keepends=True)
    out: List[str] = []
    in_appendix_a = False

    for line in lines:
        stripped = line.lstrip()
        if stripped.startswith("## "):
            in_appendix_a = stripped.startswith("## Appendix A")
            out.append(line)
            continue
        if in_appendix_a:
            m = _POLICY_APPENDIX_ROW_RE.match(line)
            if m:
                fam, num, enh = m.groups()
                cid = f"{fam.upper()}-{int(num)}" + (f"({int(enh)})" if enh else "")
                if cid not in allowed_set:
                    continue
        out.append(line)
    return "".join(out)


def filter_sctm_yaml_for_regime(
    populated_yaml_text: str,
    regime: Regime,
    baseline: Optional[Mapping[str, Any]] = None,
    open_poam_controls: Optional[Collection[str]] = None,
) -> str:
    """Filter ``SCTM_Burndown_Matrix.yaml`` controls and update summary counts for ``regime``."""
    spec = get_regime_spec(regime)
    active_baseline = baseline if baseline is not None else get_regime_baseline(regime)
    allowed_set: Set[str] = {normalize_control_id(c) for c in active_baseline["control_ids"]}
    open_base_set: Set[str] = {base_control_id(c) for c in (open_poam_controls or ())}
    try:
        parsed = yaml.safe_load(populated_yaml_text)
    except yaml.YAMLError:
        return populated_yaml_text
    if not isinstance(parsed, dict):
        return populated_yaml_text

    sys_meta = parsed.get("system_metadata")
    if isinstance(sys_meta, dict):
        sys_meta["compliance_baseline"] = regime.value

    families = parsed.get("control_families")
    kept_families: List[Dict[str, Any]] = []
    total_rep = 0
    total_imp = 0
    total_inh = 0
    total_part = 0
    total_plan = 0
    if isinstance(families, list):
        for fam in families:
            if not isinstance(fam, dict):
                continue
            ctrls = fam.get("controls")
            if not isinstance(ctrls, list):
                continue
            kept_ctrls: List[Dict[str, Any]] = []
            for c in ctrls:
                if not isinstance(c, dict):
                    continue
                cid_norm = normalize_control_id(str(c.get("id", "")))
                if cid_norm not in allowed_set:
                    continue
                c_copy = dict(c)
                if base_control_id(cid_norm) in open_base_set and str(c_copy.get("status", "")).strip().lower() == "implemented":
                    c_copy["status"] = "Planned"
                if regime == Regime.DOD_IL6:
                    for field in ("implementation_details", "codebase_evidence"):
                        if isinstance(c_copy.get(field), str):
                            c_copy[field] = (
                                c_copy[field]
                                .replace(
                                    "FedRAMP High / DoD IL5 accredited facilities",
                                    "[CONFIG_REQUIRED: Classified Cloud / GDC Secret Enclave accredited facilities]",
                                )
                                .replace(
                                    "Google Cloud FedRAMP High / DoD IL5 Authorization Package",
                                    spec.csp_inheritance_header,
                                )
                                .replace(
                                    "Google Cloud FedRAMP High / DoD IL5 Authorization",
                                    spec.csp_inheritance_header,
                                )
                            )
                kept_ctrls.append(c_copy)
            if not kept_ctrls:
                continue
            fam_copy = dict(fam)
            fam_copy["controls"] = kept_ctrls
            fam_copy["total_controls"] = len(kept_ctrls)
            fam_imp = sum(1 for c in kept_ctrls if str(c.get("status", "")).lower() == "implemented")
            fam_inh = sum(1 for c in kept_ctrls if str(c.get("status", "")).lower() == "inherited")
            fam_part = sum(1 for c in kept_ctrls if str(c.get("status", "")).lower() in ("partially implemented", "partial"))
            fam_plan = sum(
                1 for c in kept_ctrls if str(c.get("status", "")).lower() not in ("implemented", "inherited", "partially implemented", "partial")
            )
            fam_copy["implemented_iac"] = fam_imp
            fam_copy["inherited_gcp"] = fam_inh
            fam_copy["planned_rmf_manual"] = fam_part + fam_plan
            kept_families.append(fam_copy)
            total_rep += len(kept_ctrls)
            total_imp += fam_imp
            total_inh += fam_inh
            total_part += fam_part
            total_plan += fam_plan
        parsed["control_families"] = kept_families

    summary = parsed.get("overall_traceability_summary")
    if isinstance(summary, dict):
        summary["total_baseline_controls"] = len(allowed_set)
        summary["representative_summary_controls"] = total_rep
        summary["implemented"] = total_imp
        summary["inherited"] = total_inh
        summary["partially_implemented"] = total_part
        summary["planned"] = total_plan
        summary["completion_percentage"] = round(((total_imp + total_inh) / total_rep) * 100.0, 1) if total_rep else 0.0

    return yaml.safe_dump(parsed, sort_keys=False, allow_unicode=True, width=120)


def _audit_policy_appendix_baseline_parity(
    ato_path: Path, regime: Regime, expected_ids: Set[str]
) -> List[Dict[str, Any]]:
    """Checks Policy Appendix A rows against the active baseline."""
    policies_dir = ato_path / "Policies_and_Procedures"
    if not policies_dir.is_dir():
        return []
    out_of_baseline_pol: List[str] = []
    for p_file in sorted(policies_dir.glob("*.md")):
        in_app = False
        for line in p_file.read_text(encoding="utf-8").splitlines():
            if line.lstrip().startswith("## "):
                in_app = line.lstrip().startswith("## Appendix A")
                continue
            if in_app and (m := _POLICY_APPENDIX_ROW_RE.match(line)):
                fam, num, enh = m.groups()
                cid = f"{fam.upper()}-{int(num)}" + (f"({int(enh)})" if enh else "")
                if cid not in expected_ids:
                    out_of_baseline_pol.append(f"{p_file.name}:{cid}")
    if not out_of_baseline_pol:
        return []
    return [{
        "id": "CAT1-BASELINE-001",
        "component": "Policy Appendix A Baseline Parity",
        "severity": "CAT I (Critical)",
        "description": (
            f"Policy Appendix A tables contain {len(out_of_baseline_pol)} control row(s) outside the "
            f"{regime.value} baseline: {out_of_baseline_pol[:10]}."
        ),
        "remediation": f"Filter Policy Appendix A rows to the {regime.value} baseline.",
    }]


def _audit_sctm_xlsm_baseline_parity(
    ato_path: Path, regime: Regime, expected_ids: Set[str]
) -> List[Dict[str, Any]]:
    """Checks SCTM XLSM out-of-baseline rows for Not Applicable status and blank completion date."""
    sctm_xlsm = ato_path / "SCTM" / "SCTM_Security_Control_Traceability_Matrix.xlsm"
    if not sctm_xlsm.is_file():
        return []
    findings: List[Dict[str, Any]] = []
    try:
        import openpyxl  # pylint: disable=import-outside-toplevel

        wb = openpyxl.load_workbook(sctm_xlsm, read_only=True, data_only=True)
        if "Template" in wb.sheetnames:
            ws = wb["Template"]
            bad_rows: List[str] = []
            for row in ws.iter_rows(min_row=7, max_col=10, values_only=True):
                raw_cid = row[0]
                if not raw_cid or not str(raw_cid).strip():
                    continue
                norm_cid = normalize_control_id(str(raw_cid))
                if norm_cid not in expected_ids:
                    status_val = str(row[4] or "").strip()
                    na_just = str(row[8] or "").strip()
                    est_date = str(row[9] or "").strip()
                    if status_val != "Not Applicable" or not na_just or est_date:
                        bad_rows.append(norm_cid)
            if bad_rows:
                findings.append({
                    "id": "CAT1-BASELINE-001",
                    "component": "SCTM XLSM Baseline Parity",
                    "severity": "CAT I (Critical)",
                    "description": (
                        f"SCTM workbook has {len(bad_rows)} out-of-baseline control(s) without valid "
                        f"'Not Applicable' status, justification, or blank completion date: {bad_rows[:10]}."
                    ),
                    "remediation": f"Mark out-of-baseline controls in SCTM as Not Applicable for {regime.value}.",
                })
        wb.close()
    except (OSError, ValueError, KeyError):
        pass
    return findings


def audit_baseline_control_parity(
    ato_dir: Union[str, Path],
    inventory: Mapping[str, Any],
) -> List[Dict[str, Any]]:
    """Audit SSP control headings, Policy Appendix A rows, and SCTM XLSM N/A rows against the active baseline."""
    sys_info = inventory.get("system_information")
    if not isinstance(sys_info, Mapping):
        return []
    regime = classify_regime(sys_info)
    if not is_supported_regime(regime):
        return []

    baseline = get_regime_baseline(regime)
    expected_ids: Set[str] = {normalize_control_id(c) for c in baseline["control_ids"]}
    ato_path = Path(ato_dir)
    findings: List[Dict[str, Any]] = []

    ssp_candidates = (
        ato_path / "SSP" / "SSP_System_Security_Plan.md",
        ato_path / "SSP" / "System_Security_Plan.md",
    )
    ssp_file = next((p for p in ssp_candidates if p.is_file()), None)
    if ssp_file is not None:
        ssp_text = ssp_file.read_text(encoding="utf-8")
        ctrl_re = re.compile(r"(?m)^###\s+([A-Z]{2}-\d+(?:\(\d+\))?)(?=\s|$)")
        actual_ssp_ids = {normalize_control_id(m.group(1)) for m in ctrl_re.finditer(ssp_text)}
        extra_ssp = sorted(actual_ssp_ids - expected_ids)
        missing_ssp = sorted(expected_ids - actual_ssp_ids)
        if extra_ssp or (len(actual_ssp_ids) > 20 and missing_ssp):
            findings.append({
                "id": "CAT1-BASELINE-001",
                "component": "SSP Baseline Control Parity",
                "severity": "CAT I (Critical)",
                "description": (
                    f"{ssp_file.name} control set ({len(actual_ssp_ids)} controls) does not match the "
                    f"{regime.value} baseline ({len(expected_ids)} controls): "
                    f"out-of-baseline={extra_ssp[:10]}, missing={missing_ssp[:10]}."
                ),
                "remediation": (
                    f"Regenerate SSP_System_Security_Plan.md for {regime.value} so control headings match "
                    f"config/baselines/{get_regime_spec(regime).baseline_filename} 1:1."
                ),
            })

    findings.extend(_audit_policy_appendix_baseline_parity(ato_path, regime, expected_ids))
    findings.extend(_audit_sctm_xlsm_baseline_parity(ato_path, regime, expected_ids))
    return findings
