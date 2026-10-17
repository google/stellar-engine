# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

import os
import shutil
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch, MagicMock

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "scripts")))

import validate_compliance_artifacts
import oscal_generator


class _StopValidation(Exception):
    """Sentinel used to abort a run at a specific seam.

    A dedicated type keeps the test from accidentally passing on a genuine
    failure raised by the code under test.
    """


class TestHardeningValidate(unittest.TestCase):
    def test_docx_zip_bomb(self):
        with tempfile.TemporaryDirectory() as td:
            tf_path = os.path.join(td, "test.docx")
            with zipfile.ZipFile(tf_path, "w") as z:
                z.writestr("word/document.xml", "test")

            with patch("zipfile.ZipFile.getinfo") as mock_getinfo:
                mock_info = MagicMock()
                mock_info.file_size = 101 * 1024 * 1024  # > 100MB
                mock_info.compress_size = 100
                mock_getinfo.return_value = mock_info

                results = validate_compliance_artifacts.audit_docx_policies(td)
                self.assertEqual(len(results), 1)
                self.assertEqual(results[0]["status"], "UNVERIFIED")
                self.assertIn("exceeds safe size/compression", results[0]["issues"][0])

    def test_oscal_deterministic_uuid_fips(self):
        # UUID should be formatted properly as a string
        uuid_val = oscal_generator._deterministic_uuid("test-string")
        self.assertTrue(isinstance(uuid_val, str))
        self.assertEqual(len(uuid_val), 36)

        import uuid
        parsed = uuid.UUID(uuid_val)
        self.assertEqual(parsed.version, 4)

    def test_ssp_text_redos_fix(self):
        # Ensure audit_senior_compliance_quality reads SSP/System_Security_Plan.md fallback without ReDoS
        target_dir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, target_dir, True)
        ato_dir = os.path.join(target_dir, "ato_artifacts")
        ssp_dir = os.path.join(ato_dir, "SSP")
        os.makedirs(ssp_dir)
        ssp_md = os.path.join(ssp_dir, "System_Security_Plan.md")

        content = (
            "### SC-7 Boundary Protection\n"
            "Default-deny perimeter firewall rules and VPC network isolation enforce boundary defense.\n"
            "| **Implementation Status**:<br>- [x] Implemented<br>- [ ] Planned |\n"
            + ("x" * 10000)
            + "\n### SC-8\n"
        )
        with open(ssp_md, "w", encoding="utf-8") as f:
            f.write(content)

        res = validate_compliance_artifacts.audit_senior_compliance_quality(
            target_dir=target_dir,
            inventory={},
            ato_dir=ato_dir,
            alignment_res={"discrepancies": []},
            excel_results=[],
            docx_results=[],
            oscal_results=[],
            unresolved_tokens=[],
            config_required_vars=[],
            stigs_required=[],
            rmf_action_items=[]
        )
        sc7_rec = next(r for r in res["atc_records"] if r["id"] == "SC-7")
        self.assertEqual(sc7_rec["evidence_source"], "SSP")
        self.assertEqual(sc7_rec["verification"], "VERIFIED")

    def test_subprocess_command_injection_prevention(self):
        """A target directory that looks like a CLI flag must not be parsed as one."""
        # The target path is relative, so the validator materializes directories
        # beneath the current working directory. Run inside a throwaway cwd so the
        # test cannot leave a literal '-evil-flag' tree in the repository.
        original_cwd = os.getcwd()
        with tempfile.TemporaryDirectory(prefix="compliance-injection-test-") as sandbox_cwd:
            try:
                os.chdir(sandbox_cwd)
                with patch("validate_compliance_artifacts.run_bounded_command") as mock_run, \
                        patch("os.path.exists", return_value=True), \
                        patch("shutil.copytree"), \
                        patch("validate_compliance_artifacts.read_json_file", return_value={"dummy": "data"}):
                    # Abort the run at the first subprocess boundary; we only care
                    # about the argv that was about to be executed.
                    mock_run.side_effect = _StopValidation("halt after first subprocess call")
                    with self.assertRaises(_StopValidation):
                        validate_compliance_artifacts.validate_compliance_package("-evil-flag", fix_drift=True)
            finally:
                os.chdir(original_cwd)

        self.assertTrue(
            mock_run.called,
            "validate_compliance_package never reached the command runner; the argv assertion below "
            "would silently vacuous-pass, so fail loudly instead.",
        )
        call_args = mock_run.call_args[0][0]
        self.assertIn("--", call_args, f"missing end-of-options separator in {call_args!r}")
        self.assertLess(
            call_args.index("--"),
            len(call_args) - 1,
            f"'--' must precede the operand, not trail it: {call_args!r}",
        )
        self.assertTrue(call_args[-1].endswith("-evil-flag"))

    def test_injection_test_leaves_no_residue(self):
        """The injection test must not create a '-evil-flag' directory in the repo."""
        skill_root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
        for probe_root in (skill_root, os.getcwd()):
            residue = os.path.join(probe_root, "-evil-flag")
            self.assertFalse(
                os.path.exists(residue),
                f"test residue directory was left behind at {residue}",
            )

    def _make_minimal_package(self, extra_inventory=None, create_ssp=True):
        td = tempfile.mkdtemp(prefix="val_contract_")
        self.addCleanup(shutil.rmtree, td, True)
        ato_dir = os.path.join(td, "ato_artifacts")
        os.makedirs(os.path.join(ato_dir, "SSP"), exist_ok=True)
        os.makedirs(os.path.join(ato_dir, "POAM"), exist_ok=True)
        if create_ssp:
            with open(os.path.join(ato_dir, "SSP", "SSP_System_Security_Plan.md"), "w", encoding="utf-8") as f:
                f.write("# System Security Plan\n\n### AC-17 Remote Access\nCloud IAP tunnels over TLS 1.3.\n")
        with open(os.path.join(ato_dir, "POAM", "Plan_of_Action_and_Milestones.yaml"), "w", encoding="utf-8") as f:
            f.write("poam_items: []\n")
        inv = {
            "system_information": {
                "system_name": "ContractTestSys",
                "system_abbreviation": "CTS",
                "organization": "Test Agency",
                "impact_level": "IL5",
                "compliance_baseline": "DoD IL5",
            },
            "personnel_roles": {
                "authorizing_official": {"name": "AO"},
                "system_owner": {"name": "SO"},
                "issm": {"name": "ISSM"},
                "isso": {"name": "ISSO"},
            },
            "infrastructure_components": {},
            "network_architecture": {},
        }
        if extra_inventory:
            inv.update(extra_inventory)
        import json
        with open(os.path.join(td, "system_inventory.json"), "w", encoding="utf-8") as f:
            json.dump(inv, f)
        return td

    def test_exit_nonzero_cat1_assessor(self):
        td = self._make_minimal_package(
            extra_inventory={
                "network_architecture": {
                    "firewall_rules": [
                        {
                            "name": "open-ssh",
                            "direction": "INGRESS",
                            "action": "ALLOW",
                            "protocol": "tcp",
                            "ports": ["22"],
                            "source_ranges": ["0.0.0.0/0"],
                        }
                    ]
                }
            }
        )
        self.assertFalse(
            validate_compliance_artifacts.validate_compliance_package(
                td, run_linter=False, linter_strict=True
            )
        )
        # --no-linter-strict must NOT suppress assessor CAT I findings
        self.assertFalse(
            validate_compliance_artifacts.validate_compliance_package(
                td, run_linter=False, linter_strict=False
            )
        )

    def test_exit_nonzero_cat1_semantic_strict(self):
        # Missing SSP triggers LINT-SSP-MISSING (CAT I semantic finding), no assessor CAT I
        td = self._make_minimal_package(create_ssp=False)
        self.assertFalse(
            validate_compliance_artifacts.validate_compliance_package(
                td, run_linter=True, linter_strict=True
            )
        )

    def test_no_linter_strict_scope(self):
        # 1. Semantic CAT I only + --no-linter-strict -> returns True (exit 0)
        td_sem = self._make_minimal_package(create_ssp=False)
        self.assertTrue(
            validate_compliance_artifacts.validate_compliance_package(
                td_sem, run_linter=True, linter_strict=False
            )
        )
        # 2. Unresolved macro token + --no-linter-strict -> returns False (exit 1)
        td_tok = self._make_minimal_package()
        with open(os.path.join(td_tok, "ato_artifacts", "SSP", "SSP_System_Security_Plan.md"), "a", encoding="utf-8") as f:
            f.write("\nUnresolved {{ BAD_TOKEN_HERE }} in SSP.\n")
        self.assertFalse(
            validate_compliance_artifacts.validate_compliance_package(
                td_tok, run_linter=False, linter_strict=False
            )
        )
        # 3. Corrupt YAML deliverable + --no-linter-strict -> returns False (exit 1)
        td_pkg = self._make_minimal_package()
        with open(os.path.join(td_pkg, "ato_artifacts", "POAM", "Plan_of_Action_and_Milestones.yaml"), "w", encoding="utf-8") as f:
            f.write("poam_items: [unclosed\n  : :\n")
        self.assertFalse(
            validate_compliance_artifacts.validate_compliance_package(
                td_pkg, run_linter=False, linter_strict=False
            )
        )

    def test_corrupt_inventory_fix_exits_1(self):
        td = self._make_minimal_package()
        with open(os.path.join(td, "system_inventory.json"), "w", encoding="utf-8") as f:
            f.write("{corrupt json")
        self.assertFalse(
            validate_compliance_artifacts.validate_compliance_package(
                td, fix_drift=True, run_linter=False, linter_strict=False
            )
        )
        self.assertFalse(
            validate_compliance_artifacts.validate_compliance_package(
                td, fix_drift=False, run_linter=False, linter_strict=False
            )
        )

    def test_linter_crash_strict_fails(self):
        td = self._make_minimal_package()
        with patch(
            "validate_compliance_artifacts.run_semantic_linter",
            side_effect=RuntimeError("simulated linter crash"),
        ):
            self.assertFalse(
                validate_compliance_artifacts.validate_compliance_package(
                    td, run_linter=True, linter_strict=True
                )
            )
            self.assertTrue(
                validate_compliance_artifacts.validate_compliance_package(
                    td, run_linter=True, linter_strict=False
                )
            )

    def test_cli_rejects_unknown_flags(self):
        td = self._make_minimal_package()
        for bad_flag in ("--bogus-nonexistent-flag", "--update", "--fill-examples"):
            with self.subTest(flag=bad_flag):
                with self.assertRaises(SystemExit) as ctx:
                    validate_compliance_artifacts.main([bad_flag, td])
                self.assertEqual(ctx.exception.code, 2)

    def test_space_separated_values(self):
        td = self._make_minimal_package()
        with patch("validate_compliance_artifacts.validate_compliance_package", return_value=True) as mock_val:
            validate_compliance_artifacts.main(
                ["--fix", "--no-fill-example-data", "--policy-format", "markdown", td]
            )
        self.assertTrue(mock_val.called)
        self.assertEqual(mock_val.call_args.kwargs.get("policy_format"), "markdown")
        self.assertTrue(mock_val.call_args.kwargs.get("fix_drift"))

    def test_format_flags_require_fix(self):
        td = self._make_minimal_package()
        for flag, val in (
            ("--policy-format", "markdown"),
            ("--data-format", "yaml"),
            ("--oscal-format", "json"),
            ("--oscal-version", "1.2.3"),
        ):
            with self.subTest(flag=flag):
                with self.assertRaises(SystemExit) as ctx:
                    validate_compliance_artifacts.main([flag, val, td])
                self.assertEqual(ctx.exception.code, 2)

    def test_help_lists_stig_flags(self):
        import io
        from contextlib import redirect_stdout
        buf = io.StringIO()
        with redirect_stdout(buf):
            with self.assertRaises(SystemExit) as ctx:
                validate_compliance_artifacts.main(["--help"])
        self.assertEqual(ctx.exception.code, 0)
        help_text = buf.getvalue()
        self.assertIn("--stigs-mode", help_text)
        self.assertIn("--stigs-catalog", help_text)
        self.assertNotIn("--fill-examples", help_text)

    def test_reconciliation_includes_serverless(self):
        td = self._make_minimal_package()
        ato_dir = os.path.join(td, "ato_artifacts")
        with open(os.path.join(ato_dir, "SSP", "SSP_System_Security_Plan.md"), "w", encoding="utf-8") as f:
            f.write(
                "# SSP\nCloud Run: svc-api\nFunction: fn-worker\nRegistry: ar-main\n"
                "Sink: sink-audit\nTopic: topic-events\nNAT: nat-egress\nPerimeter: perim-il5\n"
            )
        inv = {
            "infrastructure_components": {
                "cloud_run_services": [{"name": "svc-api"}],
                "cloud_functions": [{"name": "fn-worker"}],
                "artifact_registries": [{"name": "ar-main"}],
                "logging_sinks": [{"name": "sink-audit"}],
                "pubsub_topics": [{"name": "topic-events"}],
                "nat_gateways": [{"name": "nat-egress"}],
                "service_perimeters": [{"name": "perim-il5"}],
            }
        }
        res = validate_compliance_artifacts.audit_inventory_artifact_alignment(td, inv, ato_dir)
        self.assertEqual(res["total_discovered_assets"], 7, res["breakdown"])
        self.assertEqual(res["reconciled_assets_count"], 7, res["breakdown"])
        self.assertEqual(res["discrepancies"], [])

    def test_sa_dict_keys(self):
        td = self._make_minimal_package()
        ato_dir = os.path.join(td, "ato_artifacts")
        with open(os.path.join(ato_dir, "SSP", "SSP_System_Security_Plan.md"), "w", encoding="utf-8") as f:
            f.write("# SSP\nService account sa-bq@proj.iam.gserviceaccount.com\n")
        inv = {
            "infrastructure_components": {
                "service_accounts": [
                    {"name": "sa-bq", "email": "sa-bq@proj.iam.gserviceaccount.com"},
                    {"name": "", "email": ""},
                ]
            }
        }
        res = validate_compliance_artifacts.audit_inventory_artifact_alignment(td, inv, ato_dir)
        self.assertEqual(res["breakdown"]["service_accounts"]["discovered"], 2)
        self.assertEqual(res["breakdown"]["service_accounts"]["matched"], 1)
        self.assertEqual(len(res["discrepancies"]), 1)

    def test_generic_labels_not_evidence(self):
        td = self._make_minimal_package()
        ato_dir = os.path.join(td, "ato_artifacts")
        with open(os.path.join(ato_dir, "SSP", "SSP_System_Security_Plan.md"), "w", encoding="utf-8") as f:
            f.write("# SSP\nBy default, the main system enforces this policy.\n")
        inv = {
            "infrastructure_components": {
                "kms_keys": [{"name": "default"}],
                "storage_buckets": [{"name": "main"}],
                "compute_instances": [{"name": "this"}],
            }
        }
        res = validate_compliance_artifacts.audit_inventory_artifact_alignment(td, inv, ato_dir)
        self.assertEqual(res["total_discovered_assets"], 3)
        self.assertEqual(res["reconciled_assets_count"], 0)
        self.assertEqual(len(res["discrepancies"]), 3)

    def test_pta_readiness_matches_verdict(self):
        td = self._make_minimal_package(create_ssp=False)
        ok = validate_compliance_artifacts.validate_compliance_package(
            td, run_linter=True, linter_strict=True, policy_format="markdown"
        )
        self.assertFalse(ok)
        with open(os.path.join(td, "ato_artifacts", "Path_to_Authorization.md"), "r", encoding="utf-8") as f:
            pta = f.read()
        self.assertIn("Remediation Required Prior to Assessment Submission", pta)
        self.assertNotIn("Conditional Package Readiness - Ready for Interim Assessment", pta)

    def test_excel_row_counts(self):
        import openpyxl
        td = self._make_minimal_package()
        ato_dir = os.path.join(td, "ato_artifacts")
        os.makedirs(os.path.join(ato_dir, "POAM"), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "POA&M"
        ws.cell(row=7, column=1, value="Control")
        ws.cell(row=7, column=2, value="POA&M Item ID")
        # Populate 2 real rows (rows 8 and 9)
        ws.cell(row=8, column=1, value="AC-2")
        ws.cell(row=8, column=2, value="POAM-001")
        ws.cell(row=9, column=1, value="SC-7")
        ws.cell(row=9, column=2, value="POAM-002")
        # Leave formatted/empty cells out to row 200 so ws.max_row is 200
        ws.cell(row=200, column=10, value="")
        xl_path = os.path.join(ato_dir, "POAM", "Plan_of_Action_and_Milestones.xlsx")
        wb.save(xl_path)
        res = validate_compliance_artifacts.audit_excel_workbooks(ato_dir)
        self.assertEqual(len(res), 1)
        self.assertEqual(res[0]["rows_count"]["POA&M Milestones"], 2)

    def test_kms_drift_recognizes_cmek_key(self):
        from compliance_engine import semantic_linter
        td = self._make_minimal_package()
        ato_dir = os.path.join(td, "ato_artifacts")
        ssp_path = os.path.join(ato_dir, "SSP", "SSP_System_Security_Plan.md")
        with open(ssp_path, "w", encoding="utf-8") as f:
            f.write("# SSP\nData is protected with Customer-Managed Encryption Keys (CMEK) under SC-28.\n")
        inv_cmek = {
            "infrastructure_components": {
                "kms_keys": [],
                "databases": [{"name": "bq-ds", "cmek_key": "projects/p/locations/us/keyRings/kr/cryptoKeys/k"}],
            }
        }
        from pathlib import Path
        drift = semantic_linter.evaluate_architectural_drift(inv_cmek, ssp_path=Path(ssp_path))
        self.assertFalse(any(d.finding_id == "DRIFT-KMS-001" for d in drift), [d.to_dict() for d in drift])
        senior = validate_compliance_artifacts.audit_senior_compliance_quality(
            target_dir=td,
            inventory=inv_cmek,
            ato_dir=ato_dir,
            alignment_res={"discrepancies": []},
            excel_results=[],
            docx_results=[],
            oscal_results=[],
            unresolved_tokens=[],
            config_required_vars=[],
        )
        self.assertFalse(any(f["id"] == "CAT2-CRYPTO-001" for f in senior["cat_2_findings"]))

    def test_placeholder_scan_xlsm_all_columns(self):
        import openpyxl
        td = self._make_minimal_package()
        ato_dir = os.path.join(td, "ato_artifacts")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "POA&M"
        ws.cell(row=8, column=2, value="POAM-001")
        ws.cell(row=8, column=35, value="[CONFIG_REQUIRED: Late Column Value]")
        wb.save(os.path.join(ato_dir, "POAM", "Plan_of_Action_and_Milestones.xlsx"))
        validate_compliance_artifacts.validate_compliance_package(
            td, run_linter=False, linter_strict=False, policy_format="markdown"
        )
        with open(os.path.join(ato_dir, "Path_to_Authorization.md"), "r", encoding="utf-8") as f:
            pta = f.read()
        self.assertIn("[CONFIG_REQUIRED: Late Column Value]", pta)

    def test_placeholder_scan_runbook_tokens(self):
        td = self._make_minimal_package()
        ato_dir = os.path.join(td, "ato_artifacts")
        rb_dir = os.path.join(ato_dir, "Incident_Response_Runbooks")
        os.makedirs(rb_dir, exist_ok=True)
        with open(os.path.join(rb_dir, "IR_Compute_Resource_Compromise_Runbook.md"), "w", encoding="utf-8") as f:
            f.write("# Runbook\ngcloud compute instances stop <INSTANCE_ID: fill in>\n")
        validate_compliance_artifacts.validate_compliance_package(
            td, run_linter=False, linter_strict=False, policy_format="markdown"
        )
        with open(os.path.join(ato_dir, "Path_to_Authorization.md"), "r", encoding="utf-8") as f:
            pta = f.read()
        self.assertIn("<INSTANCE_ID: fill in>", pta)

    def test_assignment_slots_counted(self):
        td = self._make_minimal_package()
        ato_dir = os.path.join(td, "ato_artifacts")
        with open(os.path.join(ato_dir, "SSP", "SSP_System_Security_Plan.md"), "w", encoding="utf-8") as f:
            f.write("# SSP\nReview accounts [Assignment: organization-defined frequency].\n")
        validate_compliance_artifacts.validate_compliance_package(
            td, run_linter=False, linter_strict=False, policy_format="markdown"
        )
        with open(os.path.join(ato_dir, "Path_to_Authorization.md"), "r", encoding="utf-8") as f:
            pta = f.read()
        self.assertIn("[Assignment: organization-defined frequency]", pta)

    def test_sheet_scan_failure_reported(self):
        bad_ws = MagicMock()
        bad_ws.max_row = 10
        bad_ws.max_column = 10
        bad_ws.cell.side_effect = ValueError("corrupted worksheet cell XML")
        item_res = {"file": "test.xlsx", "status": "PASS", "issues": []}
        validate_compliance_artifacts._scan_sheet_for_unresolved_tokens(bad_ws, "POA&M", item_res)
        self.assertEqual(item_res["status"], "FAIL")
        self.assertTrue(any("corrupted worksheet cell XML" in i for i in item_res["issues"]))

    def test_cross_artifact_parity_check(self):
        td = self._make_minimal_package()
        ato_dir = os.path.join(td, "ato_artifacts")
        os.makedirs(os.path.join(ato_dir, "SCTM"), exist_ok=True)
        with open(os.path.join(ato_dir, "POAM", "Plan_of_Action_and_Milestones.yaml"), "w", encoding="utf-8") as f:
            f.write(
                "system_metadata:\n  system_name: OtherSystemName\n"
                "poam_items:\n  - item_id: POAM-001\n    control_identifier: AC-02\n    status: Ongoing\n"
            )
        with open(os.path.join(ato_dir, "SCTM", "SCTM_Burndown_Matrix.yaml"), "w", encoding="utf-8") as f:
            f.write(
                "system_metadata:\n  system_name: ContractTestSys\n"
                "control_families:\n  - controls:\n      - id: AC-2\n        status: Implemented\n"
            )
        import json
        inv = json.load(open(os.path.join(td, "system_inventory.json"), encoding="utf-8"))
        parity = validate_compliance_artifacts.audit_cross_artifact_parity(ato_dir, inv)
        self.assertFalse(parity["system_name_consistent"])
        self.assertTrue(len(parity["status_conflicts"]) >= 1)
        self.assertIn("DoD IL5", validate_compliance_artifacts.REGIME_CONTROL_EXPECTATIONS)

    def test_b3a_characterization_before_and_after_decomp(self):
        from compliance_engine import poam_rules, semantic_linter
        # 1. evaluate_control_substance across all control rule branches
        samples = [
            ("AC-17", "Cloud IAP zero-trust tunnels with TLS 1.3 encryption.", True),
            ("IA-2", "Hardware token MFA via FIDO2 security keys.", True),
            ("IA-5", "Password length and PKI certificate lifecycle.", True),
            ("SC-7", "Default-deny VPC firewall ingress rules.", True),
            ("SC-8", "TLS encryption in transit across endpoints.", True),
            ("SC-28", "Cloud KMS CMEK encryption at rest.", True),
            ("AU-2", "Cloud Logging audit sink to external SIEM.", True),
            ("IR-4", "Incident response runbook and 1-hour CISA reporting.", True),
            ("IR-8", "Incident response plan annual review and roles.", True),
            ("IR-9", "Information spillage quarantine and sanitization.", True),
            ("RA-5", "Automated vulnerability scan and patch SLA.", True),
            ("AC-17", "Users login normally.", False),
        ]
        for cid, text, expected_ok in samples:
            ok, _, _ = semantic_linter.evaluate_control_substance(cid, text)
            self.assertEqual(ok, expected_ok, f"{cid}: {text}")

        # 2. derive_poam_findings + consolidate_poam_items characterization
        inv = {
            "system_information": {"system_abbreviation": "CHR", "effective_date": "2026-06-01"},
            "poam_items": [
                {"control": "AC-2", "weakness_name": "Item A", "check_id": "CKV_1", "severity": "Moderate"},
                {"control": "AC-2", "weakness_name": "Item B", "check_id": "CKV_1", "severity": "High"},
            ],
            "infrastructure_components": {
                "unparsed_terraform_files": [{"path": "bad.tf", "error": "syntax error"}],
                "databases": [{"name": "db1", "backup_enabled": False}],
            },
        }
        poam = poam_rules.derive_poam_findings(inv, target_dir=None)
        self.assertEqual(len(poam), 3)
        self.assertEqual(poam[0]["severity"], "High")
        self.assertIn("2 affected locations", poam[0]["title"])

    def test_b3a_functions_within_loc_and_cc_limits(self):
        import ast
        src_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "src", "compliance_engine"))
        branch_nodes = (
            ast.If, ast.For, ast.AsyncFor, ast.While, ast.ExceptHandler,
            ast.With, ast.AsyncWith, ast.Assert, ast.IfExp, ast.comprehension,
        )
        violations = []
        for mod_name in ("validate_compliance_artifacts.py", "semantic_linter.py", "poam_rules.py"):
            mod_path = os.path.join(src_dir, mod_name)
            tree = ast.parse(open(mod_path, encoding="utf-8").read())
            for node in ast.walk(tree):
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    loc = node.end_lineno - node.lineno + 1
                    cc = 1
                    for child in ast.walk(node):
                        if isinstance(child, branch_nodes):
                            cc += 1
                        elif isinstance(child, ast.BoolOp):
                            cc += len(child.values) - 1
                    if loc > 150 or cc > 30:
                        violations.append(f"{mod_name}:{node.name} (LOC={loc}, CC={cc})")
        self.assertEqual(violations, [], f"Oversized/complex functions remain: {violations}")

    def test_out_of_baseline_ssp_control_rejected(self):
        """Injecting an out-of-baseline control into a FedRAMP Low SSP triggers CAT1-BASELINE-001."""
        from compliance_engine import baseline_registry
        from compliance_engine.file_helpers import Regime

        with tempfile.TemporaryDirectory() as td:
            ato_dir = os.path.join(td, "ato_artifacts")
            ssp_dir = os.path.join(ato_dir, "SSP")
            os.makedirs(ssp_dir)
            baseline = baseline_registry.get_regime_baseline(Regime.FEDRAMP_LOW)
            raw_ssp = open(
                os.path.join(os.path.dirname(__file__), "..", "templates", "ssp", "SSP_FedRAMP_High_Template.md"),
                encoding="utf-8",
            ).read()
            rendered = baseline_registry.render_ssp_for_regime(raw_ssp, Regime.FEDRAMP_LOW, baseline)
            ssp_path = os.path.join(ssp_dir, "SSP_System_Security_Plan.md")
            with open(ssp_path, "w", encoding="utf-8") as fh:
                fh.write(rendered)
            inv = {
                "system_information": {
                    "impact_level": "FedRAMP Low",
                    "compliance_baseline": "FedRAMP Low",
                }
            }
            findings_clean = validate_compliance_artifacts.audit_baseline_control_parity(ato_dir, inv)
            self.assertEqual(findings_clean, [])

            with open(ssp_path, "a", encoding="utf-8") as fh:
                fh.write("\n### AC-4 Information Flow Enforcement\n\nOut of baseline.\n")
            findings_bad = validate_compliance_artifacts.audit_baseline_control_parity(ato_dir, inv)
            self.assertTrue(
                any(f.get("id") == "CAT1-BASELINE-001" for f in findings_bad),
                f"Expected CAT1-BASELINE-001 finding, got {findings_bad}",
            )

    def test_atc_controls_scoped_to_dod_baseline(self):
        """Civilian and DoD IL4 packages are not flagged for NSS-only AU-4(1) or CA-3(6) ATC controls."""
        with tempfile.TemporaryDirectory() as td:
            ato_dir = os.path.join(td, "ato_artifacts")
            ssp_dir = os.path.join(ato_dir, "SSP")
            os.makedirs(ssp_dir)
            with open(os.path.join(ssp_dir, "SSP_System_Security_Plan.md"), "w", encoding="utf-8") as fh:
                fh.write("# 1. System Identification\n# 2. Minimum Security Controls\n")
            for impact in ("FedRAMP Moderate", "DoD IL4"):
                inv = {"system_information": {"impact_level": impact, "compliance_baseline": impact}}
                res = validate_compliance_artifacts.audit_senior_compliance_quality(
                    target_dir=td,
                    inventory=inv,
                    ato_dir=ato_dir,
                    alignment_res={"discrepancies": []},
                    excel_results=[],
                    docx_results=[],
                    oscal_results=[],
                    unresolved_tokens=[],
                    config_required_vars=[],
                    stigs_required=[],
                    rmf_action_items=[],
                )
                checked_ids = {r["id"] for r in res.get("atc_records", [])}
                self.assertNotIn("AU-4(1)", checked_ids, impact)
                self.assertNotIn("CA-3(6)", checked_ids, impact)


if __name__ == "__main__":
    unittest.main()
