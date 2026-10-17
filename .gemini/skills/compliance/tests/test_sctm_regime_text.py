# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""SCTM workbook wording and providers follow the regime and the evidence."""

import os
import tempfile
import unittest

from compliance_engine import excel_hydrator
from compliance_engine import file_helpers


def _hydrate(impact: str, **sys_extra):
    import openpyxl  # pylint: disable=import-outside-toplevel

    inventory = {
        "system_information": dict({"system_name": "Demo", "system_abbreviation": "DMO", "impact_level": impact}, **sys_extra),
        "infrastructure_components": {},
        "network_architecture": {},
    }
    tmp = tempfile.mkdtemp()
    out = os.path.join(tmp, "sctm.xlsm")
    tpl = file_helpers.get_templates_dir() / "sctm" / "ControlInfoExport_Template.xlsm"
    excel_hydrator.SCTMHydrator(str(tpl)).hydrate(inventory, out)
    wb = openpyxl.load_workbook(out, keep_vba=True)
    ws = wb["Template"]
    rows = [[ws.cell(row=r, column=c).value for c in range(1, 30)] for r in range(7, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    wb.close()
    return rows


class SctmRegimeTextTest(unittest.TestCase):
    """Column AC, Common Control Provider and Implementation Status."""

    def test_il5_no_scc_recommendation(self) -> None:
        rows = _hydrate("IL5")
        self.assertTrue(rows)
        self.assertFalse([r[0] for r in rows if any("SCC" in str(v) or "Security Command Center" in str(v) for v in r)])

    def test_fedramp_provider_not_dod(self) -> None:
        rows = _hydrate("FedRAMP High")
        self.assertFalse([r[0] for r in rows if r[5] == "DoD"])

    def test_empty_inventory_rows_not_implemented(self) -> None:
        rows = _hydrate("IL5")
        self.assertFalse([r[0] for r in rows if r[4] == "Implemented"])


if __name__ == "__main__":
    unittest.main()
