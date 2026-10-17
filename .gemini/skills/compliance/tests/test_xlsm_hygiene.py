#!/usr/bin/env python3
# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Structural hygiene, sanitization, and integrity tests for .xlsm templates."""

import hashlib
import io
import json
import math
from pathlib import Path
import re
import tempfile
import unittest
import warnings
import xml.etree.ElementTree as ET
import zipfile

import openpyxl


SKILL_DIR = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = SKILL_DIR / "templates"
CONFIG_DIR = SKILL_DIR / "config"

XLSM_REL_PATHS = (
    "hwsw/HWSWList_Template.xlsm",
    "poam/POAM_Export_Template.xlsm",
    "ppsm/PPSMBoundariesInformationExport_Template.xlsm",
    "sctm/ControlInfoExport_Template.xlsm",
)

EXPECTED_DATA_VALIDATION_COUNTS = {
    "hwsw/HWSWList_Template.xlsm": {"Hardware": 13, "Software": 14},
    "poam/POAM_Export_Template.xlsm": {"POA&M": 28, "Example": 26},
    "ppsm/PPSMBoundariesInformationExport_Template.xlsm": {"PPSM": 7},
    "sctm/ControlInfoExport_Template.xlsm": {"Template": 34, "Example": 28},
}

WORKSTATION_PATH_PATTERNS = (
    b"/Users/",
    b"/home/",
    b"C:\\Users\\",
    b"\\Users\\",
    b"AppData\\Local",
    b"Library/Containers",
    b".exd",
)

FORBIDDEN_ZIP_PART_SUBSTRINGS = (
    "docmetadata/",
    "labelinfo.xml",
    "customxml/",
    "docprops/custom.xml",
)

ALLOWED_PUBLIC_DOMAINS = frozenset({
    "example.com",
    "example.org",
    "example.net",
    "agency.example",
    "nist.gov",
    "csrc.nist.gov",
    "dodcio.defense.gov",
    "ise.gov",
    "www.ise.gov",
})

ALLOWED_ASSIGNED_FROM_TIERS = frozenset({
    "Enterprise",
    "DoW Enterprise",
    "[r5] Privacy (DoW - Risk Level: Low)",
})

MANGLED_CITATION_RE = re.compile(
    r"(?:DoDI|DoDD|DoDM|NIST\s+SP\s+[0-9-]+),\s+(?:U\.S\.\s+|The\s+)?Enterprise\s+"
    r"|NIST\s+SP\s+800-53,\s+System\s+Maintenance\s+Policy"
    r"|\b(?:Chapters\s+5\s+and\s+7|sections\s+3-25|OMB\s+20A-130|DDoDI|DoWI|Organizationsand|SEC\s+Rule\s+17a-4|DA-PAM|DA\s+PAM)\b"
    r"|DoDI\s+8530\.01\(3\)|DODI\s+8500\.01-\s*ISSM"
)

STAFF_SECTION_RE = re.compile(r"\b(?:CIO/|J-6|G-6|G-2|G-3|S-6)\b")
SAMPLE_IDENTITY_RE = re.compile(
    r"(?:\b(?:Jane\s+Doe|John\s+Smith|John\s+Doe|Admin\s+Cert)\b|\b\d{3}-\d{3}-\d{4}\b)",
    re.IGNORECASE,
)
EMAIL_RE = re.compile(r"\b([a-zA-Z0-9._%+-]+)@([a-zA-Z0-9.-]+\.[a-zA-Z]{2,})\b")
HOSTNAME_RE = re.compile(
    r"\b([a-zA-Z0-9-]+(?:\.[a-zA-Z0-9-]+)+\.(?:com|org|net|gov|mil|edu|int|io|us|uk|example|eit|local|internal|corp))\b",
    re.IGNORECASE,
)
CANONICAL_LIST_VALUE_RE = re.compile(r"^[A-Za-z0-9 /().,&-]+$")


def decompress_ovba_stream(compressed: bytes) -> bytes:
    """Decompresses an MS-OVBA compressed stream (e.g., VBA/dir) using pure Python."""
    if not compressed or compressed[0] != 0x01:
        return b""
    out = bytearray()
    pos = 1
    n = len(compressed)
    while pos + 2 <= n:
        header = compressed[pos] | (compressed[pos + 1] << 8)
        pos += 2
        chunk_size = (header & 0x0FFF) + 3
        chunk_is_compressed = (header & 0x8000) != 0
        chunk_end = min(pos + chunk_size - 2, n)
        if not chunk_is_compressed:
            out.extend(compressed[pos:pos + 4096])
            pos += 4096
            continue
        chunk_start_out = len(out)
        while pos < chunk_end:
            flags = compressed[pos]
            pos += 1
            for bit in range(8):
                if pos >= chunk_end:
                    break
                if (flags & (1 << bit)) == 0:
                    out.append(compressed[pos])
                    pos += 1
                else:
                    if pos + 2 > chunk_end:
                        pos = chunk_end
                        break
                    token = compressed[pos] | (compressed[pos + 1] << 8)
                    pos += 2
                    decompressed_chunk_len = len(out) - chunk_start_out
                    bit_count = max(
                        4,
                        math.ceil(math.log2(max(decompressed_chunk_len, 1))),
                    )
                    if (1 << (bit_count - 1)) < decompressed_chunk_len:
                        bit_count = max(4, int(math.ceil(math.log2(decompressed_chunk_len))))
                    len_mask = (1 << (16 - bit_count)) - 1
                    length = (token & len_mask) + 3
                    offset = (token >> (16 - bit_count)) + 1
                    if offset <= 0 or offset > len(out):
                        return b""
                    for _ in range(length):
                        out.append(out[-offset])
    return bytes(out)


def find_workstation_path_violations(xlsm_bytes: bytes) -> list[str]:
    """Scans an .xlsm archive for workstation paths, HelpFile paths, or disallowed parts."""
    violations: list[str] = []
    with zipfile.ZipFile(io.BytesIO(xlsm_bytes), "r") as zf:
        for name in zf.namelist():
            low_name = name.lower()
            for forbidden_part in FORBIDDEN_ZIP_PART_SUBSTRINGS:
                if forbidden_part in low_name:
                    violations.append(f"disallowed archive part: {name}")

        for member in zf.namelist():
            if not (member == "xl/vbaProject.bin" or member.endswith((".xml", ".rels"))):
                continue
            raw = zf.read(member)
            for pat in WORKSTATION_PATH_PATTERNS:
                if pat in raw:
                    violations.append(f"{member}: contains {pat!r}")
                if member == "xl/vbaProject.bin":
                    utf16_pat = pat.decode("ascii").encode("utf-16le")
                    if utf16_pat in raw:
                        violations.append(f"{member}: contains UTF-16LE {pat!r}")

        if "xl/vbaProject.bin" in zf.namelist():
            vba_raw = zf.read("xl/vbaProject.bin")
            for m in re.finditer(rb'HelpFile="([^"]*)"', vba_raw):
                if m.group(1):
                    violations.append(f"xl/vbaProject.bin: non-empty HelpFile={m.group(1)!r}")
            for m in re.finditer(rb"\x01[\x00-\xff][\xb0-\xbf]", vba_raw):
                decomp = decompress_ovba_stream(vba_raw[m.start():m.start() + 4096])
                if decomp:
                    for pat in WORKSTATION_PATH_PATTERNS:
                        if pat in decomp:
                            violations.append(f"decompressed VBA stream: contains {pat!r}")
    return violations


def validate_hwsw_lists_values(values_by_col: dict[str, list[str]]) -> list[str]:
    """Validates HWSW (U) Lists columns against structural format and reference mappings."""
    errors: list[str] = []
    ref_map_path = CONFIG_DIR / "reference_mappings.json"
    ref_map = json.loads(ref_map_path.read_text(encoding="utf-8"))

    for col_name, items in values_by_col.items():
        if not items:
            errors.append(f"{col_name}: empty list")
        if len(items) > 100:
            errors.append(f"{col_name}: excessive row count ({len(items)})")
        for val in items:
            if not CANONICAL_LIST_VALUE_RE.match(val):
                errors.append(f"{col_name}: non-canonical value {val!r}")
            if HOSTNAME_RE.search(val):
                errors.append(f"{col_name}: contains hostname in {val!r}")

    hw_set = set(values_by_col.get("Hardware Type", []))
    sw_set = set(values_by_col.get("Software Type", []))
    for mapped_hw in ref_map.get("hardware_type_exact_map", {}).values():
        if mapped_hw not in hw_set:
            errors.append(f"Hardware Type missing mapped value {mapped_hw!r}")
    for mapped_sw in ref_map.get("software_type_exact_map", {}).values():
        if mapped_sw not in sw_set:
            errors.append(f"Software Type missing mapped value {mapped_sw!r}")
    return errors


def find_identity_and_domain_violations(text: str, context: str) -> list[str]:
    """Checks cell text for sample identities, non-placeholder emails, and unapproved domains."""
    issues: list[str] = []
    m_id = SAMPLE_IDENTITY_RE.search(text)
    if m_id:
        issues.append(f"{context}: sample identity {m_id.group(0)!r}")

    for m_email in EMAIL_RE.finditer(text):
        issues.append(f"{context}: email address {m_email.group(0)!r}")

    for m_host in HOSTNAME_RE.finditer(text):
        host = m_host.group(1).lower().rstrip(".")
        if host not in ALLOWED_PUBLIC_DOMAINS:
            issues.append(f"{context}: unapproved domain {host!r}")
    return issues


class TestXlsmTemplateHygiene(unittest.TestCase):
    """Verifies .xlsm templates are structurally sanitized, warning-free, and hash-verified."""

    def test_no_workstation_paths_generic(self) -> None:
        """Verifies OLE streams, XML parts, and archive structure contain no workstation paths."""
        for rel in XLSM_REL_PATHS:
            fpath = TEMPLATES_DIR / rel
            self.assertTrue(fpath.is_file(), f"Missing template: {fpath}")
            violations = find_workstation_path_violations(fpath.read_bytes())
            self.assertEqual(
                violations,
                [],
                f"Workstation path or archive hygiene violations in {rel}: {violations}",
            )

        sample_bytes = (TEMPLATES_DIR / XLSM_REL_PATHS[0]).read_bytes()
        buf = io.BytesIO()
        injected_helpfile = b'HelpFile="' + WORKSTATION_PATH_PATTERNS[0] + b'testuser/vba.hlp"'
        with zipfile.ZipFile(io.BytesIO(sample_bytes), "r") as zin, zipfile.ZipFile(
            buf, "w", zipfile.ZIP_DEFLATED
        ) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "xl/vbaProject.bin":
                    data = data.replace(b'HelpFile=""', injected_helpfile)
                zout.writestr(item, data)
        neg_violations = find_workstation_path_violations(buf.getvalue())
        self.assertTrue(
            any("HelpFile" in v for v in neg_violations),
            f"Expected HelpFile violation on negative fixture, got: {neg_violations}",
        )
        self.assertGreaterEqual(len(neg_violations), 2)

    def test_vba_project_and_docprops_free_of_personal_workstation_paths(self) -> None:
        """Alias verifying VBA project streams and document properties are clean."""
        self.test_no_workstation_paths_generic()

    def test_lists_allowlist(self) -> None:
        """Verifies HWSW (U) Lists contains only canonical category values and rejects bad entries."""
        fpath = TEMPLATES_DIR / "hwsw/HWSWList_Template.xlsm"
        wb = openpyxl.load_workbook(fpath, data_only=True, keep_vba=True)
        try:
            ws = wb["(U) Lists"]
            col_map = {
                "Hardware Type": 1,
                "Software Type": 3,
                "Approval": 5,
                "Yes Or No": 7,
            }
            values_by_col: dict[str, list[str]] = {}
            for title, col_idx in col_map.items():
                self.assertEqual(ws.cell(1, col_idx).value, title)
                col_vals: list[str] = []
                for r in range(2, ws.max_row + 1):
                    v = ws.cell(r, col_idx).value
                    if v is not None and str(v).strip():
                        col_vals.append(str(v).strip())
                values_by_col[title] = col_vals

            errors = validate_hwsw_lists_values(values_by_col)
            self.assertEqual(errors, [], f"HWSW (U) Lists validation errors: {errors}")

            corrupted = dict(values_by_col)
            corrupted["Hardware Type"] = list(values_by_col["Hardware Type"]) + [
                "node01.internal.corp"
            ]
            neg_errors = validate_hwsw_lists_values(corrupted)
            self.assertNotEqual(neg_errors, [], "Negative fixture for (U) Lists must fail")
        finally:
            wb.close()

    def test_no_header_footer_markings(self) -> None:
        """Verifies no worksheet across any .xlsm template contains header or footer text."""
        ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        for rel in XLSM_REL_PATHS:
            fpath = TEMPLATES_DIR / rel
            with zipfile.ZipFile(fpath, "r") as zf:
                for name in zf.namelist():
                    if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
                        continue
                    root = ET.fromstring(zf.read(name))
                    for hf in root.findall(".//m:headerFooter", ns):
                        for child in list(hf):
                            text = (child.text or "").strip()
                            self.assertEqual(
                                text,
                                "",
                                f"Non-empty header/footer element <{child.tag}> in {rel}:{name}: {text!r}",
                            )

    def test_no_army_hierarchy_strings(self) -> None:
        """Verifies SCTM Assignment Values uses only generic baseline tiers and valid citations."""
        fpath = TEMPLATES_DIR / "sctm/ControlInfoExport_Template.xlsm"
        wb = openpyxl.load_workbook(fpath, data_only=False, keep_vba=True)
        try:
            ws = wb["Assignment Values"]
            for r in range(6, ws.max_row + 1):
                assigned_from = str(ws.cell(r, 9).value or "").strip()
                assignment_val = str(ws.cell(r, 10).value or "").strip()
                self.assertIn(
                    assigned_from,
                    ALLOWED_ASSIGNED_FROM_TIERS,
                    f"Row {r} in Assignment Values has non-generic Assigned From tier: {assigned_from!r}",
                )
                self.assertIsNone(
                    STAFF_SECTION_RE.search(assignment_val),
                    f"Row {r} in Assignment Values contains staff-section designator: {assignment_val!r}",
                )
                self.assertIsNone(
                    MANGLED_CITATION_RE.search(assignment_val),
                    f"Row {r} in Assignment Values contains mangled citation: {assignment_val!r}",
                )
                self.assertEqual(
                    assignment_val.count("("),
                    assignment_val.count(")"),
                    f"Row {r} in Assignment Values has unbalanced parentheses: {assignment_val!r}",
                )
                self.assertEqual(
                    assignment_val.count('"') % 2,
                    0,
                    f"Row {r} in Assignment Values has unbalanced quotes: {assignment_val!r}",
                )
        finally:
            wb.close()

    def test_no_sample_identities(self) -> None:
        """Verifies no sample identities, emails, or unapproved domains exist in .xlsm XML parts."""
        all_issues: list[str] = []
        for rel in XLSM_REL_PATHS:
            fpath = TEMPLATES_DIR / rel
            with zipfile.ZipFile(fpath, "r") as zf:
                for name in zf.namelist():
                    if not (
                        name == "xl/sharedStrings.xml"
                        or (name.startswith("xl/worksheets/") and name.endswith(".xml"))
                    ):
                        continue
                    root = ET.fromstring(zf.read(name))
                    for text in root.itertext():
                        cleaned = (text or "").strip()
                        if cleaned:
                            all_issues.extend(
                                find_identity_and_domain_violations(cleaned, f"{rel}:{name}")
                            )
        self.assertEqual(all_issues, [], f"Identity or domain violations found: {all_issues}")

        self.assertNotEqual(
            find_identity_and_domain_violations("Contact Jane Doe at 555-0199", "fixture"),
            [],
        )
        self.assertNotEqual(
            find_identity_and_domain_violations("host.sub.internal", "fixture"),
            [],
        )

    def test_worksheets_and_shared_strings_free_of_non_public_or_army_strings(self) -> None:
        """Alias running header/footer, hierarchy/citation, and sample-identity checks."""
        self.test_no_header_footer_markings()
        self.test_no_army_hierarchy_strings()
        self.test_no_sample_identities()

    def test_data_validations_preserved(self) -> None:
        """Verifies all 150 dataValidation rules and sqref ranges are preserved without warnings."""
        total_dvs = 0
        for rel in XLSM_REL_PATHS:
            fpath = TEMPLATES_DIR / rel
            with warnings.catch_warnings(record=True) as caught:
                warnings.simplefilter("always")
                wb = openpyxl.load_workbook(fpath, data_only=False, keep_vba=True)
            try:
                ext_warnings = [
                    str(w.message)
                    for w in caught
                    if "Data Validation extension is not supported" in str(w.message)
                ]
                self.assertEqual(
                    ext_warnings,
                    [],
                    f"Unexpected openpyxl Data Validation extension warning on {rel}: {ext_warnings}",
                )

                expected_sheets = EXPECTED_DATA_VALIDATION_COUNTS[rel]
                for sheet_name, expected_count in expected_sheets.items():
                    ws = wb[sheet_name]
                    dvs = ws.data_validations.dataValidation
                    self.assertEqual(
                        len(dvs),
                        expected_count,
                        f"{rel}:{sheet_name} expected {expected_count} dataValidations, got {len(dvs)}",
                    )
                    for dv in dvs:
                        self.assertTrue(str(dv.sqref).strip(), f"Empty sqref in {rel}:{sheet_name}")
                    total_dvs += len(dvs)

                if rel == "poam/POAM_Export_Template.xlsm":
                    with tempfile.TemporaryDirectory() as tmpdir:
                        out_file = Path(tmpdir) / "poam_roundtrip.xlsm"
                        wb.save(out_file)
                        wb2 = openpyxl.load_workbook(out_file, data_only=False, keep_vba=True)
                        try:
                            dvs_poam = wb2["POA&M"].data_validations.dataValidation
                            dvs_example = wb2["Example"].data_validations.dataValidation
                            self.assertEqual(len(dvs_poam), 28)
                            self.assertEqual(len(dvs_example), 26)
                            formulas = {dv.formula1 for dv in dvs_poam if dv.type == "list"}
                            self.assertIn("Sheet1!$B$1:$B$6", formulas)
                            self.assertIn("Sheet1!$D$1:$D$6", formulas)
                            self.assertIn("Sheet1!$H$1:$H$6", formulas)
                        finally:
                            wb2.close()
            finally:
                wb.close()
        self.assertEqual(total_dvs, 150)

    def test_poam_and_all_templates_load_and_roundtrip_without_data_validation_warnings(self) -> None:
        """Alias verifying openpyxl load/save preserves all dropdown validations."""
        self.test_data_validations_preserved()

    def test_provenance_documents_boilerplate_policy(self) -> None:
        """Verifies PROVENANCE.md documents the boilerplate policy and makes no copyright claim."""
        prov_path = TEMPLATES_DIR / "PROVENANCE.md"
        self.assertTrue(prov_path.is_file(), "PROVENANCE.md must exist")
        prov_text = prov_path.read_text(encoding="utf-8")
        low_text = prov_text.lower()

        self.assertIsNone(
            re.search(r"17\s+u\.s\.c\.|public[\s-]+domain|unmodified\s+upstream", low_text)
        )
        self.assertIn("skip boilerplate check", low_text)

    def test_provenance_md_sha256_checksums_match_templates(self) -> None:
        """Verifies PROVENANCE.md checksums and byte sizes match the .xlsm files on disk."""
        prov_path = TEMPLATES_DIR / "PROVENANCE.md"
        self.assertTrue(prov_path.is_file(), "PROVENANCE.md must exist")
        prov_text = prov_path.read_text(encoding="utf-8")

        for rel in XLSM_REL_PATHS:
            fpath = TEMPLATES_DIR / rel
            raw = fpath.read_bytes()
            actual_sha = hashlib.sha256(raw).hexdigest()
            self.assertIn(
                actual_sha,
                prov_text,
                f"SHA-256 for {rel} ({actual_sha}) not found in PROVENANCE.md",
            )
            self.assertIn(
                str(len(raw)),
                prov_text,
                f"Byte size for {rel} ({len(raw)}) not found in PROVENANCE.md",
            )


if __name__ == "__main__":
    unittest.main()
