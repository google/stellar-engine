# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Characterization snapshots for inventory, HW/SW + PPSM records and OSCAL requirements.

Snapshots pin current behavior so refactors are detected. A snapshot may only be
regenerated in a commit that cites the work item intending the change:

    COMPLIANCE_UPDATE_SNAPSHOTS=1 python -m pytest tests/test_characterization_snapshots.py
"""

import contextlib
import datetime
import json
import os
from pathlib import Path
import re
import shutil
import tempfile
from typing import Any, Dict, Iterator, List
import unittest
from unittest import mock

import yaml

from compliance_engine import excel_hydrator
from compliance_engine import extract_system_data
from compliance_engine import file_helpers
from compliance_engine import generate_compliance_artifacts
from compliance_engine import oscal_generator

_TESTS_DIR = Path(__file__).resolve().parent
_SNAPSHOT_DIR = _TESTS_DIR / "fixtures" / "snapshots"
_REPO_ROOT = _TESTS_DIR.parents[3]
_UPDATE = os.environ.get("COMPLIANCE_UPDATE_SNAPSHOTS") == "1"

# (snapshot name, repo-relative blueprint path or None for the synthetic fixture,
#  path of the copy inside the mirror)
_TARGETS = (
    ("il5_bigquery", "blueprints/il5/bigquery", "blueprints/il5/bigquery"),
    ("il5_gke_hardened", "blueprints/il5/gke-hardened", "blueprints/il5/gke-hardened"),
    ("fedramp_high_cloud_run", "blueprints/fedramp-high/cloud-run", "blueprints/fedramp-high/cloud-run"),
    ("synthetic_il5", None, "blueprints/il5/synthetic"),
)

_DATE_FORMATS = ("%Y-%m-%d", "%B %d, %Y", "%m/%d/%Y", "%Y%m%d")


def _date_tokens() -> Dict[str, str]:
    """Maps date strings derived from today (offsets up to two years) to stable tokens."""
    today = datetime.date.today()
    tokens: Dict[str, str] = {}
    for offset in range(-1, 800):
        day = today + datetime.timedelta(days=offset)
        for fmt in _DATE_FORMATS:
            tokens.setdefault(day.strftime(fmt), f"<TODAY{offset:+d}>")
    return tokens


_DATE_RE = re.compile(r"\d{4}-\d{2}-\d{2}|[A-Z][a-z]+ \d{2}, \d{4}|\d{2}/\d{2}/\d{4}|\b\d{8}\b")
_TIME_RE = re.compile(r"T\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:\d{2})?")


def _normalize(value: Any, root: str, tokens: Dict[str, str]) -> Any:
    """Replaces the temporary root path and today-relative dates with stable tokens."""
    if isinstance(value, dict):
        return {str(k): _normalize(v, root, tokens) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_normalize(v, root, tokens) for v in value]
    if isinstance(value, (datetime.date, datetime.datetime)):
        value = value.isoformat()
    if isinstance(value, str):
        value = value.replace(root, "<ROOT>")
        value = _TIME_RE.sub("T<TIME>", value)
        return _DATE_RE.sub(lambda m: tokens.get(m.group(0), m.group(0)), value)
    return value


@contextlib.contextmanager
def _mirror(blueprint: Any, rel_copy: str) -> Iterator[Path]:
    """Copies a blueprint and the repo modules/ tree into a temporary git-like root."""
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp).resolve()
        (root / ".git").mkdir()
        shutil.copytree(_REPO_ROOT / "modules", root / "modules")
        dest = root / rel_copy
        if blueprint is None:
            shutil.copytree(_SNAPSHOT_DIR / "synthetic", dest)
        else:
            shutil.copytree(_REPO_ROOT / blueprint, dest, symlinks=False)
        cwd = os.getcwd()
        empty_cwd = root / "empty_cwd"
        empty_cwd.mkdir()
        os.chdir(empty_cwd)
        try:
            yield dest
        finally:
            os.chdir(cwd)


def _cell_value(value: Any) -> Any:
    """Renders formula objects (which lack value equality) as comparable text."""
    text = getattr(value, "text", None)
    if text is not None and not isinstance(value, str):
        return f"={text} @{getattr(value, 'ref', '')}"
    return value


def _workbook_cell_map(template: Path, output: Path) -> Dict[str, Dict[str, Any]]:
    """Returns every cell whose value differs from the template, keyed by sheet and coordinate."""
    import openpyxl  # pylint: disable=import-outside-toplevel

    tpl = openpyxl.load_workbook(template, keep_vba=True)
    out = openpyxl.load_workbook(output, keep_vba=True)
    cells: Dict[str, Dict[str, Any]] = {}
    for ws in out.worksheets:
        base = tpl[ws.title] if ws.title in tpl.sheetnames else None
        changed: Dict[str, Any] = {}
        for row in ws.iter_rows():
            for cell in row:
                before = _cell_value(base[cell.coordinate].value) if base is not None else None
                after = _cell_value(cell.value)
                if after != before:
                    changed[cell.coordinate] = after
        if changed:
            cells[ws.title] = changed
    tpl.close()
    out.close()
    return cells


def _build_snapshot(blueprint: Any, rel_copy: str) -> Dict[str, Any]:
    """Runs extraction and the record builders against a mirrored target."""
    env = {"COMPLIANCE_DISABLE_LIVE_SCANNERS": "1", "COMPLIANCE_OFFLINE": "1"}
    with mock.patch.dict(os.environ, env), _mirror(blueprint, rel_copy) as target:
        root = str(target.parents[len(Path(rel_copy).parts) - 1])
        inventory = extract_system_data.extract_system_inventory(str(target))
        hwsw_yaml = yaml.safe_load(generate_compliance_artifacts.generate_hwsw_inventory_yaml(inventory))
        ppsm_yaml = yaml.safe_load(generate_compliance_artifacts.generate_ppsm_matrix_yaml(inventory))
        templates = file_helpers.get_templates_dir()
        workbooks: Dict[str, Any] = {}
        for key, cls, tpl in (
            ("hwsw", excel_hydrator.HWSWHydrator, templates / "hwsw" / "HWSWList_Template.xlsm"),
            ("ppsm", excel_hydrator.PPSMHydrator, templates / "ppsm" / "PPSMBoundariesInformationExport_Template.xlsm"),
        ):
            out_path = target / f"{key}.xlsm"
            cls(str(tpl)).hydrate(inventory, str(out_path))
            workbooks[key] = _workbook_cell_map(tpl, out_path)
        ssp = oscal_generator.generate_oscal_ssp(inventory)["system-security-plan"]
        requirements: List[Dict[str, Any]] = []
        for req in ssp["control-implementation"]["implemented-requirements"]:
            requirements.append({
                "control-id": req.get("control-id"),
                "by-components": [
                    {
                        "state": bc.get("implementation-status", {}).get("state"),
                        "description": bc.get("description"),
                    }
                    for bc in req.get("by-components", [])
                ],
            })
        snapshot = {
            "inventory": inventory,
            "hwsw_yaml": hwsw_yaml,
            "ppsm_yaml": ppsm_yaml,
            "workbooks": workbooks,
            "oscal": {"import_profile": ssp["import-profile"]["href"], "implemented_requirements": requirements},
        }
        return _normalize(snapshot, root, _date_tokens())


class CharacterizationSnapshotTest(unittest.TestCase):
    """Compares current engine output with the recorded snapshots."""

    maxDiff = None
    _cache: Dict[str, Dict[str, Any]] = {}

    @classmethod
    def _snapshot(cls, name: str, blueprint: Any, rel_copy: str) -> Dict[str, Any]:
        if name not in cls._cache:
            cls._cache[name] = _build_snapshot(blueprint, rel_copy)
        return cls._cache[name]

    def _check(self, section: str) -> None:
        for name, blueprint, rel_copy in _TARGETS:
            with self.subTest(target=name, section=section):
                actual = json.loads(json.dumps(self._snapshot(name, blueprint, rel_copy)[section], sort_keys=True, default=str))
                path = _SNAPSHOT_DIR / f"{name}.{section}.json"
                if _UPDATE:
                    path.write_text(json.dumps(actual, indent=1, sort_keys=True) + "\n", encoding="utf-8")
                    continue
                self.assertTrue(path.is_file(), f"missing snapshot {path}; regenerate with COMPLIANCE_UPDATE_SNAPSHOTS=1")
                expected = json.loads(path.read_text(encoding="utf-8"))
                self.assertEqual(expected, actual, f"{path.name} changed; regenerate only in a commit citing the intended work item")

    def test_inventory_snapshot(self) -> None:
        self._check("inventory")

    def test_hwsw_ppsm_records_snapshot(self) -> None:
        self._check("hwsw_yaml")
        self._check("ppsm_yaml")
        self._check("workbooks")

    def test_oscal_requirements_snapshot(self) -> None:
        self._check("oscal")


if __name__ == "__main__":
    unittest.main()
