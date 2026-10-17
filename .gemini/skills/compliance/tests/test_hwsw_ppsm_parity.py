# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""YAML and XLSM hardware/PPSM exports report the same values."""

from pathlib import Path
import tempfile
import unittest

import openpyxl
import yaml

from compliance_engine import excel_hydrator
from compliance_engine import generate_compliance_artifacts

_TEMPLATES = Path(__file__).resolve().parents[1] / "templates"
_NOT_DETERMINED = "[NOT DETERMINED FROM SOURCE]"
# Hardware sheet columns (1-based) for the fields compared against YAML keys.
_HW_COLUMNS = {"asset_ip_address": 6, "manufacturer": 12, "model_number": 13, "serial_number": 14, "os_fw_version": 17, "location": 19}


def _inventory():
    return {
        "system_information": {"system_name": "Demo", "system_abbreviation": "DMO", "impact_level": "IL5", "emass_id": "1234"},
        "infrastructure_components": {
            "gke_clusters": ["plain-cluster", {"name": "gke-a"}, {"name": "gke-b", "master_version": "1.30.2"}],
            "compute_instances": [{"name": "vm-a", "zone": None, "image": "projects/debian-cloud/global/images/family/debian-12", "machine_type": "e2-medium"}],
            "kms_keys": [{"name": "key-hsm", "protection_level": "HSM"}, {"name": "key-sw", "protection_level": "SOFTWARE"}],
        },
        "network_architecture": {},
    }


class HardwareParityTest(unittest.TestCase):
    """Every hardware record carries identical field values in both formats."""

    @classmethod
    def setUpClass(cls) -> None:
        inv = _inventory()
        cls.yaml_rows = {r["nickname"]: r for r in yaml.safe_load(generate_compliance_artifacts.generate_hwsw_inventory_yaml(inv))["hardware_assets"]}
        cls.tmp = tempfile.TemporaryDirectory()
        out = str(Path(cls.tmp.name) / "hwsw.xlsm")
        excel_hydrator.HWSWHydrator(str(_TEMPLATES / "hwsw" / "HWSWList_Template.xlsm")).hydrate(inv, out)
        ws = openpyxl.load_workbook(out, keep_vba=True)["Hardware"]
        cls.xlsm_rows = {}
        for r in range(8, ws.max_row + 1):
            nick = ws.cell(row=r, column=4).value
            if nick:
                cls.xlsm_rows[str(nick)] = {k: ws.cell(row=r, column=c).value for k, c in _HW_COLUMNS.items()}

        ppsm_out = str(Path(cls.tmp.name) / "ppsm.xlsm")
        excel_hydrator.PPSMHydrator(str(_TEMPLATES / "ppsm" / "PPSMBoundariesInformationExport_Template.xlsm")).hydrate(inv, ppsm_out)
        cls.ppsm_ws = openpyxl.load_workbook(ppsm_out, keep_vba=True)["PPSM"]
        cls.ppsm_yaml = yaml.safe_load(generate_compliance_artifacts.generate_ppsm_matrix_yaml(inv))

    @classmethod
    def tearDownClass(cls) -> None:
        cls.tmp.cleanup()

    def test_same_records(self) -> None:
        self.assertEqual(set(self.yaml_rows), set(self.xlsm_rows))

    def test_field_values_match(self) -> None:
        for nick, xrow in self.xlsm_rows.items():
            yrow = self.yaml_rows[nick]
            for key, xval in xrow.items():
                if key in yrow or xval not in (None, "N/A"):
                    with self.subTest(asset=nick, field=key):
                        self.assertEqual(str(yrow.get(key)), str(xval))

    def test_tenant_ip_not_invented(self) -> None:
        self.assertEqual(self.xlsm_rows["GCP Cloud Tenant (DMO)"]["asset_ip_address"], _NOT_DETERMINED)

    def test_vm_os_and_zone(self) -> None:
        vm = self.yaml_rows["vm-a"]
        self.assertEqual(vm["os_fw_version"], "debian-12")
        self.assertNotIn("None", vm["serial_number"])

    def test_gke_version_not_invented(self) -> None:
        self.assertIn(_NOT_DETERMINED, self.yaml_rows["gke-a"]["model_number"])
        self.assertIn("1.30.2", self.yaml_rows["gke-b"]["os_fw_version"])
        self.assertNotIn("1.28", str(self.xlsm_rows["plain-cluster"]))

    def test_ppsm_emass_id_alias(self) -> None:
        self.assertEqual(self.ppsm_yaml["system_metadata"]["emass_system_id"], "1234")
        self.assertEqual(str(self.ppsm_ws["C6"].value), "1234")

    def test_single_builder_consolidation_and_truthful_db_hsm_oscal(self) -> None:
        """Shared builders: shared builders exist, unset DB version is not PostgreSQL 15, no LiquidSecurity or Andromeda."""
        from compliance_engine import hardware_records
        from compliance_engine import oscal_generator

        self.assertTrue(hasattr(hardware_records, "build_hwsw_inventory_records"))
        self.assertTrue(hasattr(hardware_records, "build_ppsm_records"))

        # HSM vendor uses Marvell (formerly Cavium), not undocumented Marvell LiquidSecurity
        self.assertNotIn("LiquidSecurity", self.yaml_rows["key-hsm"]["manufacturer"])
        self.assertIn("Marvell", self.yaml_rows["key-hsm"]["manufacturer"])

        # Unset Cloud SQL version must not invent PostgreSQL 15
        inv_db = {
            "system_information": {"system_name": "DbSys", "system_abbreviation": "DBS", "impact_level": "IL5"},
            "infrastructure_components": {
                "databases": [{"name": "sql-unknown-ver", "type": "google_sql_database_instance"}],
            },
            "network_architecture": {},
        }
        db_yaml = yaml.safe_load(generate_compliance_artifacts.generate_hwsw_inventory_yaml(inv_db))
        db_row = next(r for r in db_yaml["hardware_assets"] if r["nickname"] == "sql-unknown-ver")
        self.assertNotIn("PostgreSQL 15", db_row["os_fw_version"])
        self.assertIn(_NOT_DETERMINED, db_row["os_fw_version"])

        # OSCAL VPC component description must not mention internal Google 'Andromeda' SDN name
        comp_def = oscal_generator.generate_oscal_component_definition(inv_db)
        self.assertNotIn("Andromeda", str(comp_def))


class PpsmParityTest(unittest.TestCase):
    """PPSM rows carry the same record, protocol, port and boundary in both formats."""

    def test_rows_match(self) -> None:
        inv = {
            "system_information": {"system_name": "Demo", "system_abbreviation": "DMO", "impact_level": "IL5"},
            "infrastructure_components": {"services_enabled": ["run.googleapis.com", "storage.googleapis.com"]},
            "network_architecture": {"firewall_rules": [{"name": "allow-https", "action": "allow", "direction": "INGRESS",
                                                         "allow": [{"protocol": "tcp", "ports": ["443"]}], "source_ranges": ["0.0.0.0/0"]}]},
            "application_components": {},
        }
        yaml_rows = [(r["record_name"], r["protocol"], str(r["port"]), r["boundary"])
                     for r in yaml.safe_load(generate_compliance_artifacts.generate_ppsm_matrix_yaml(inv))["ppsm_matrix"]]
        with tempfile.TemporaryDirectory() as tmp:
            out = str(Path(tmp) / "ppsm.xlsm")
            excel_hydrator.PPSMHydrator(str(_TEMPLATES / "ppsm" / "PPSMBoundariesInformationExport_Template.xlsm")).hydrate(inv, out)
            ws = openpyxl.load_workbook(out, keep_vba=True)["PPSM"]
            xlsm_rows = [tuple(str(ws.cell(row=r, column=c).value) for c in (3, 4, 6, 7))
                         for r in range(9, ws.max_row + 1) if ws.cell(row=r, column=3).value]
        self.assertTrue(yaml_rows)
        self.assertEqual(yaml_rows, xlsm_rows)

    def test_ppsm_vpn_and_civilian_logical_point_from_inventory(self) -> None:
        """Firewall and port PPSM rows reflect actual VPN resources and DoD vs civilian regime."""
        from compliance_engine import hardware_records

        inv_civ_no_vpn = {
            "system_information": {"system_name": "CivSys", "system_abbreviation": "CVS", "impact_level": "FedRAMP High"},
            "infrastructure_components": {"all_resources": [{"type": "google_compute_firewall", "name": "fw-https"}]},
            "network_architecture": {"firewall_rules": [{"name": "fw-https", "protocol": "TCP", "ports": "443", "direction": "INGRESS"}]},
            "application_components": {"exposed_ports": [{"port": "8080", "protocol": "TCP", "service_name": "web"}]},
        }
        rows_no_vpn = hardware_records.build_ppsm_records(inv_civ_no_vpn)
        self.assertEqual(len(rows_no_vpn), 2)
        for row in rows_no_vpn:
            self.assertEqual(row["vpn_encrypted_traffic"], "No")
            self.assertEqual(row["vpn_tunnel_type"], "N/A")
            self.assertEqual(row["logical_source_point"], "Agency Network / Authorized User Boundary")
            self.assertEqual(row["logical_destination_point"], "Agency Network / Authorized User Boundary")

        inv_dod_with_vpn = {
            "system_information": {"system_name": "DodSys", "system_abbreviation": "DDS", "impact_level": "IL5"},
            "infrastructure_components": {
                "all_resources": [{"type": "google_compute_vpn_tunnel", "name": "tun-1"}],
            },
            "network_architecture": {"firewall_rules": [{"name": "fw-https", "protocol": "TCP", "ports": "443", "direction": "INGRESS"}]},
            "application_components": {},
        }
        rows_with_vpn = hardware_records.build_ppsm_records(inv_dod_with_vpn)
        self.assertEqual(len(rows_with_vpn), 1)
        self.assertEqual(rows_with_vpn[0]["vpn_encrypted_traffic"], "Yes")
        self.assertEqual(rows_with_vpn[0]["vpn_tunnel_type"], "Cloud Layer 3 VPN")
        self.assertEqual(rows_with_vpn[0]["logical_source_point"], "DoD Enclave (DoD Network)")
        self.assertEqual(rows_with_vpn[0]["logical_destination_point"], "DoD Enclave (DoD Network)")


if __name__ == "__main__":
    unittest.main()
